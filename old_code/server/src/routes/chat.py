"""
聊天路由 — SSE流式端点 + 非流式兼容 + 后台事件轮询。

关键改进（vs项目B v2）：
- SSE流在main response完成后立即关闭，不等后台任务
- 后台事件通过 /api/chat/events 轮询端点获取
"""

from __future__ import annotations

import base64
import csv
import io
import json
import logging
import os
import re
import time
import uuid
import zlib
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import httpx
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

import src.config as cfg
from src.models import (
    ChatRequest,
    ChatResponse,
    TeachingSubmissionRequest,
    TeachingEvaluationRequest,
    TeachingClassCreateRequest,
    TeachingClassMemberAddRequest,
    TeachingAssignmentCreateRequest,
    TeachingAssignmentSubmitRequest,
    TeachingTemplateCreateRequest,
    TeachingClassTemplateInstantiateRequest,
    TeachingAssignmentTemplateApplyRequest,
    TeachingTemplateRollbackRequest,
    TeachingTemplateStatusUpdateRequest,
    TeachingInterventionActionCreateRequest,
    TeachingClassGoalCreateRequest,
    TeachingClassGoalStatusUpdateRequest,
    TeachingGoalTermArchiveRequest,
    TeachingInterventionExperimentPlanRequest,
    TeachingInterventionExperimentStatusUpdateRequest,
)
from src.routes.auth import get_current_user
from src.database import get_db
from src.config import UPLOAD_DIR, UPLOAD_MAX_SIZE_MB

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/chat", tags=["chat"])

# ── 对话历史路由（独立前缀 /api/conversations）──────────────────
conv_router = APIRouter(prefix="/api/conversations", tags=["conversations"])


_RESPONSE_MODE_EXECUTION = "execution"
_RESPONSE_MODE_LEARNING = "learning"
_VALID_RESPONSE_MODES = {_RESPONSE_MODE_EXECUTION, _RESPONSE_MODE_LEARNING}


def _normalize_response_mode_input(raw_mode: str | None) -> tuple[str, bool]:
    mode = str(raw_mode or "").strip().lower()
    if mode in _VALID_RESPONSE_MODES:
        return mode, True
    return _RESPONSE_MODE_EXECUTION, False


def _sanitize_response_mode(raw_mode: str | None) -> str | None:
    mode = str(raw_mode or "").strip().lower()
    return mode if mode in _VALID_RESPONSE_MODES else None


_VALID_ACCOUNT_ROLES = {"general", "teacher", "student"}


def _normalize_account_role(raw_role: str | None) -> str:
    role = str(raw_role or "").strip().lower()
    return role if role in _VALID_ACCOUNT_ROLES else "general"


def _trim_context_text(raw: Any, limit: int) -> str:
    text = str(raw or "").replace("\x00", "").strip()
    if limit > 0 and len(text) > limit:
        return text[:limit]
    return text


_CHAT_MESSAGE_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _normalize_chat_message(raw: Any) -> str:
    text = str(raw or "")
    if not text:
        return ""

    normalized = (
        text
        .replace("\ufeff", "")
        .replace("\u200b", "")
        .replace("\u200c", "")
        .replace("\u200d", "")
        .replace("\xa0", " ")
    )
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
    normalized = _CHAT_MESSAGE_CONTROL_RE.sub("", normalized)
    return normalized.strip()


_PAGE_CONTEXT_ATTACHMENT_NAME = "页面上下文快照"
_PAGE_CONTEXT_ROUTE_LIMIT = 120
_PAGE_CONTEXT_TITLE_LIMIT = 160
_PAGE_CONTEXT_METRIC_LIMIT = 10
_PAGE_CONTEXT_SELECTION_LIMIT = 8
_PAGE_CONTEXT_TEXT_LIMIT = 1600


def _normalize_page_context(raw_context: Any) -> dict[str, Any]:
    if not isinstance(raw_context, dict):
        return {}

    route = _trim_context_text(raw_context.get("route"), _PAGE_CONTEXT_ROUTE_LIMIT)
    page_title = _trim_context_text(raw_context.get("page_title"), _PAGE_CONTEXT_TITLE_LIMIT)
    summary = _trim_context_text(raw_context.get("summary"), 240)
    source = _trim_context_text(raw_context.get("source"), 64)

    metrics: list[str] = []
    raw_metrics = raw_context.get("visible_metrics")
    if isinstance(raw_metrics, list):
        for item in raw_metrics:
            text = _trim_context_text(item, 120)
            if text:
                metrics.append(text)
            if len(metrics) >= _PAGE_CONTEXT_METRIC_LIMIT:
                break

    selections: list[str] = []
    raw_selections = raw_context.get("selection_hints")
    if isinstance(raw_selections, list):
        for item in raw_selections:
            text = _trim_context_text(item, 120)
            if text:
                selections.append(text)
            if len(selections) >= _PAGE_CONTEXT_SELECTION_LIMIT:
                break

    snapshot = _trim_context_text(raw_context.get("snapshot"), _PAGE_CONTEXT_TEXT_LIMIT)
    captured_at = _trim_context_text(raw_context.get("captured_at"), 64)

    if not (route or page_title or summary or metrics or selections or snapshot):
        return {}

    normalized: dict[str, Any] = {}
    if route:
        normalized["route"] = route
    if page_title:
        normalized["page_title"] = page_title
    if summary:
        normalized["summary"] = summary
    if source:
        normalized["source"] = source
    if metrics:
        normalized["visible_metrics"] = metrics
    if selections:
        normalized["selection_hints"] = selections
    if snapshot:
        normalized["snapshot"] = snapshot
    if captured_at:
        normalized["captured_at"] = captured_at

    return normalized


def _build_page_context_attachment(raw_context: Any) -> dict[str, Any] | None:
    context = _normalize_page_context(raw_context)
    if not context:
        return None

    lines = ["# 页面上下文快照"]

    route = str(context.get("route") or "")
    page_title = str(context.get("page_title") or "")
    summary = str(context.get("summary") or "")
    source = str(context.get("source") or "")
    captured_at = str(context.get("captured_at") or "")

    if route:
        lines.append(f"- 路由: {route}")
    if page_title:
        lines.append(f"- 页面标题: {page_title}")
    if summary:
        lines.append(f"- 页面摘要: {summary}")
    if source:
        lines.append(f"- 来源: {source}")
    if captured_at:
        lines.append(f"- 采集时间: {captured_at}")

    metrics = context.get("visible_metrics") if isinstance(context.get("visible_metrics"), list) else []
    if metrics:
        lines.append("\n## 当前可见指标")
        for item in metrics:
            lines.append(f"- {item}")

    selections = context.get("selection_hints") if isinstance(context.get("selection_hints"), list) else []
    if selections:
        lines.append("\n## 当前选中信息")
        for item in selections:
            lines.append(f"- {item}")

    snapshot = str(context.get("snapshot") or "")
    if snapshot:
        lines.append("\n## 页面文本快照")
        lines.append(snapshot)

    content = "\n".join(lines).strip()
    if len(content) > _PAGE_CONTEXT_TEXT_LIMIT:
        content = content[:_PAGE_CONTEXT_TEXT_LIMIT]

    summary_text = page_title or route or _PAGE_CONTEXT_ATTACHMENT_NAME
    return {
        "name": _PAGE_CONTEXT_ATTACHMENT_NAME,
        "type": "page_context",
        "status": "parsed",
        "summary": f"页面上下文：{summary_text}",
        "content": content,
        "parser": "client-page-context",
        "extracted_at": captured_at or datetime.now(timezone.utc).isoformat(),
    }


def _merge_attachments_with_page_context(raw_attachments: Any, raw_page_context: Any) -> list[dict[str, Any]]:
    merged: list[dict[str, Any]] = []

    if isinstance(raw_attachments, list):
        for item in raw_attachments:
            if isinstance(item, dict):
                merged.append(item)

    page_context_attachment = _build_page_context_attachment(raw_page_context)
    if page_context_attachment:
        merged.append(page_context_attachment)

    return merged


async def _resolve_requester_account_role(db: Any, user: dict) -> str:
    token_role = _normalize_account_role(user.get("account_role"))
    if token_role != "general":
        return token_role

    row = await db.execute_fetchone(
        "SELECT account_role FROM users WHERE id = ?",
        (user["id"],),
    )
    if not row:
        return "general"
    raw = row["account_role"] if "account_role" in row.keys() else "general"
    return _normalize_account_role(raw)


def _can_access_teaching_submission(*, requester_id: int, requester_role: str, submission_row: Any) -> bool:
    teacher_id = int(submission_row["teacher_user_id"] or 0)
    student_id = int(submission_row["student_user_id"] or 0)
    if requester_role == "teacher":
        return requester_id == teacher_id
    if requester_role == "student":
        return requester_id == student_id
    return requester_id in {teacher_id, student_id}


async def _fetch_teaching_class(db: Any, class_id: int) -> Any | None:
    return await db.execute_fetchone(
        """
        SELECT id, teacher_user_id, name, description, status, created_at, updated_at
        FROM teaching_classes
        WHERE id = ?
        """,
        (class_id,),
    )


async def _require_teacher_owned_class(db: Any, *, class_id: int, teacher_user_id: int) -> Any:
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")
    if int(class_row["teacher_user_id"] or 0) != int(teacher_user_id):
        raise HTTPException(status_code=403, detail="无权管理该班级")
    return class_row


async def _is_teaching_class_member(db: Any, *, class_id: int, student_user_id: int) -> bool:
    row = await db.execute_fetchone(
        "SELECT 1 FROM teaching_class_members WHERE class_id = ? AND student_user_id = ? LIMIT 1",
        (class_id, student_user_id),
    )
    return bool(row)


_VALID_TEACHING_TEMPLATE_TYPES = {"class", "assignment", "rubric"}
_VALID_TEACHING_TEMPLATE_STATUSES = {"active", "draft", "review", "approved"}
_PUBLISHABLE_TEACHING_TEMPLATE_STATUSES = {"active", "approved"}


def _normalize_teaching_template_status(raw_status: str | None, *, allow_active_alias: bool = True) -> str:
    token = str(raw_status or "").strip().lower()
    if not token:
        return "active" if allow_active_alias else "approved"
    if token == "active" and not allow_active_alias:
        return "approved"
    if token not in _VALID_TEACHING_TEMPLATE_STATUSES:
        raise HTTPException(status_code=400, detail="template_status 仅支持 active/draft/review/approved")
    return token


def _canonical_teaching_template_status(raw_status: str | None) -> str:
    token = _normalize_teaching_template_status(raw_status, allow_active_alias=True)
    return "approved" if token == "active" else token


def _is_teaching_template_publishable(raw_status: str | None) -> bool:
    token = _normalize_teaching_template_status(raw_status, allow_active_alias=True)
    return token in _PUBLISHABLE_TEACHING_TEMPLATE_STATUSES


def _ensure_teaching_template_status_transition(current_status: str | None, target_status: str | None) -> str:
    current = _canonical_teaching_template_status(current_status)
    target = _normalize_teaching_template_status(target_status, allow_active_alias=False)
    if current == target:
        raise HTTPException(status_code=400, detail="模板状态未变化")

    allowed: dict[str, set[str]] = {
        "draft": {"review", "approved"},
        "review": {"draft", "approved"},
        "approved": {"draft", "review"},
    }
    if target not in allowed.get(current, set()):
        raise HTTPException(status_code=400, detail=f"模板状态不允许从 {current} 变更到 {target}")
    return target


def _normalize_teaching_template_type(raw_type: str | None) -> str:
    token = str(raw_type or "").strip().lower()
    if token not in _VALID_TEACHING_TEMPLATE_TYPES:
        raise HTTPException(status_code=400, detail="template_type 仅支持 class/assignment/rubric")
    return token


async def _fetch_teaching_template(db: Any, template_id: int) -> Any | None:
    return await db.execute_fetchone(
        """
        SELECT id, teacher_user_id, template_type, name, description, payload_json, status, created_at, updated_at
        FROM teaching_templates
        WHERE id = ?
        """,
        (template_id,),
    )


async def _require_teacher_owned_template(db: Any, *, template_id: int, teacher_user_id: int) -> Any:
    template_row = await _fetch_teaching_template(db, template_id)
    if not template_row:
        raise HTTPException(status_code=404, detail="模板不存在")
    if int(template_row["teacher_user_id"] or 0) != int(teacher_user_id):
        raise HTTPException(status_code=403, detail="无权管理该模板")
    return template_row


def _parse_teaching_template_payload(payload_json: str | None) -> dict[str, Any]:
    try:
        parsed = json.loads(payload_json or "{}")
    except Exception:
        return {}
    return parsed if isinstance(parsed, dict) else {}


async def _next_teaching_template_version_no(db: Any, *, template_id: int) -> int:
    row = await db.execute_fetchone(
        "SELECT COALESCE(MAX(version_no), 0) + 1 AS next_version_no FROM teaching_template_versions WHERE template_id = ?",
        (template_id,),
    )
    return int(row["next_version_no"] or 1) if row else 1


async def _append_teaching_template_version(
    db: Any,
    *,
    template_row: Any,
    change_type: str,
    change_note: str = "",
    source_version_id: int | None = None,
) -> int:
    template_id = int(template_row["id"] or 0)
    if template_id <= 0:
        raise HTTPException(status_code=500, detail="模板版本记录失败：无效模板ID")

    version_no = await _next_teaching_template_version_no(db, template_id=template_id)
    await db.execute(
        """
        INSERT INTO teaching_template_versions
            (template_id, teacher_user_id, template_type, version_no, change_type, change_note, source_version_id,
             name, description, payload_json, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        """,
        (
            template_id,
            int(template_row["teacher_user_id"] or 0),
            str(template_row["template_type"] or "").strip().lower(),
            version_no,
            str(change_type or "update").strip().lower() or "update",
            str(change_note or "")[:500],
            int(source_version_id) if source_version_id else None,
            str(template_row["name"] or "")[:120],
            str(template_row["description"] or "")[:1000],
            str(template_row["payload_json"] or "{}"),
        ),
    )
    return version_no


def _serialize_teaching_template_row(
    row: Any,
    *,
    latest_version_no: int = 0,
    version_count: int = 0,
    usage_count: int = 0,
    class_usage_count: int = 0,
    assignment_usage_count: int = 0,
    usage_recent_7d: int = 0,
    last_used_at: str = "",
    effect_snapshot: dict[str, Any] | None = None,
) -> dict[str, Any]:
    raw_status = str(row["status"] or "").strip().lower()
    if raw_status not in _VALID_TEACHING_TEMPLATE_STATUSES:
        raw_status = "active"
    canonical_status = _canonical_teaching_template_status(raw_status)
    return {
        "id": row["id"],
        "teacher_user_id": row["teacher_user_id"],
        "template_type": row["template_type"],
        "name": row["name"] or "",
        "description": row["description"] or "",
        "payload": _parse_teaching_template_payload(row["payload_json"]),
        "status": raw_status,
        "status_canonical": canonical_status,
        "is_publishable": _is_teaching_template_publishable(raw_status),
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
        "latest_version_no": int(latest_version_no or 0),
        "version_count": int(version_count or 0),
        "usage_count": int(usage_count or 0),
        "class_usage_count": int(class_usage_count or 0),
        "assignment_usage_count": int(assignment_usage_count or 0),
        "usage_recent_7d": int(usage_recent_7d or 0),
        "last_used_at": str(last_used_at or ""),
        "effect_snapshot": effect_snapshot if isinstance(effect_snapshot, dict) else {},
    }


async def _list_teaching_template_version_meta(
    db: Any,
    *,
    teacher_user_id: int,
) -> dict[int, dict[str, int]]:
    rows = await db.execute_fetchall(
        """
        SELECT template_id, MAX(version_no) AS latest_version_no, COUNT(1) AS version_count
        FROM teaching_template_versions
        WHERE teacher_user_id = ?
        GROUP BY template_id
        """,
        (teacher_user_id,),
    )
    meta: dict[int, dict[str, int]] = {}
    for row in rows:
        tid = int(row["template_id"] or 0)
        if tid <= 0:
            continue
        meta[tid] = {
            "latest_version_no": int(row["latest_version_no"] or 0),
            "version_count": int(row["version_count"] or 0),
        }
    return meta


async def _append_teaching_template_usage_log(
    db: Any,
    *,
    template_row: Any,
    usage_scene: str,
    target_id: int = 0,
    metadata: dict[str, Any] | None = None,
) -> None:
    template_id = int(template_row["id"] or 0)
    teacher_user_id = int(template_row["teacher_user_id"] or 0)
    template_type = str(template_row["template_type"] or "").strip().lower()
    scene = str(usage_scene or "").strip().lower()
    if template_id <= 0 or teacher_user_id <= 0 or not template_type or not scene:
        return

    await db.execute(
        """
        INSERT INTO teaching_template_usage_logs
            (template_id, teacher_user_id, template_type, usage_scene, target_id, metadata_json, created_at)
        VALUES (?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        """,
        (
            template_id,
            teacher_user_id,
            template_type,
            scene,
            int(target_id or 0),
            json.dumps(metadata or {}, ensure_ascii=False),
        ),
    )


async def _list_teaching_template_usage_meta(
    db: Any,
    *,
    teacher_user_id: int,
    template_ids: list[int] | None = None,
) -> dict[int, dict[str, Any]]:
    params: list[Any] = [teacher_user_id]
    sql = """
        SELECT
            template_id,
            COUNT(1) AS usage_count,
            SUM(CASE WHEN usage_scene = 'class_instantiate' THEN 1 ELSE 0 END) AS class_usage_count,
            SUM(CASE WHEN usage_scene = 'assignment_apply' THEN 1 ELSE 0 END) AS assignment_usage_count,
            SUM(CASE WHEN created_at >= datetime('now', '-7 day') THEN 1 ELSE 0 END) AS usage_recent_7d,
            MAX(created_at) AS last_used_at
        FROM teaching_template_usage_logs
        WHERE teacher_user_id = ?
    """

    normalized_ids = sorted({int(x) for x in (template_ids or []) if int(x) > 0})
    if template_ids is not None:
        if not normalized_ids:
            return {}
        placeholders = ",".join("?" * len(normalized_ids))
        sql += f" AND template_id IN ({placeholders})"
        params.extend(normalized_ids)

    sql += " GROUP BY template_id"

    try:
        rows = await db.execute_fetchall(sql, tuple(params))
    except Exception:
        return {}

    meta: dict[int, dict[str, Any]] = {}
    for row in rows:
        tid = int(row["template_id"] or 0)
        if tid <= 0:
            continue
        meta[tid] = {
            "usage_count": int(row["usage_count"] or 0),
            "class_usage_count": int(row["class_usage_count"] or 0),
            "assignment_usage_count": int(row["assignment_usage_count"] or 0),
            "usage_recent_7d": int(row["usage_recent_7d"] or 0),
            "last_used_at": row["last_used_at"] or "",
        }
    return meta


_VALID_INTERVENTION_SEVERITIES = {"high", "medium", "low"}


def _normalize_teaching_intervention_code(raw_code: str | None) -> str:
    token = re.sub(r"[^a-z0-9_-]", "_", str(raw_code or "").strip().lower())
    token = re.sub(r"_+", "_", token).strip("_")
    if not token:
        raise HTTPException(status_code=400, detail="intervention_code 不能为空")
    return token[:80]


def _normalize_teaching_intervention_severity(raw_severity: str | None) -> str:
    token = str(raw_severity or "").strip().lower()
    return token if token in _VALID_INTERVENTION_SEVERITIES else "medium"


def _serialize_teaching_intervention_action_row(row: Any) -> dict[str, Any]:
    assignment_id = int(row["assignment_id"] or 0)
    metadata = _parse_teaching_template_payload(row["metadata_json"])
    return {
        "id": int(row["id"] or 0),
        "class_id": int(row["class_id"] or 0),
        "teacher_user_id": int(row["teacher_user_id"] or 0),
        "intervention_code": str(row["intervention_code"] or ""),
        "intervention_title": str(row["intervention_title"] or ""),
        "intervention_severity": _normalize_teaching_intervention_severity(row["intervention_severity"]),
        "assignment_id": assignment_id,
        "note": str(row["note"] or ""),
        "strategy_variant": _extract_teaching_strategy_variant(metadata),
        "strategy_experiment_id": str(metadata.get("strategy_experiment_id") or ""),
        "strategy_note": str(metadata.get("strategy_note") or ""),
        "metadata": metadata,
        "created_at": row["created_at"] or "",
    }


async def _build_teaching_intervention_action_summary(
    db: Any,
    *,
    class_id: int,
    teacher_user_id: int,
) -> dict[str, Any]:
    overview_row = await db.execute_fetchone(
        """
        SELECT
            COUNT(1) AS total_action_count,
            SUM(CASE WHEN created_at >= datetime('now', '-7 day') THEN 1 ELSE 0 END) AS recent_7d_action_count,
            MAX(created_at) AS last_action_at
        FROM teaching_intervention_actions
        WHERE class_id = ? AND teacher_user_id = ?
        """,
        (class_id, teacher_user_id),
    )
    code_rows = await db.execute_fetchall(
        """
        SELECT
            intervention_code,
            COUNT(1) AS action_count,
            SUM(CASE WHEN created_at >= datetime('now', '-7 day') THEN 1 ELSE 0 END) AS recent_7d_count,
            MAX(created_at) AS last_action_at
        FROM teaching_intervention_actions
        WHERE class_id = ? AND teacher_user_id = ?
        GROUP BY intervention_code
        ORDER BY action_count DESC, last_action_at DESC
        """,
        (class_id, teacher_user_id),
    )

    code_stats: list[dict[str, Any]] = []
    for row in code_rows:
        code = str(row["intervention_code"] or "").strip().lower()
        if not code:
            continue
        code_stats.append(
            {
                "intervention_code": code,
                "action_count": int(row["action_count"] or 0),
                "recent_7d_count": int(row["recent_7d_count"] or 0),
                "last_action_at": row["last_action_at"] or "",
            }
        )

    return {
        "total_action_count": int(overview_row["total_action_count"] or 0) if overview_row else 0,
        "recent_7d_action_count": int(overview_row["recent_7d_action_count"] or 0) if overview_row else 0,
        "last_action_at": (overview_row["last_action_at"] if overview_row else "") or "",
        "code_stats": code_stats,
    }


async def _build_assignment_template_effect_snapshots(
    db: Any,
    *,
    teacher_user_id: int,
    template_ids: list[int] | None = None,
) -> dict[int, dict[str, Any]]:
    normalized_ids = sorted({int(x) for x in (template_ids or []) if int(x) > 0})
    if not normalized_ids:
        return {}

    placeholders = ",".join(["?"] * len(normalized_ids))
    params: list[Any] = [teacher_user_id]
    params.extend(normalized_ids)

    try:
        usage_rows = await db.execute_fetchall(
            f"""
            SELECT template_id, target_id AS assignment_id
            FROM teaching_template_usage_logs
            WHERE teacher_user_id = ?
              AND usage_scene = 'assignment_apply'
              AND target_id > 0
              AND template_id IN ({placeholders})
            GROUP BY template_id, target_id
            """,
            tuple(params),
        )
    except Exception:
        return {}

    if not usage_rows:
        return {}

    template_assignment_ids: dict[int, set[int]] = {}
    assignment_ids: set[int] = set()
    for row in usage_rows:
        tid = int(row["template_id"] or 0)
        aid = int(row["assignment_id"] or 0)
        if tid <= 0 or aid <= 0:
            continue
        assignment_ids.add(aid)
        template_assignment_ids.setdefault(tid, set()).add(aid)

    if not assignment_ids:
        return {}

    assignment_id_list = sorted(assignment_ids)
    assignment_placeholders = ",".join(["?"] * len(assignment_id_list))
    assignment_params: list[Any] = [teacher_user_id]
    assignment_params.extend(assignment_id_list)

    assignment_rows = await db.execute_fetchall(
        f"""
        SELECT
            ta.id,
            ta.class_id,
            (SELECT COUNT(*) FROM teaching_class_members tcm WHERE tcm.class_id = ta.class_id) AS member_count,
            (SELECT COUNT(*) FROM teaching_assignment_submissions tas WHERE tas.assignment_id = ta.id) AS submitted_count,
            (SELECT COUNT(*) FROM teaching_assignment_submissions tas
               JOIN teaching_submissions ts ON ts.id = tas.teaching_submission_id
             WHERE tas.assignment_id = ta.id AND ts.status = 'reviewed') AS reviewed_count,
            (SELECT AVG(te.score) FROM teaching_assignment_submissions tas
               JOIN teaching_evaluations te ON te.submission_id = tas.teaching_submission_id
             WHERE tas.assignment_id = ta.id) AS avg_score
        FROM teaching_assignments ta
        WHERE ta.teacher_user_id = ? AND ta.id IN ({assignment_placeholders})
        """,
        tuple(assignment_params),
    )

    assignment_stat_map: dict[int, dict[str, Any]] = {}
    for row in assignment_rows:
        aid = int(row["id"] or 0)
        if aid <= 0:
            continue
        assignment_stat_map[aid] = {
            "member_count": int(row["member_count"] or 0),
            "submitted_count": int(row["submitted_count"] or 0),
            "reviewed_count": int(row["reviewed_count"] or 0),
            "avg_score": float(row["avg_score"] or 0),
        }

    snapshot_map: dict[int, dict[str, Any]] = {}
    for tid, aid_set in template_assignment_ids.items():
        expected = 0
        submitted = 0
        reviewed = 0
        score_weight_sum = 0.0
        score_weight_count = 0
        valid_assignment_count = 0

        for aid in sorted(aid_set):
            stat = assignment_stat_map.get(aid)
            if not stat:
                continue
            valid_assignment_count += 1
            member_count = int(stat.get("member_count") or 0)
            submitted_count = int(stat.get("submitted_count") or 0)
            reviewed_count = int(stat.get("reviewed_count") or 0)
            expected += member_count
            submitted += submitted_count
            reviewed += reviewed_count
            avg_score = float(stat.get("avg_score") or 0)
            if reviewed_count > 0:
                score_weight_sum += avg_score * reviewed_count
                score_weight_count += reviewed_count

        snapshot_map[tid] = {
            "assignment_count": valid_assignment_count,
            "expected_submission_count": expected,
            "submitted_count": submitted,
            "reviewed_count": reviewed,
            "submission_rate": round((submitted / expected), 4) if expected > 0 else 0.0,
            "review_completion_rate": round((reviewed / submitted), 4) if submitted > 0 else 0.0,
            "avg_score": round((score_weight_sum / score_weight_count), 2) if score_weight_count > 0 else 0.0,
        }

    return snapshot_map


_VALID_TEACHING_CLASS_GOAL_METRICS = {
    "submission_rate",
    "review_completion_rate",
    "avg_score",
}
_VALID_TEACHING_CLASS_GOAL_STATUSES = {"active", "completed", "archived"}


def _normalize_teaching_class_goal_code(raw_code: str | None) -> str:
    token = re.sub(r"[^a-z0-9_-]", "_", str(raw_code or "").strip().lower())
    token = re.sub(r"_+", "_", token).strip("_")
    if not token:
        raise HTTPException(status_code=400, detail="goal_code 不能为空")
    return token[:80]


def _normalize_teaching_class_goal_metric_type(raw_metric_type: str | None) -> str:
    token = str(raw_metric_type or "").strip().lower()
    if token not in _VALID_TEACHING_CLASS_GOAL_METRICS:
        raise HTTPException(
            status_code=400,
            detail="metric_type 仅支持 submission_rate/review_completion_rate/avg_score",
        )
    return token


def _normalize_teaching_class_goal_status(raw_status: str | None) -> str:
    token = str(raw_status or "").strip().lower()
    if token not in _VALID_TEACHING_CLASS_GOAL_STATUSES:
        raise HTTPException(status_code=400, detail="status 仅支持 active/completed/archived")
    return token


def _build_teaching_class_goal_metric_values(
    *,
    expected_submission_count: int,
    submitted_count: int,
    reviewed_count: int,
    score_distribution: dict[str, Any],
) -> dict[str, float]:
    expected = int(expected_submission_count or 0)
    submitted = int(submitted_count or 0)
    reviewed = int(reviewed_count or 0)
    avg_score = float((score_distribution or {}).get("avg_score") or 0)
    return {
        "submission_rate": round((submitted / expected), 4) if expected > 0 else 0.0,
        "review_completion_rate": round((reviewed / submitted), 4) if submitted > 0 else 0.0,
        "avg_score": round(avg_score, 2),
    }


def _serialize_teaching_class_goal_row(
    row: Any,
    *,
    metric_values: dict[str, float],
) -> dict[str, Any]:
    metric_type = str(row["metric_type"] or "").strip().lower()
    if metric_type not in _VALID_TEACHING_CLASS_GOAL_METRICS:
        metric_type = "submission_rate"

    status = str(row["status"] or "").strip().lower()
    if status not in _VALID_TEACHING_CLASS_GOAL_STATUSES:
        status = "active"

    target_value = float(row["target_value"] or 0)
    current_value = float(metric_values.get(metric_type) or 0)
    progress_ratio = 0.0
    if target_value > 0:
        progress_ratio = current_value / target_value
    is_achieved = bool(target_value > 0 and current_value >= target_value)

    return {
        "id": int(row["id"] or 0),
        "class_id": int(row["class_id"] or 0),
        "teacher_user_id": int(row["teacher_user_id"] or 0),
        "goal_code": str(row["goal_code"] or ""),
        "goal_name": str(row["goal_name"] or ""),
        "metric_type": metric_type,
        "target_value": target_value,
        "current_value": round(current_value, 4) if metric_type != "avg_score" else round(current_value, 2),
        "progress_ratio": round(progress_ratio, 4),
        "is_achieved": is_achieved,
        "status": status,
        "note": str(row["note"] or ""),
        "due_at": str(row["due_at"] or ""),
        "created_at": row["created_at"] or "",
        "updated_at": row["updated_at"] or "",
    }


def _build_teaching_class_goal_summary(goals: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(goals)
    active_count = len([g for g in goals if str(g.get("status") or "") == "active"])
    completed_count = len([g for g in goals if str(g.get("status") or "") == "completed"])
    achieved_count = len([g for g in goals if bool(g.get("is_achieved"))])
    overdue_count = len(
        [
            g
            for g in goals
            if str(g.get("status") or "") == "active"
            and str(g.get("due_at") or "").strip()
            and str(g.get("due_at") or "") < datetime.now(timezone.utc).date().isoformat()
            and not bool(g.get("is_achieved"))
        ]
    )
    return {
        "total_goal_count": total,
        "active_goal_count": active_count,
        "completed_goal_count": completed_count,
        "achieved_goal_count": achieved_count,
        "overdue_goal_count": overdue_count,
    }


async def _list_teaching_class_goal_rows(
    db: Any,
    *,
    class_id: int,
    teacher_user_id: int,
    include_archived: bool = False,
) -> list[Any]:
    sql = """
        SELECT id, class_id, teacher_user_id, goal_code, goal_name, metric_type, target_value,
               status, note, due_at, created_at, updated_at
        FROM teaching_class_goals
        WHERE class_id = ? AND teacher_user_id = ?
    """
    params: list[Any] = [class_id, teacher_user_id]
    if not include_archived:
        sql += " AND status != 'archived'"
    sql += " ORDER BY created_at DESC, id DESC"
    return await db.execute_fetchall(sql, tuple(params))


async def _compute_teaching_class_goal_metric_values(
    db: Any,
    *,
    class_id: int,
) -> tuple[dict[str, float], dict[str, int]]:
    member_row = await db.execute_fetchone(
        "SELECT COUNT(*) AS cnt FROM teaching_class_members WHERE class_id = ?",
        (class_id,),
    )
    assignment_row = await db.execute_fetchone(
        "SELECT COUNT(*) AS cnt FROM teaching_assignments WHERE class_id = ?",
        (class_id,),
    )
    submitted_row = await db.execute_fetchone(
        "SELECT COUNT(*) AS cnt FROM teaching_assignment_submissions WHERE class_id = ?",
        (class_id,),
    )
    reviewed_row = await db.execute_fetchone(
        """
        SELECT COUNT(*) AS cnt
        FROM teaching_assignment_submissions tas
        JOIN teaching_submissions ts ON ts.id = tas.teaching_submission_id
        WHERE tas.class_id = ? AND ts.status = 'reviewed'
        """,
        (class_id,),
    )
    avg_score_row = await db.execute_fetchone(
        """
        SELECT AVG(te.score) AS avg_score
        FROM teaching_assignment_submissions tas
        JOIN teaching_evaluations te ON te.submission_id = tas.teaching_submission_id
        WHERE tas.class_id = ?
        """,
        (class_id,),
    )

    member_count = int(member_row["cnt"] or 0) if member_row else 0
    assignment_count = int(assignment_row["cnt"] or 0) if assignment_row else 0
    submitted_count = int(submitted_row["cnt"] or 0) if submitted_row else 0
    reviewed_count = int(reviewed_row["cnt"] or 0) if reviewed_row else 0
    expected_submission_count = member_count * assignment_count

    metric_values = {
        "submission_rate": round((submitted_count / expected_submission_count), 4)
        if expected_submission_count > 0
        else 0.0,
        "review_completion_rate": round((reviewed_count / submitted_count), 4)
        if submitted_count > 0
        else 0.0,
        "avg_score": round(float(avg_score_row["avg_score"] or 0), 2) if avg_score_row else 0.0,
    }
    context = {
        "member_count": member_count,
        "assignment_count": assignment_count,
        "expected_submission_count": expected_submission_count,
        "submitted_count": submitted_count,
        "reviewed_count": reviewed_count,
    }
    return metric_values, context


_DEFAULT_TEACHING_GOAL_TREND_WINDOW_DAYS = 30
_MIN_TEACHING_GOAL_TREND_WINDOW_DAYS = 7
_MAX_TEACHING_GOAL_TREND_WINDOW_DAYS = 180


def _normalize_teaching_goal_trend_window_days(raw_days: int | None) -> int:
    value = int(raw_days or _DEFAULT_TEACHING_GOAL_TREND_WINDOW_DAYS)
    if value < _MIN_TEACHING_GOAL_TREND_WINDOW_DAYS:
        return _MIN_TEACHING_GOAL_TREND_WINDOW_DAYS
    if value > _MAX_TEACHING_GOAL_TREND_WINDOW_DAYS:
        return _MAX_TEACHING_GOAL_TREND_WINDOW_DAYS
    return value


def _parse_iso_date_token(raw_value: str | None) -> Any | None:
    token = str(raw_value or "").strip()
    if not token:
        return None
    token = token[:10]
    try:
        return datetime.strptime(token, "%Y-%m-%d").date()
    except Exception:
        return None


def _normalize_teaching_goal_trend_direction(metric_type: str, delta_value: float) -> str:
    threshold = 0.005 if metric_type in {"submission_rate", "review_completion_rate"} else 0.2
    if delta_value > threshold:
        return "up"
    if delta_value < -threshold:
        return "down"
    return "flat"


def _normalize_teaching_goal_trend_risk_level(
    *,
    status: str,
    is_achieved: bool,
    target_value: float,
    current_value: float,
    velocity_per_day: float,
    due_at: str,
    estimated_days_to_target: float | None,
) -> tuple[str, str]:
    if status == "archived":
        return "low", "目标已归档"
    if status == "completed" or is_achieved:
        return "low", "目标已达成"

    target_gap = max(target_value - current_value, 0.0)
    if target_gap <= 0:
        return "low", "当前指标已达到目标值"

    due_date = _parse_iso_date_token(due_at)
    today = datetime.now(timezone.utc).date()

    if due_date is None:
        if velocity_per_day <= 0:
            return "medium", "近期进展停滞，建议补充干预动作"
        return "low", "未设置截止日期，当前趋势可控"

    days_left = (due_date - today).days
    if days_left < 0:
        return "high", "目标已逾期且尚未达成"

    if velocity_per_day <= 0:
        return "high", "截止日期临近但当前进展停滞"

    if estimated_days_to_target is None:
        return "high", "缺少可预测速度，建议尽快补充执行动作"

    if estimated_days_to_target > days_left * 1.2:
        return "high", "按当前速度预计无法按期达成"
    if estimated_days_to_target > days_left * 0.85:
        return "medium", "按当前速度存在延期风险"
    return "low", "按当前速度可在截止前达成"


def _build_teaching_goal_trend_item(
    goal: dict[str, Any],
    points: list[dict[str, Any]],
    *,
    window_days: int,
) -> dict[str, Any]:
    goal_id = int(goal.get("id") or 0)
    metric_type = str(goal.get("metric_type") or "").strip().lower()
    target_value = float(goal.get("target_value") or 0)
    current_value = float(goal.get("current_value") or 0)
    status = str(goal.get("status") or "active").strip().lower()
    is_achieved = bool(goal.get("is_achieved"))

    normalized_points: list[dict[str, Any]] = []
    for point in points:
        snapshot_date = str(point.get("snapshot_date") or "")
        parsed_date = _parse_iso_date_token(snapshot_date)
        if not parsed_date:
            continue
        normalized_points.append(
            {
                "snapshot_date": snapshot_date,
                "parsed_date": parsed_date,
                "current_value": float(point.get("current_value") or 0),
                "progress_ratio": float(point.get("progress_ratio") or 0),
                "is_achieved": bool(int(point.get("is_achieved") or 0)),
            }
        )

    if not normalized_points:
        today = datetime.now(timezone.utc).date().isoformat()
        normalized_points = [
            {
                "snapshot_date": today,
                "parsed_date": _parse_iso_date_token(today),
                "current_value": current_value,
                "progress_ratio": float(goal.get("progress_ratio") or 0),
                "is_achieved": is_achieved,
            }
        ]

    normalized_points.sort(key=lambda item: item["parsed_date"])
    first_point = normalized_points[0]
    last_point = normalized_points[-1]

    span_days = 1
    if first_point["parsed_date"] and last_point["parsed_date"]:
        span_days = max((last_point["parsed_date"] - first_point["parsed_date"]).days, 1)

    delta_window = float(last_point["current_value"] or 0) - float(first_point["current_value"] or 0)
    velocity_per_day = delta_window / span_days if span_days > 0 else 0.0

    delta_recent_7d = 0.0
    if last_point["parsed_date"]:
        recent_anchor_date = last_point["parsed_date"] - timedelta(days=6)
        recent_points = [item for item in normalized_points if item["parsed_date"] and item["parsed_date"] >= recent_anchor_date]
        recent_base = recent_points[0] if recent_points else first_point
        delta_recent_7d = float(last_point["current_value"] or 0) - float(recent_base["current_value"] or 0)

    estimated_days_to_target: float | None = None
    if target_value > current_value and velocity_per_day > 0:
        estimated_days_to_target = (target_value - current_value) / velocity_per_day

    risk_level, risk_reason = _normalize_teaching_goal_trend_risk_level(
        status=status,
        is_achieved=is_achieved,
        target_value=target_value,
        current_value=current_value,
        velocity_per_day=velocity_per_day,
        due_at=str(goal.get("due_at") or ""),
        estimated_days_to_target=estimated_days_to_target,
    )

    trend_direction = _normalize_teaching_goal_trend_direction(metric_type, delta_window)
    compact_points = [
        {
            "snapshot_date": item["snapshot_date"],
            "current_value": round(float(item["current_value"] or 0), 4),
            "progress_ratio": round(float(item["progress_ratio"] or 0), 4),
            "is_achieved": bool(item["is_achieved"]),
        }
        for item in normalized_points[-window_days:]
    ]

    return {
        "goal_id": goal_id,
        "goal_code": str(goal.get("goal_code") or ""),
        "goal_name": str(goal.get("goal_name") or ""),
        "metric_type": metric_type,
        "status": status,
        "due_at": str(goal.get("due_at") or ""),
        "target_value": round(target_value, 4),
        "current_value": round(current_value, 4),
        "delta_window": round(delta_window, 4),
        "delta_recent_7d": round(delta_recent_7d, 4),
        "velocity_per_day": round(velocity_per_day, 6),
        "estimated_days_to_target": round(estimated_days_to_target, 2) if estimated_days_to_target is not None else None,
        "trend_direction": trend_direction,
        "risk_level": risk_level,
        "risk_reason": risk_reason,
        "point_count": len(compact_points),
        "points": compact_points,
    }


async def _upsert_teaching_class_goal_snapshots(
    db: Any,
    *,
    goals: list[dict[str, Any]],
    source: str,
) -> None:
    if not goals:
        return

    snapshot_date = datetime.now(timezone.utc).date().isoformat()
    touched = False
    for goal in goals:
        goal_id = int(goal.get("id") or 0)
        if goal_id <= 0:
            continue
        await db.execute(
            """
            INSERT INTO teaching_class_goal_snapshots
                (class_id, goal_id, teacher_user_id, goal_code, metric_type,
                 target_value, current_value, progress_ratio, is_achieved, goal_status,
                 due_at, snapshot_date, source, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(goal_id, snapshot_date)
            DO UPDATE SET
                class_id = excluded.class_id,
                teacher_user_id = excluded.teacher_user_id,
                goal_code = excluded.goal_code,
                metric_type = excluded.metric_type,
                target_value = excluded.target_value,
                current_value = excluded.current_value,
                progress_ratio = excluded.progress_ratio,
                is_achieved = excluded.is_achieved,
                goal_status = excluded.goal_status,
                due_at = excluded.due_at,
                source = excluded.source,
                created_at = CURRENT_TIMESTAMP
            """,
            (
                int(goal.get("class_id") or 0),
                goal_id,
                int(goal.get("teacher_user_id") or 0),
                str(goal.get("goal_code") or ""),
                str(goal.get("metric_type") or ""),
                float(goal.get("target_value") or 0),
                float(goal.get("current_value") or 0),
                float(goal.get("progress_ratio") or 0),
                1 if bool(goal.get("is_achieved")) else 0,
                str(goal.get("status") or "active"),
                str(goal.get("due_at") or ""),
                snapshot_date,
                str(source or "system")[:64],
            ),
        )
        touched = True

    if touched:
        await db.commit()


async def _list_teaching_class_goal_snapshot_rows(
    db: Any,
    *,
    class_id: int,
    teacher_user_id: int,
    goal_ids: list[int],
    window_days: int,
) -> list[Any]:
    normalized_goal_ids = [int(item) for item in goal_ids if int(item) > 0]
    if not normalized_goal_ids:
        return []

    placeholders = ",".join(["?"] * len(normalized_goal_ids))
    params: list[Any] = [class_id, teacher_user_id]
    params.extend(normalized_goal_ids)
    params.append(f"-{window_days - 1} day")

    return await db.execute_fetchall(
        f"""
        SELECT goal_id, snapshot_date, current_value, progress_ratio, is_achieved, goal_status
        FROM teaching_class_goal_snapshots
        WHERE class_id = ?
          AND teacher_user_id = ?
          AND goal_id IN ({placeholders})
          AND snapshot_date >= date('now', ?)
        ORDER BY snapshot_date ASC, id ASC
        """,
        tuple(params),
    )


def _build_teaching_class_goal_trend_summary(
    goals: list[dict[str, Any]],
    snapshot_rows: list[Any],
    *,
    window_days: int,
) -> dict[str, Any]:
    goal_point_map: dict[int, list[dict[str, Any]]] = {}
    for row in snapshot_rows:
        gid = int(row["goal_id"] or 0)
        if gid <= 0:
            continue
        goal_point_map.setdefault(gid, []).append(
            {
                "snapshot_date": str(row["snapshot_date"] or ""),
                "current_value": float(row["current_value"] or 0),
                "progress_ratio": float(row["progress_ratio"] or 0),
                "is_achieved": int(row["is_achieved"] or 0),
            }
        )

    goal_trends: list[dict[str, Any]] = []
    risk_counts = {"high": 0, "medium": 0, "low": 0}
    direction_counts = {"up": 0, "flat": 0, "down": 0}

    for goal in goals:
        gid = int(goal.get("id") or 0)
        trend_item = _build_teaching_goal_trend_item(
            goal,
            goal_point_map.get(gid, []),
            window_days=window_days,
        )
        goal_trends.append(trend_item)

        risk_level = str(trend_item.get("risk_level") or "low")
        if risk_level not in risk_counts:
            risk_level = "low"
        risk_counts[risk_level] += 1

        direction = str(trend_item.get("trend_direction") or "flat")
        if direction not in direction_counts:
            direction = "flat"
        direction_counts[direction] += 1

    goal_trends.sort(
        key=lambda item: (
            0 if item.get("risk_level") == "high" else 1 if item.get("risk_level") == "medium" else 2,
            0 if item.get("trend_direction") == "down" else 1 if item.get("trend_direction") == "flat" else 2,
            str(item.get("goal_code") or ""),
        )
    )

    return {
        "window_days": window_days,
        "goal_trends": goal_trends,
        "risk_summary": {
            "high_risk_count": int(risk_counts["high"]),
            "medium_risk_count": int(risk_counts["medium"]),
            "low_risk_count": int(risk_counts["low"]),
        },
        "overview": {
            "goal_count": len(goal_trends),
            "up_count": int(direction_counts["up"]),
            "flat_count": int(direction_counts["flat"]),
            "down_count": int(direction_counts["down"]),
        },
    }


def _latest_daily_value_at_or_before(
    daily_series: list[tuple[Any, float]],
    *,
    target_date: Any,
) -> float | None:
    last_value: float | None = None
    for day, value in daily_series:
        if day > target_date:
            break
        last_value = value
    return last_value


def _normalize_teaching_strategy_variant_token(raw_variant: str | None) -> str:
    token = str(raw_variant or "").strip().upper()
    if not token:
        return ""
    token = re.sub(r"[^A-Z0-9_-]", "", token)
    return token[:32]


def _extract_teaching_strategy_variant(metadata: dict[str, Any] | None) -> str:
    if not isinstance(metadata, dict):
        return ""
    for key in ("strategy_variant", "variant", "ab_variant", "experiment_variant"):
        variant = _normalize_teaching_strategy_variant_token(str(metadata.get(key) or ""))
        if variant:
            return variant
    return ""


async def _build_teaching_intervention_outcome_correlation(
    db: Any,
    *,
    class_id: int,
    teacher_user_id: int,
    window_days: int,
) -> dict[str, Any]:
    action_rows = await db.execute_fetchall(
        """
        SELECT intervention_code, metadata_json, created_at
        FROM teaching_intervention_actions
        WHERE class_id = ?
          AND teacher_user_id = ?
          AND created_at >= datetime('now', ?)
        ORDER BY created_at ASC, id ASC
        """,
        (class_id, teacher_user_id, f"-{window_days} day"),
    )

    if not action_rows:
        return {
            "window_days": window_days,
            "action_count": 0,
            "observed_action_count": 0,
            "code_effects": [],
            "variant_effects": [],
            "note": "窗口期内暂无干预执行记录",
        }

    snapshot_rows = await db.execute_fetchall(
        """
        SELECT snapshot_date, progress_ratio
        FROM teaching_class_goal_snapshots
        WHERE class_id = ?
          AND teacher_user_id = ?
          AND snapshot_date >= date('now', ?)
        ORDER BY snapshot_date ASC, id ASC
        """,
        (class_id, teacher_user_id, f"-{window_days + 7} day"),
    )

    progress_daily_map: dict[Any, list[float]] = {}
    for row in snapshot_rows:
        day = _parse_iso_date_token(row["snapshot_date"])
        if not day:
            continue
        progress_daily_map.setdefault(day, []).append(float(row["progress_ratio"] or 0))

    daily_series: list[tuple[Any, float]] = []
    for day in sorted(progress_daily_map.keys()):
        values = progress_daily_map.get(day) or []
        if not values:
            continue
        daily_series.append((day, sum(values) / len(values)))

    code_effect_map: dict[str, dict[str, Any]] = {}
    variant_effect_map: dict[str, dict[str, Any]] = {}
    observed_action_count = 0

    for row in action_rows:
        code = str(row["intervention_code"] or "").strip().lower() or "unknown"
        metadata = _parse_teaching_template_payload(row["metadata_json"])
        variant = _extract_teaching_strategy_variant(metadata)
        created_at = str(row["created_at"] or "")
        action_day = _parse_iso_date_token(created_at)
        if action_day is None:
            continue

        code_stat = code_effect_map.setdefault(
            code,
            {
                "intervention_code": code,
                "action_count": 0,
                "observed_action_count": 0,
                "delta_sum": 0.0,
            },
        )
        code_stat["action_count"] += 1

        if variant:
            variant_key = f"{code}::{variant}"
            variant_stat = variant_effect_map.setdefault(
                variant_key,
                {
                    "intervention_code": code,
                    "strategy_variant": variant,
                    "action_count": 0,
                    "observed_action_count": 0,
                    "delta_sum": 0.0,
                },
            )
            variant_stat["action_count"] += 1
        else:
            variant_key = ""

        before_value = _latest_daily_value_at_or_before(daily_series, target_date=action_day)
        after_value = _latest_daily_value_at_or_before(daily_series, target_date=action_day + timedelta(days=7))
        if before_value is None or after_value is None:
            continue

        delta = float(after_value - before_value)
        observed_action_count += 1
        code_stat["observed_action_count"] += 1
        code_stat["delta_sum"] += delta

        if variant_key:
            variant_effect_map[variant_key]["observed_action_count"] += 1
            variant_effect_map[variant_key]["delta_sum"] += delta

    def _signal_from_delta(avg_delta: float) -> str:
        if avg_delta >= 0.03:
            return "positive"
        if avg_delta <= -0.01:
            return "negative"
        return "neutral"

    code_effects: list[dict[str, Any]] = []
    for code, stat in code_effect_map.items():
        observed = int(stat["observed_action_count"] or 0)
        avg_delta = (float(stat["delta_sum"] or 0) / observed) if observed > 0 else 0.0
        code_effects.append(
            {
                "intervention_code": code,
                "action_count": int(stat["action_count"] or 0),
                "observed_action_count": observed,
                "avg_progress_delta_7d": round(avg_delta, 4),
                "effect_signal": _signal_from_delta(avg_delta),
            }
        )

    variant_effects: list[dict[str, Any]] = []
    for _, stat in variant_effect_map.items():
        observed = int(stat["observed_action_count"] or 0)
        avg_delta = (float(stat["delta_sum"] or 0) / observed) if observed > 0 else 0.0
        variant_effects.append(
            {
                "intervention_code": str(stat["intervention_code"] or ""),
                "strategy_variant": str(stat["strategy_variant"] or ""),
                "action_count": int(stat["action_count"] or 0),
                "observed_action_count": observed,
                "avg_progress_delta_7d": round(avg_delta, 4),
                "effect_signal": _signal_from_delta(avg_delta),
            }
        )

    code_effects.sort(key=lambda item: (-int(item["action_count"]), str(item["intervention_code"])))
    variant_effects.sort(
        key=lambda item: (
            str(item.get("intervention_code") or ""),
            str(item.get("strategy_variant") or ""),
        )
    )
    return {
        "window_days": window_days,
        "action_count": len(action_rows),
        "observed_action_count": observed_action_count,
        "code_effects": code_effects,
        "variant_effects": variant_effects,
        "daily_progress_point_count": len(daily_series),
        "note": "基于目标进度快照的相关性分析，非严格因果结论",
    }


async def _build_teaching_goal_trend_bundle(
    db: Any,
    *,
    class_id: int,
    teacher_user_id: int,
    goals: list[dict[str, Any]],
    window_days: int,
    snapshot_source: str,
) -> dict[str, Any]:
    normalized_window_days = _normalize_teaching_goal_trend_window_days(window_days)
    await _upsert_teaching_class_goal_snapshots(
        db,
        goals=goals,
        source=snapshot_source,
    )

    goal_ids = [int(item.get("id") or 0) for item in goals if int(item.get("id") or 0) > 0]
    snapshot_rows = await _list_teaching_class_goal_snapshot_rows(
        db,
        class_id=class_id,
        teacher_user_id=teacher_user_id,
        goal_ids=goal_ids,
        window_days=normalized_window_days,
    )
    trend_summary = _build_teaching_class_goal_trend_summary(
        goals,
        snapshot_rows,
        window_days=normalized_window_days,
    )
    intervention_outcome_correlation = await _build_teaching_intervention_outcome_correlation(
        db,
        class_id=class_id,
        teacher_user_id=teacher_user_id,
        window_days=normalized_window_days,
    )
    trend_summary["intervention_outcome_correlation"] = intervention_outcome_correlation
    return trend_summary


def _build_teaching_intervention_strategy_recommendations(
    *,
    interventions: list[dict[str, Any]],
    intervention_action_summary: dict[str, Any],
    intervention_outcome_correlation: dict[str, Any],
    goal_risk_summary: dict[str, Any],
) -> dict[str, Any]:
    code_stat_map = {
        str(item.get("intervention_code") or "").strip().lower(): item
        for item in (intervention_action_summary.get("code_stats") or [])
        if str(item.get("intervention_code") or "").strip()
    }
    effect_map = {
        str(item.get("intervention_code") or "").strip().lower(): item
        for item in (intervention_outcome_correlation.get("code_effects") or [])
        if str(item.get("intervention_code") or "").strip()
    }

    high_risk_count = int(goal_risk_summary.get("high_risk_count") or 0)
    medium_risk_count = int(goal_risk_summary.get("medium_risk_count") or 0)

    recommendations: list[dict[str, Any]] = []
    for item in interventions:
        code = str(item.get("code") or "").strip().lower()
        if not code:
            continue

        severity = str(item.get("severity") or "medium").strip().lower()
        executed_recent_7d = int(item.get("executed_recent_7d") or 0)
        executed_count = int(item.get("executed_count") or 0)

        effect = effect_map.get(code, {})
        effect_signal = str(effect.get("effect_signal") or "neutral")
        observed_count = int(effect.get("observed_action_count") or 0)
        avg_delta_7d = float(effect.get("avg_progress_delta_7d") or 0)

        urgency_score = 40
        if severity == "high":
            urgency_score = 90
        elif severity == "medium":
            urgency_score = 65

        if high_risk_count > 0:
            urgency_score += 8
        elif medium_risk_count > 0:
            urgency_score += 3

        if executed_recent_7d > 0:
            urgency_score -= 14
        elif executed_count > 0:
            urgency_score += 4

        strategy_mode = "explore"
        strategy_title = "探索试验"
        strategy_tip = "缺少历史效果样本，建议先小范围验证。"
        recommended_variant = "A"

        if effect_signal == "positive":
            strategy_mode = "scale"
            strategy_title = "扩量执行"
            strategy_tip = "历史相关性偏正向，建议提高执行频次。"
            recommended_variant = "A"
            urgency_score += 12
        elif effect_signal == "negative":
            strategy_mode = "switch"
            strategy_title = "替代策略"
            strategy_tip = "历史相关性偏负向，建议切换策略并保留对照。"
            recommended_variant = "B"
            urgency_score += 6
        elif observed_count > 0:
            strategy_mode = "experiment"
            strategy_title = "双轨实验"
            strategy_tip = "效果不稳定，建议并行 A/B 验证。"
            recommended_variant = "A"

        priority_bucket = "P2"
        if urgency_score >= 90:
            priority_bucket = "P0"
        elif urgency_score >= 70:
            priority_bucket = "P1"

        code_stat = code_stat_map.get(code, {})
        experiment_id = f"{code}-{datetime.now(timezone.utc).date().strftime('%Y%m%d')}"

        recommendations.append(
            {
                "intervention_code": code,
                "intervention_title": str(item.get("title") or ""),
                "severity": severity,
                "priority": priority_bucket,
                "urgency_score": int(urgency_score),
                "strategy_mode": strategy_mode,
                "strategy_title": strategy_title,
                "strategy_tip": strategy_tip,
                "recommended_variant": recommended_variant,
                "suggested_variants": ["A", "B"],
                "experiment_id": experiment_id,
                "effect_signal": effect_signal,
                "effect_observed_count": observed_count,
                "effect_avg_progress_delta_7d": round(avg_delta_7d, 4),
                "executed_recent_7d": executed_recent_7d,
                "executed_count": int(code_stat.get("action_count") or executed_count),
                "last_executed_at": str(code_stat.get("last_action_at") or item.get("last_executed_at") or ""),
                "can_mark_executed": bool(item.get("can_mark_executed")),
            }
        )

    recommendations.sort(
        key=lambda rec: (
            0 if rec.get("priority") == "P0" else 1 if rec.get("priority") == "P1" else 2,
            -int(rec.get("urgency_score") or 0),
            str(rec.get("intervention_code") or ""),
        )
    )

    top = recommendations[:6]
    summary = {
        "recommendation_count": len(top),
        "p0_count": len([x for x in top if x.get("priority") == "P0"]),
        "p1_count": len([x for x in top if x.get("priority") == "P1"]),
        "p2_count": len([x for x in top if x.get("priority") == "P2"]),
        "has_high_risk_goal": high_risk_count > 0,
        "window_days": int(intervention_outcome_correlation.get("window_days") or 30),
    }
    return {
        "recommendations": top,
        "summary": summary,
    }


def _normalize_teaching_term_code(raw_code: str | None) -> str:
    token = re.sub(r"[^a-z0-9_-]", "_", str(raw_code or "").strip().lower())
    token = re.sub(r"_+", "_", token).strip("_")
    if not token:
        raise HTTPException(status_code=400, detail="term_code 不能为空")
    return token[:80]


_VALID_TEACHING_EXPERIMENT_STATUSES = {"planned", "running", "paused", "completed", "cancelled"}
_VALID_TEACHING_EXPERIMENT_MODES = {"scale", "experiment", "switch", "explore"}


def _normalize_teaching_experiment_status(raw_status: str | None) -> str:
    token = str(raw_status or "").strip().lower()
    if token not in _VALID_TEACHING_EXPERIMENT_STATUSES:
        raise HTTPException(status_code=400, detail="status 仅支持 planned/running/paused/completed/cancelled")
    return token


def _normalize_teaching_experiment_mode(raw_mode: str | None) -> str:
    token = str(raw_mode or "").strip().lower()
    return token if token in _VALID_TEACHING_EXPERIMENT_MODES else "experiment"


def _normalize_teaching_experiment_target_metric(raw_metric: str | None) -> str:
    token = str(raw_metric or "").strip().lower()
    if token in {"submission_rate", "review_completion_rate", "avg_score"}:
        return token
    return "submission_rate"


def _normalize_teaching_experiment_code(
    raw_code: str | None,
    *,
    fallback_intervention_code: str = "experiment",
) -> str:
    token = re.sub(r"[^a-z0-9_-]", "_", str(raw_code or "").strip().lower())
    token = re.sub(r"_+", "_", token).strip("_")
    if token:
        return token[:120]

    fallback = re.sub(r"[^a-z0-9_-]", "_", str(fallback_intervention_code or "experiment").strip().lower())
    fallback = re.sub(r"_+", "_", fallback).strip("_") or "experiment"
    auto_code = f"{fallback}-{datetime.now(timezone.utc).strftime('%Y%m%d')}"
    return auto_code[:120]


def _normalize_teaching_experiment_variants(
    raw_variants: list[Any] | None,
    *,
    recommended_variant: str,
) -> list[str]:
    normalized: list[str] = []
    for raw in raw_variants or []:
        token = _normalize_teaching_strategy_variant_token(str(raw or ""))
        if token and token not in normalized:
            normalized.append(token)

    rec = _normalize_teaching_strategy_variant_token(recommended_variant) or "A"
    if rec not in normalized:
        normalized.append(rec)

    if not normalized:
        normalized = [rec]
    if len(normalized) == 1:
        normalized.append("B" if normalized[0] != "B" else "A")

    return normalized[:6]


def _parse_json_value(payload_json: str | None) -> Any:
    try:
        return json.loads(payload_json or "null")
    except Exception:
        return None


def _serialize_teaching_intervention_experiment_row(
    row: Any,
    *,
    can_manage: bool = False,
) -> dict[str, Any]:
    raw_plan = _parse_json_value(row["variant_plan_json"])
    raw_variants: list[Any] = []
    source_recommendation: dict[str, Any] = {}
    if isinstance(raw_plan, list):
        raw_variants = raw_plan
    elif isinstance(raw_plan, dict):
        maybe_variants = raw_plan.get("variants")
        raw_variants = maybe_variants if isinstance(maybe_variants, list) else []
        maybe_source = raw_plan.get("source_recommendation")
        source_recommendation = maybe_source if isinstance(maybe_source, dict) else {}

    recommended_variant = _normalize_teaching_strategy_variant_token(str(row["recommended_variant"] or "")) or "A"
    variants = _normalize_teaching_experiment_variants(raw_variants, recommended_variant=recommended_variant)

    status = str(row["status"] or "planned").strip().lower()
    if status not in _VALID_TEACHING_EXPERIMENT_STATUSES:
        status = "planned"

    return {
        "id": int(row["id"] or 0),
        "class_id": int(row["class_id"] or 0),
        "teacher_user_id": int(row["teacher_user_id"] or 0),
        "experiment_code": _normalize_teaching_experiment_code(
            str(row["experiment_code"] or ""),
            fallback_intervention_code=str(row["intervention_code"] or "experiment"),
        ),
        "intervention_code": str(row["intervention_code"] or "").strip().lower(),
        "intervention_title": str(row["intervention_title"] or ""),
        "strategy_mode": _normalize_teaching_experiment_mode(row["strategy_mode"]),
        "target_metric": _normalize_teaching_experiment_target_metric(row["target_metric"]),
        "recommended_variant": recommended_variant,
        "variants": variants,
        "window_days": max(7, min(90, int(row["window_days"] or 14))),
        "status": status,
        "note": str(row["note"] or ""),
        "source_recommendation": source_recommendation,
        "created_at": row["created_at"] or "",
        "updated_at": row["updated_at"] or "",
        "can_manage": bool(can_manage),
    }


def _build_teaching_intervention_experiment_status_summary(
    experiments: list[dict[str, Any]],
) -> dict[str, Any]:
    counts = {status: 0 for status in _VALID_TEACHING_EXPERIMENT_STATUSES}
    for item in experiments:
        status = str(item.get("status") or "planned").strip().lower()
        if status not in counts:
            status = "planned"
        counts[status] += 1

    active_count = int(counts["planned"] + counts["running"] + counts["paused"])
    return {
        "experiment_count": len(experiments),
        "active_count": active_count,
        "planned_count": int(counts["planned"]),
        "running_count": int(counts["running"]),
        "paused_count": int(counts["paused"]),
        "completed_count": int(counts["completed"]),
        "cancelled_count": int(counts["cancelled"]),
    }


async def _list_teaching_intervention_experiment_rows(
    db: Any,
    *,
    class_id: int,
    teacher_user_id: int,
    limit: int,
    offset: int,
    statuses: list[str] | None = None,
    experiment_codes: list[str] | None = None,
) -> tuple[list[Any], int]:
    normalized_statuses = [
        _normalize_teaching_experiment_status(status)
        for status in (statuses or [])
        if str(status or "").strip()
    ]
    normalized_codes = [
        _normalize_teaching_experiment_code(code)
        for code in (experiment_codes or [])
        if str(code or "").strip()
    ]

    where_parts = ["class_id = ?", "teacher_user_id = ?"]
    params: list[Any] = [class_id, teacher_user_id]

    if normalized_statuses:
        placeholders = ",".join(["?"] * len(normalized_statuses))
        where_parts.append(f"status IN ({placeholders})")
        params.extend(normalized_statuses)

    if normalized_codes:
        placeholders = ",".join(["?"] * len(normalized_codes))
        where_parts.append(f"experiment_code IN ({placeholders})")
        params.extend(normalized_codes)

    where_sql = " AND ".join(where_parts)

    rows = await db.execute_fetchall(
        f"""
        SELECT id, class_id, teacher_user_id, experiment_code, intervention_code, intervention_title,
               strategy_mode, target_metric, recommended_variant, variant_plan_json,
               window_days, status, note, created_at, updated_at
        FROM teaching_intervention_experiments
        WHERE {where_sql}
        ORDER BY updated_at DESC, id DESC
        LIMIT ? OFFSET ?
        """,
        tuple(params + [limit, offset]),
    )
    total_row = await db.execute_fetchone(
        f"""
        SELECT COUNT(1) AS cnt
        FROM teaching_intervention_experiments
        WHERE {where_sql}
        """,
        tuple(params),
    )
    total = int(total_row["cnt"] or 0) if total_row else len(rows)
    return rows, total


async def _build_teaching_intervention_experiment_result_map(
    db: Any,
    *,
    class_id: int,
    teacher_user_id: int,
    experiments: list[dict[str, Any]],
    analysis_window_days: int,
) -> dict[str, dict[str, Any]]:
    normalized_window_days = max(7, min(180, int(analysis_window_days or 60)))
    if not experiments:
        return {}

    snapshot_rows = await db.execute_fetchall(
        """
        SELECT snapshot_date, progress_ratio
        FROM teaching_class_goal_snapshots
        WHERE class_id = ?
          AND teacher_user_id = ?
          AND snapshot_date >= date('now', ?)
        ORDER BY snapshot_date ASC, id ASC
        """,
        (class_id, teacher_user_id, f"-{normalized_window_days + 7} day"),
    )

    progress_daily_map: dict[Any, list[float]] = {}
    for row in snapshot_rows:
        day = _parse_iso_date_token(row["snapshot_date"])
        if not day:
            continue
        progress_daily_map.setdefault(day, []).append(float(row["progress_ratio"] or 0))

    daily_series: list[tuple[Any, float]] = []
    for day in sorted(progress_daily_map.keys()):
        values = progress_daily_map.get(day) or []
        if not values:
            continue
        daily_series.append((day, sum(values) / len(values)))

    action_rows = await db.execute_fetchall(
        """
        SELECT intervention_code, metadata_json, created_at
        FROM teaching_intervention_actions
        WHERE class_id = ?
          AND teacher_user_id = ?
          AND created_at >= datetime('now', ?)
        ORDER BY created_at ASC, id ASC
        """,
        (class_id, teacher_user_id, f"-{normalized_window_days} day"),
    )

    parsed_actions: list[dict[str, Any]] = []
    for row in action_rows:
        intervention_code = str(row["intervention_code"] or "").strip().lower()
        if not intervention_code:
            continue

        metadata = _parse_teaching_template_payload(row["metadata_json"])
        experiment_code_raw = str(metadata.get("strategy_experiment_id") or "").strip()
        if not experiment_code_raw:
            continue

        experiment_code = _normalize_teaching_experiment_code(
            experiment_code_raw,
            fallback_intervention_code=intervention_code,
        )
        strategy_variant = _extract_teaching_strategy_variant(metadata)
        action_day = _parse_iso_date_token(row["created_at"])
        if action_day is None:
            continue

        parsed_actions.append(
            {
                "intervention_code": intervention_code,
                "experiment_code": experiment_code,
                "strategy_variant": strategy_variant,
                "action_day": action_day,
            }
        )

    def _signal_from_delta(avg_delta: float) -> str:
        if avg_delta >= 0.03:
            return "positive"
        if avg_delta <= -0.01:
            return "negative"
        return "neutral"

    result_map: dict[str, dict[str, Any]] = {}
    for exp in experiments:
        intervention_code = str(exp.get("intervention_code") or "").strip().lower()
        experiment_code = _normalize_teaching_experiment_code(
            str(exp.get("experiment_code") or ""),
            fallback_intervention_code=intervention_code or "experiment",
        )
        recommended_variant = _normalize_teaching_strategy_variant_token(str(exp.get("recommended_variant") or "")) or "A"
        variants = _normalize_teaching_experiment_variants(exp.get("variants") or [], recommended_variant=recommended_variant)

        variant_stat_map: dict[str, dict[str, Any]] = {
            variant: {
                "strategy_variant": variant,
                "action_count": 0,
                "observed_action_count": 0,
                "delta_sum": 0.0,
            }
            for variant in variants
        }

        action_count = 0
        observed_action_count = 0
        for action in parsed_actions:
            if action.get("intervention_code") != intervention_code:
                continue
            if action.get("experiment_code") != experiment_code:
                continue

            variant = _normalize_teaching_strategy_variant_token(str(action.get("strategy_variant") or ""))
            if not variant:
                continue

            stat = variant_stat_map.setdefault(
                variant,
                {
                    "strategy_variant": variant,
                    "action_count": 0,
                    "observed_action_count": 0,
                    "delta_sum": 0.0,
                },
            )
            stat["action_count"] += 1
            action_count += 1

            action_day = action.get("action_day")
            before_value = _latest_daily_value_at_or_before(daily_series, target_date=action_day)
            after_value = _latest_daily_value_at_or_before(daily_series, target_date=action_day + timedelta(days=7))
            if before_value is None or after_value is None:
                continue

            delta = float(after_value - before_value)
            observed_action_count += 1
            stat["observed_action_count"] += 1
            stat["delta_sum"] += delta

        variant_results: list[dict[str, Any]] = []
        for stat in variant_stat_map.values():
            observed = int(stat["observed_action_count"] or 0)
            avg_delta = (float(stat["delta_sum"] or 0) / observed) if observed > 0 else 0.0
            variant_results.append(
                {
                    "strategy_variant": str(stat["strategy_variant"] or ""),
                    "action_count": int(stat["action_count"] or 0),
                    "observed_action_count": observed,
                    "avg_progress_delta_7d": round(avg_delta, 4),
                    "effect_signal": _signal_from_delta(avg_delta),
                }
            )
        variant_results.sort(
            key=lambda item: (
                -float(item.get("avg_progress_delta_7d") or 0),
                -int(item.get("observed_action_count") or 0),
                str(item.get("strategy_variant") or ""),
            )
        )

        winner_variant = ""
        winner_source = "insufficient"
        confidence_score = 0.0
        guardrail_note = "暂无可用样本，建议先执行实验并记录分组。"

        observed_items = [item for item in variant_results if int(item.get("observed_action_count") or 0) > 0]
        if observed_items:
            winner_variant = str(observed_items[0].get("strategy_variant") or "")
            winner_source = "observed"
            winner_delta = float(observed_items[0].get("avg_progress_delta_7d") or 0)
            second_delta = float(observed_items[1].get("avg_progress_delta_7d") or 0) if len(observed_items) > 1 else 0.0
            margin = winner_delta - second_delta

            sample_factor = min(1.0, float(observed_action_count) / 10.0)
            margin_factor = min(1.0, max(margin, 0.0) / 0.03)
            confidence_score = round(min(0.95, 0.35 + 0.45 * sample_factor + 0.20 * margin_factor), 2)

            if observed_action_count < 3:
                guardrail_note = "样本量不足，建议继续采样后再扩大执行。"
            elif margin < 0.01:
                guardrail_note = "分组差异接近，建议延长观察窗口并保持对照组。"
            elif winner_delta <= 0:
                guardrail_note = "当前优胜分组仍未体现正向增益，建议切换干预策略。"
            else:
                guardrail_note = "可按优胜分组扩量执行，并保留 10%-20% 对照组持续校验。"
        else:
            winner_variant = recommended_variant
            winner_source = "recommended_fallback"

        result_map[experiment_code] = {
            "analysis_window_days": normalized_window_days,
            "daily_progress_point_count": len(daily_series),
            "action_count": int(action_count),
            "observed_action_count": int(observed_action_count),
            "variant_results": variant_results,
            "winner_variant": winner_variant,
            "winner_source": winner_source,
            "confidence_score": confidence_score,
            "guardrail_note": guardrail_note,
        }

    return result_map

async def _record_invalid_response_mode_metric(
    *,
    user_id: int,
    endpoint: str,
    raw_mode: str | None,
    normalized_mode: str,
) -> None:
    try:
        db = await get_db()
        await db.execute(
            "INSERT INTO metrics (metric_type, metric_key, metric_value, metadata) VALUES (?, ?, ?, ?)",
            (
                "chat_guard",
                "invalid_response_mode",
                1.0,
                json.dumps(
                    {
                        "user_id": user_id,
                        "endpoint": endpoint,
                        "raw_mode": raw_mode,
                        "fallback_mode": normalized_mode,
                        "recorded_at": datetime.now(timezone.utc).isoformat(),
                    },
                    ensure_ascii=False,
                ),
            ),
        )
        await db.commit()
    except Exception as e:
        logger.warning("Failed to record invalid response_mode metric: %s", e)


_DOCX_XML_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
_XLSX_XML_NS = {
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}
_PPTX_XML_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

_CHAT_ATTACHMENT_TEXT_CHAR_LIMIT = 2400
_CHAT_ZIP_MAX_ENTRIES = 36
_CHAT_ZIP_MAX_PREVIEW_FILES = 8
_CHAT_ZIP_MAX_BYTES_PER_FILE = 256 * 1024
_CHAT_ZIP_MAX_NESTED_DEPTH = 1
_CHAT_ZIP_TEXT_SUFFIXES = {
    ".txt", ".md", ".csv", ".json", ".yaml", ".yml",
    ".log", ".sql", ".xml", ".html", ".py", ".js", ".ts", ".tsv", ".ini",
}
_CHAT_ATTACHMENT_ALLOWED_SUFFIXES = {
    ".txt", ".md", ".csv", ".json", ".yaml", ".yml",
    ".log", ".sql", ".xml", ".html", ".py", ".js", ".ts",
    ".xlsx", ".xls", ".docx", ".pdf", ".pptx", ".zip", ".png", ".jpg", ".jpeg", ".webp",
}
_CHAT_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
_CHAT_AUDIO_ALLOWED_SUFFIXES = {".webm", ".wav", ".mp3", ".m4a", ".ogg", ".aac", ".flac", ".mp4"}
_CHAT_AUDIO_ALLOWED_MIME_TYPES = {
    "audio/webm",
    "audio/wav",
    "audio/x-wav",
    "audio/mpeg",
    "audio/mp3",
    "audio/mp4",
    "audio/ogg",
    "audio/aac",
    "audio/flac",
    "video/webm",
}
_CHAT_UPLOAD_DIR = UPLOAD_DIR / "chat_attachments"
_CHAT_ASR_DEFAULT_MODEL = "whisper-1"
_CHAT_ASR_DEFAULT_TIMEOUT_SECONDS = 45
_CHAT_PDF_PREVIEW_MAX_PAGES = 8
_CHAT_PDF_PREVIEW_MAX_LINES_PER_PAGE = 220
_CHAT_PDF_HEADING_MAX_LEN = 72
_CHAT_PDF_NOISE_LINE_MAX_LEN = 42
_CHAT_OCR_MAX_VARIANTS = 4
_CHAT_OCR_PASS_CONFIGS = [
    ("chi_sim+eng", "--oem 3 --psm 6"),
    ("chi_sim+eng", "--oem 3 --psm 4"),
    ("chi_sim+eng", "--oem 3 --psm 11"),
    ("eng", "--oem 3 --psm 6"),
]


def _read_env_int(name: str, default: int, *, minimum: int = 0) -> int:
    try:
        value = int(str(os.getenv(name, str(default)) or str(default)).strip())
    except Exception:
        value = int(default)
    return max(minimum, value)


_CHAT_DOUBAO_IMAGE_ENABLED = os.getenv("DOUBAO_IMAGE_ENABLED", "1") == "1"
_CHAT_DOUBAO_IMAGE_MODEL = str(
    os.getenv("DOUBAO_IMAGE_MODEL", "") or ""
).strip()
_CHAT_DOUBAO_IMAGE_TIMEOUT_SECONDS = _read_env_int("DOUBAO_IMAGE_TIMEOUT_SECONDS", 18, minimum=6)
_CHAT_DOUBAO_IMAGE_MAX_TOKENS = _read_env_int("DOUBAO_IMAGE_MAX_TOKENS", 1200, minimum=300)
_CHAT_DOUBAO_IMAGE_RETRY_PER_MODEL = _read_env_int("DOUBAO_IMAGE_RETRY_PER_MODEL", 1, minimum=1)
_CHAT_DOUBAO_IMAGE_MAX_MODELS_PER_REQUEST = _read_env_int("DOUBAO_IMAGE_MAX_MODELS_PER_REQUEST", 3, minimum=1)
_CHAT_DOUBAO_IMAGE_TOTAL_BUDGET_SECONDS = _read_env_int("DOUBAO_IMAGE_TOTAL_BUDGET_SECONDS", 90, minimum=8)
_CHAT_DOUBAO_IMAGE_FAIL_FAST_NETWORK = os.getenv("DOUBAO_IMAGE_FAIL_FAST_NETWORK", "0") == "1"
_CHAT_DOUBAO_IMAGE_DISCOVER_MODELS = os.getenv("DOUBAO_IMAGE_DISCOVER_MODELS", "1") == "1"
_CHAT_DOUBAO_IMAGE_DISCOVER_TTL_SECONDS = _read_env_int("DOUBAO_IMAGE_DISCOVER_TTL_SECONDS", 900, minimum=60)
_CHAT_DOUBAO_IMAGE_DISCOVER_MAX_MODELS = _read_env_int("DOUBAO_IMAGE_DISCOVER_MAX_MODELS", 8, minimum=1)
_CHAT_DOUBAO_IMAGE_DISCOVER_CACHE: dict[str, Any] = {"ts": 0.0, "models": []}
_CHAT_OCR_SINGLE_PASS_TIMEOUT_SECONDS = _read_env_int("CHAT_OCR_SINGLE_PASS_TIMEOUT_SECONDS", 5, minimum=1)
_CHAT_OCR_TOTAL_TIMEOUT_SECONDS = _read_env_int("CHAT_OCR_TOTAL_TIMEOUT_SECONDS", 18, minimum=6)
_CHAT_DOUBAO_PLACEHOLDER_KEY_MARKERS = (
    "placeholder",
    "your_api_key",
    "change_me",
    "changeme",
)
_CHAT_DOUBAO_IMAGE_PROMPT = (
    "你是电商图像解析助手。请准确提取图片中的文字、数字、表格字段和值，"
    "按‘标题/字段/数据行/备注’输出为纯文本，不要编造图片中没有的信息。"
)

_CHAT_LAYOUT_PIPE_VERTICAL_CHARS = "│┃║¦｜"
_CHAT_LAYOUT_PIPE_NORMALIZED = "|"


def _safe_filename(filename: str) -> str:
    raw = str(filename or "upload.bin").strip() or "upload.bin"
    safe = raw.replace("/", "_").replace("\\", "_").replace("..", "_")
    return safe[:120] or "upload.bin"


def _is_allowed_voice_upload(*, suffix: str, content_type: str) -> bool:
    if suffix in _CHAT_AUDIO_ALLOWED_SUFFIXES:
        return True
    normalized = str(content_type or "").strip().lower()
    if normalized.startswith("audio/"):
        return True
    return normalized in _CHAT_AUDIO_ALLOWED_MIME_TYPES


def _derive_asr_url(raw_url: str | None) -> str:
    url = str(raw_url or "").strip().rstrip("/")
    if not url:
        return ""

    for suffix in ("/v1/audio/transcriptions", "/audio/transcriptions"):
        if url.endswith(suffix):
            return url

    base = url
    for suffix in ("/v1/chat/completions", "/chat/completions", "/v1/completions", "/completions"):
        if url.endswith(suffix):
            base = url[: -len(suffix)].rstrip("/")
            break

    if base.endswith(("/v1", "/v2", "/v3", "/v4")):
        return f"{base}/audio/transcriptions"
    return f"{base}/v1/audio/transcriptions"


def _resolve_asr_runtime() -> tuple[str, str, str, int]:
    env_url = str(os.getenv("ASR_API_URL", "") or "").strip()
    env_key = str(os.getenv("ASR_API_KEY", "") or "").strip()
    env_model = str(os.getenv("ASR_MODEL", "") or "").strip()
    env_timeout = str(os.getenv("ASR_TIMEOUT_SECONDS", "") or "").strip()

    llm_url = str(getattr(cfg, "LLM_API_URL", "") or "").strip()
    llm_key = str(getattr(cfg, "LLM_API_KEY", "") or "").strip()
    llm_timeout = int(getattr(cfg, "LLM_TIMEOUT_SECONDS", _CHAT_ASR_DEFAULT_TIMEOUT_SECONDS) or _CHAT_ASR_DEFAULT_TIMEOUT_SECONDS)

    timeout_seconds = llm_timeout if llm_timeout > 0 else _CHAT_ASR_DEFAULT_TIMEOUT_SECONDS
    if env_timeout:
        try:
            timeout_seconds = max(5, int(env_timeout))
        except Exception:
            timeout_seconds = max(5, timeout_seconds)

    # Ark chat/completions 并不等价于可直接推导的 ASR 端点。
    # 未显式配置 ASR_API_URL 时，避免将 /api/v3/chat/completions 误拼成
    # /api/v3/audio/transcriptions 后得到迷惑性的 404。
    if not env_url and "/api/v3/" in llm_url:
        return (
            "",
            env_key or llm_key,
            env_model or _CHAT_ASR_DEFAULT_MODEL,
            timeout_seconds,
        )

    return (
        _derive_asr_url(env_url or llm_url),
        env_key or llm_key,
        env_model or _CHAT_ASR_DEFAULT_MODEL,
        timeout_seconds,
    )


def _extract_text_from_asr_payload(payload: Any) -> str:
    if isinstance(payload, str):
        return payload.strip()

    if isinstance(payload, dict):
        for key in ("text", "transcript"):
            val = payload.get(key)
            if isinstance(val, str) and val.strip():
                return val.strip()

        for key in ("result", "data"):
            nested = payload.get(key)
            if isinstance(nested, (dict, str)):
                text = _extract_text_from_asr_payload(nested)
                if text:
                    return text

        for key in ("results", "segments"):
            nested_list = payload.get(key)
            if isinstance(nested_list, list):
                for item in nested_list:
                    text = _extract_text_from_asr_payload(item)
                    if text:
                        return text

    return ""


async def _transcribe_audio_bytes(
    *,
    raw: bytes,
    filename: str,
    content_type: str,
    language: str | None = None,
) -> dict[str, str]:
    asr_url, asr_key, asr_model, timeout_seconds = _resolve_asr_runtime()
    if not asr_url:
        raise HTTPException(status_code=503, detail="未配置语音转写服务，请设置 ASR_API_URL")

    headers: dict[str, str] = {}
    if asr_key:
        headers["Authorization"] = f"Bearer {asr_key}"

    data: dict[str, str] = {"model": asr_model}
    normalized_language = str(language or "").strip().lower()
    if normalized_language:
        data["language"] = normalized_language

    try:
        async with httpx.AsyncClient(timeout=timeout_seconds, trust_env=False) as client:
            response = await client.post(
                asr_url,
                headers=headers,
                data=data,
                files={
                    "file": (
                        filename,
                        raw,
                        content_type or "application/octet-stream",
                    )
                },
            )
    except httpx.TimeoutException:
        raise HTTPException(status_code=504, detail="语音转写服务超时，请稍后重试")
    except Exception as e:
        logger.warning("Voice ASR request failed: %s", e)
        raise HTTPException(status_code=502, detail="语音转写服务不可用，请稍后重试")

    if response.status_code >= 400:
        resp_text = str(response.text or "").strip().replace("\n", " ")
        logger.warning(
            "Voice ASR upstream error status=%s body=%s",
            response.status_code,
            resp_text[:220],
        )
        raise HTTPException(status_code=502, detail=f"语音转写服务调用失败（{response.status_code}）")

    try:
        payload: Any = response.json()
    except Exception:
        payload = {"text": response.text}

    transcript = _extract_text_from_asr_payload(payload).strip()
    if not transcript:
        raise HTTPException(status_code=502, detail="语音转写结果为空，请重试或更换录音格式")

    return {
        "text": transcript,
        "model": asr_model,
        "provider_url": asr_url,
    }


def _decode_text_bytes(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return raw.decode("utf-8", errors="ignore")


def _extract_csv_preview(raw: bytes) -> str:
    text = _decode_text_bytes(raw)
    reader = csv.reader(io.StringIO(text))
    lines: list[str] = []
    for idx, row in enumerate(reader):
        if idx >= 40:
            break
        lines.append(",".join(str(col).strip() for col in row[:12]))
    return "\n".join(lines)


def _extract_xlsx_preview_xml_fallback(raw: bytes) -> tuple[str, str, str]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            workbook_root = ET.fromstring(zf.read("xl/workbook.xml"))
            first_sheet = workbook_root.find(".//s:sheets/s:sheet", _XLSX_XML_NS)
            if first_sheet is None:
                return "", "partial", ""

            sheet_name = str(first_sheet.attrib.get("name") or "").strip()
            rel_id = str(
                first_sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                or first_sheet.attrib.get("r:id")
                or ""
            ).strip()

            rels_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
            target_path = ""
            for rel in rels_root.findall("pr:Relationship", _XLSX_XML_NS):
                rid = str(rel.attrib.get("Id") or "").strip()
                if rel_id and rid != rel_id:
                    continue
                target_path = str(rel.attrib.get("Target") or "").strip()
                if target_path:
                    break

            if not target_path:
                target_path = "worksheets/sheet1.xml"

            normalized_target = target_path.replace("\\", "/").lstrip("/")
            if not normalized_target.startswith("xl/"):
                normalized_target = f"xl/{normalized_target}"

            sheet_root = ET.fromstring(zf.read(normalized_target))

            shared_strings: list[str] = []
            if "xl/sharedStrings.xml" in zf.namelist():
                ss_root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
                for si in ss_root.findall("s:si", _XLSX_XML_NS):
                    parts = [str(node.text or "") for node in si.findall(".//s:t", _XLSX_XML_NS)]
                    shared_strings.append("".join(parts).strip())

            lines: list[str] = []
            if sheet_name:
                lines.append(f"# Sheet: {sheet_name}")

            row_nodes = sheet_root.findall(".//s:sheetData/s:row", _XLSX_XML_NS)
            for row_idx, row_node in enumerate(row_nodes):
                if row_idx >= 40:
                    break
                cells: list[str] = []
                for cell_node in row_node.findall("s:c", _XLSX_XML_NS)[:12]:
                    cell_type = str(cell_node.attrib.get("t") or "").strip().lower()
                    value = ""
                    if cell_type == "inlineStr":
                        t_node = cell_node.find(".//s:t", _XLSX_XML_NS)
                        value = str(t_node.text or "") if t_node is not None else ""
                    else:
                        v_node = cell_node.find("s:v", _XLSX_XML_NS)
                        raw_value = str(v_node.text or "").strip() if v_node is not None else ""
                        if cell_type == "s" and raw_value.isdigit():
                            index = int(raw_value)
                            if 0 <= index < len(shared_strings):
                                value = shared_strings[index]
                            else:
                                value = raw_value
                        else:
                            value = raw_value

                    value = re.sub(r"\s+", " ", str(value or "")).strip()
                    cells.append(value)

                row_text = "	".join(cells).strip()
                if row_text:
                    lines.append(row_text)

            rendered = "\n".join(lines).strip()
            if rendered:
                return rendered, "parsed", "xlsx-xml-fallback"
    except Exception:
        pass

    return "", "partial", ""


def _extract_xlsx_preview(raw: bytes) -> tuple[str, str, str]:
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        lines: list[str] = []
        sheet_name = str(ws.title or "").strip()
        if sheet_name:
            lines.append(f"# Sheet: {sheet_name}")
        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if idx >= 40:
                break
            cells = []
            for col in row[:12]:
                val = "" if col is None else str(col).strip()
                cells.append(val)
            row_text = "	".join(cells).strip()
            if row_text:
                lines.append(row_text)

        rendered = "\n".join(lines).strip()
        if rendered:
            return rendered, "parsed", "openpyxl"
    except Exception:
        pass

    return _extract_xlsx_preview_xml_fallback(raw)

def _extract_docx_preview_xml_fallback(raw: bytes) -> tuple[str, str, str]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            root = ET.fromstring(zf.read("word/document.xml"))

        lines: list[str] = []
        for p_node in root.findall(".//w:p", _DOCX_XML_NS):
            style_node = p_node.find("./w:pPr/w:pStyle", _DOCX_XML_NS)
            style_value = ""
            if style_node is not None:
                style_value = str(
                    style_node.attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
                    or style_node.attrib.get("w:val")
                    or style_node.attrib.get("val")
                    or ""
                ).strip().lower()

            fragments = [str(t_node.text or "") for t_node in p_node.findall(".//w:t", _DOCX_XML_NS)]
            paragraph = re.sub(r"\s+", " ", "".join(fragments)).strip()
            if not paragraph:
                continue

            if style_value.startswith("heading"):
                lines.append(f"# {paragraph}")
            else:
                lines.append(paragraph)

            if len(lines) >= 60:
                break

        rendered = "\n".join(lines).strip()
        if rendered:
            return rendered, "parsed", "docx-xml-fallback"
    except Exception:
        pass

    return "", "partial", ""


def _extract_docx_preview(raw: bytes) -> tuple[str, str, str]:
    try:
        import docx  # type: ignore

        doc = docx.Document(io.BytesIO(raw))
        lines: list[str] = []
        for p in doc.paragraphs:
            text = str(p.text or "").strip()
            if not text:
                continue
            try:
                style_name = str(getattr(p.style, "name", "") or "")
            except Exception:
                style_name = ""
            if style_name.lower().startswith("heading"):
                lines.append(f"# {text}")
            else:
                lines.append(text)
            if len(lines) >= 60:
                break

        rendered = "\n".join(lines).strip()
        if rendered:
            return rendered, "parsed", "python-docx"
    except Exception:
        pass

    return _extract_docx_preview_xml_fallback(raw)


def _extract_pptx_preview_xml_fallback(raw: bytes) -> tuple[str, str, str]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            slide_names = sorted(
                name for name in zf.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )
            if not slide_names:
                return "", "partial", ""

            lines: list[str] = []
            for slide_index, slide_name in enumerate(slide_names[:20], start=1):
                try:
                    slide_root = ET.fromstring(zf.read(slide_name))
                except Exception:
                    continue

                chunks: list[str] = []
                for text_node in slide_root.findall(".//a:t", _PPTX_XML_NS):
                    chunk = re.sub(r"\s+", " ", str(text_node.text or "")).strip()
                    if chunk:
                        chunks.append(chunk)

                if not chunks:
                    for raw_text in slide_root.itertext():
                        chunk = re.sub(r"\s+", " ", str(raw_text or "")).strip()
                        if chunk:
                            chunks.append(chunk)

                merged = re.sub(r"\s+", " ", " ".join(chunks)).strip()
                if not merged:
                    continue

                lines.append(f"# Slide {slide_index}")
                lines.append(merged)
                if len(lines) >= 80:
                    break

            rendered = "\n".join(lines).strip()
            if rendered:
                return rendered, "parsed", "pptx-xml-fallback"
    except Exception:
        pass

    return "", "partial", ""


def _extract_pptx_preview(raw: bytes) -> tuple[str, str, str]:
    try:
        from pptx import Presentation  # type: ignore

        presentation = Presentation(io.BytesIO(raw))
        lines: list[str] = []
        for slide_idx, slide in enumerate(presentation.slides, start=1):
            if slide_idx > 20:
                break

            snippets: list[str] = []
            for shape in slide.shapes:
                text_value = ""
                try:
                    if bool(getattr(shape, "has_text_frame", False)):
                        text_value = str(getattr(shape.text_frame, "text", "") or "")
                    elif hasattr(shape, "text"):
                        text_value = str(getattr(shape, "text", "") or "")
                except Exception:
                    text_value = ""

                normalized = re.sub(r"\s+", " ", text_value).strip()
                if not normalized:
                    continue
                snippets.append(normalized)
                if len(snippets) >= 10:
                    break

            if not snippets:
                continue

            lines.append(f"# Slide {slide_idx}")
            lines.extend(snippets)
            if len(lines) >= 80:
                break

        rendered = "\n".join(lines).strip()
        if rendered:
            return rendered, "parsed", "python-pptx"
    except Exception:
        pass

    return _extract_pptx_preview_xml_fallback(raw)


def _extract_zip_entry_preview(entry_name: str, entry_raw: bytes, *, depth: int) -> tuple[str, str]:
    suffix = Path(str(entry_name or "").strip()).suffix.lower()

    if suffix == ".csv":
        return _truncate_text(_extract_csv_preview(entry_raw), limit=720), "csv"
    if suffix in _CHAT_ZIP_TEXT_SUFFIXES:
        return _truncate_text(_decode_text_bytes(entry_raw), limit=720), "text-decoder"

    if suffix in {".xlsx", ".xls"}:
        preview_text, preview_status, preview_parser = _extract_xlsx_preview(entry_raw)
        if preview_status == "parsed" and preview_text:
            return _truncate_text(preview_text, limit=720), preview_parser or "openpyxl"
        return "", ""

    if suffix == ".docx":
        preview_text, preview_status, preview_parser = _extract_docx_preview(entry_raw)
        if preview_status == "parsed" and preview_text:
            return _truncate_text(preview_text, limit=720), preview_parser or "python-docx"
        return "", ""

    if suffix == ".pdf":
        preview_text, preview_status, preview_parser = _extract_pdf_preview(entry_raw)
        if preview_status == "parsed" and preview_text:
            return _truncate_text(preview_text, limit=720), preview_parser or "pypdf"
        return "", ""

    if suffix == ".pptx":
        preview_text, preview_status, preview_parser = _extract_pptx_preview(entry_raw)
        if preview_status == "parsed" and preview_text:
            return _truncate_text(preview_text, limit=720), preview_parser or "pptx-xml-fallback"
        return "", ""

    if suffix == ".zip" and depth < _CHAT_ZIP_MAX_NESTED_DEPTH:
        nested_text, nested_status, nested_parser = _extract_zip_preview(entry_raw, depth=depth + 1)
        if nested_text and nested_status in {"parsed", "partial"}:
            wrapped = "# 嵌套 ZIP 预览\n" + nested_text
            return _truncate_text(wrapped, limit=720), nested_parser or "zip-manifest"
        return "", ""

    return "", ""


def _extract_zip_preview(raw: bytes, *, depth: int = 0) -> tuple[str, str, str]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            entries = [entry for entry in zf.infolist() if not bool(entry.is_dir())]
            if not entries:
                return "", "partial", ""

            lines: list[str] = ["# ZIP 文件清单"]
            for entry in entries[:_CHAT_ZIP_MAX_ENTRIES]:
                entry_name = str(entry.filename or "").strip()
                lines.append(f"- {entry_name} ({int(entry.file_size or 0)} bytes)")

            preview_count = 0
            advanced_preview_count = 0
            for entry in entries:
                if preview_count >= _CHAT_ZIP_MAX_PREVIEW_FILES:
                    break

                entry_name = str(entry.filename or "").strip()
                file_size = int(entry.file_size or 0)
                if file_size <= 0 or file_size > _CHAT_ZIP_MAX_BYTES_PER_FILE:
                    continue

                try:
                    entry_raw = zf.read(entry)
                except Exception:
                    continue

                preview_text, preview_parser = _extract_zip_entry_preview(entry_name, entry_raw, depth=depth)
                if not preview_text:
                    continue

                lines.append("")
                lines.append(f"[文件] {entry_name}")
                if preview_parser:
                    lines.append(f"(parser: {preview_parser})")
                lines.append(preview_text)
                preview_count += 1

                parser_token = str(preview_parser or "").lower()
                if parser_token not in {"", "text-decoder", "csv"}:
                    advanced_preview_count += 1

            rendered = "\n".join(lines).strip()
            if preview_count > 0:
                parser = "zip-hybrid-bundle" if advanced_preview_count > 0 else "zip-text-bundle"
                return rendered, "parsed", parser
            if rendered:
                return rendered, "partial", "zip-manifest"
    except Exception:
        pass

    return "", "partial", ""

def _normalize_preview_line(raw: str) -> str:
    text = str(raw or "").replace("\x00", "").strip()
    if not text:
        return ""
    return re.sub(r"\s+", " ", text)


def _is_pdf_page_marker_line(line: str) -> bool:
    text = _normalize_preview_line(line)
    if not text:
        return False

    compact = text.lower()
    if re.match(r"^\[?第\s*\d+\s*页(?:\s*/\s*共\s*\d+\s*页)?\]?$", text):
        return True
    if re.match(r"^(?:page|p)\s*\d+(?:\s*(?:/|of)\s*\d+)?$", compact):
        return True
    if re.match(r"^第\s*\d+\s*页\s*共\s*\d+\s*页$", text):
        return True
    return False


def _looks_like_pdf_heading(line: str) -> bool:
    text = _normalize_preview_line(line).lstrip("#").strip()
    if not text:
        return False
    if len(text) > _CHAT_PDF_HEADING_MAX_LEN:
        return False
    if _is_pdf_page_marker_line(text):
        return False

    if re.match(r"^(第[一二三四五六七八九十百千0-9]+[章节篇卷部])", text):
        return True
    if re.match(r"^(chapter|section)\s*\d+", text, re.IGNORECASE):
        return True
    if re.match(r"^\d{1,2}(?:\.\d{1,3}){0,3}\s+", text):
        return True
    if re.match(r"^[（(]?\d+[)）][\s　]", text):
        return True

    if text.endswith((":", "：")) and len(text) <= 32:
        return True
    return False


def _build_structured_pdf_preview(pages_text: list[str]) -> tuple[str, dict[str, int]]:
    pages: list[list[str]] = []
    for page_text in pages_text[:_CHAT_PDF_PREVIEW_MAX_PAGES]:
        lines: list[str] = []
        for raw_line in str(page_text or "").splitlines()[:_CHAT_PDF_PREVIEW_MAX_LINES_PER_PAGE]:
            line = str(raw_line or "").replace("\x00", "").strip()
            if line:
                lines.append(line)
        pages.append(lines)

    edge_counter: Counter[str] = Counter()
    for lines in pages:
        if not lines:
            continue
        edge_candidates = lines[:2] + lines[-2:]
        for candidate in edge_candidates:
            if len(candidate) <= _CHAT_PDF_NOISE_LINE_MAX_LEN:
                edge_counter[candidate] += 1

    repeated_noise = {
        line
        for line, count in edge_counter.items()
        if count >= 2 and not _looks_like_pdf_heading(line) and not _is_pdf_page_marker_line(line)
    }

    output: list[str] = []
    section_count = 0
    removed_noise_lines = 0
    table_block_count = 0
    table_row_count = 0

    for page_idx, lines in enumerate(pages, 1):
        if not lines:
            continue

        output.append(f"[第{page_idx}页]")

        cleaned_lines: list[str] = []
        for line in lines:
            if line in repeated_noise or _is_pdf_page_marker_line(line):
                removed_noise_lines += 1
                continue
            cleaned_lines.append(line)

            if _looks_like_pdf_heading(line):
                heading = _normalize_preview_line(line).lstrip("#").strip()
                if heading:
                    output.append(f"## {heading}")
                    section_count += 1
                continue

            output.append(line)

        table_rows = _extract_layout_table_rows(cleaned_lines)
        if table_rows:
            table_block_count += 1
            table_row_count += len(table_rows)
            output.append("[结构化表格候选]")
            for row in table_rows[:24]:
                output.append("\t".join(row[:8]))

        if output and output[-1] != "":
            output.append("")

    structured = "\n".join(output).strip()
    return structured, {
        "page_count": sum(1 for lines in pages if lines),
        "section_count": section_count,
        "removed_noise_lines": removed_noise_lines,
        "table_block_count": table_block_count,
        "table_row_count": table_row_count,
    }


def _pick_preview_summary_line(text: str) -> str:
    for raw_line in str(text or "").splitlines():
        line = _normalize_preview_line(raw_line)
        if not line:
            continue
        if re.match(r"^\[第\d+页\]$", line):
            continue

        stripped = line.lstrip("#").strip()
        if not stripped:
            continue
        if _is_pdf_page_marker_line(stripped):
            continue
        return stripped[:120]
    return ""


def _split_space_aligned_cells(line: str) -> list[str]:
    raw = str(line or "").replace("\x00", "")
    if not raw.strip():
        return []

    text = raw.strip()
    cells = [seg.strip() for seg in re.split(r"[ 	　]{2,}", text) if seg and seg.strip()]
    if len(cells) < 2:
        return []
    return cells[:10]


def _is_layout_numeric_like_cell(value: str) -> bool:
    text = str(value or "").strip()
    if not text:
        return False

    if re.fullmatch(r"[+\-]?\d[\d,，._%:/-]*", text):
        return True
    if re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", text):
        return True
    return False


def _is_layout_sentence_like_cell(value: str) -> bool:
    text = str(value or "").strip()
    if not text:
        return False

    if len(text) >= 28:
        return True
    if len(text) >= 18 and re.search(r"[，。；：,.!?！？]", text):
        return True
    return False


def _normalize_space_aligned_segment(segment: list[list[str]]) -> list[list[str]]:
    if not segment:
        return []

    col_counter: Counter[int] = Counter()
    for row in segment:
        if not isinstance(row, list):
            continue
        col_counter[len(row)] += 1
    if not col_counter:
        return []

    common_col_count = col_counter.most_common(1)[0][0]
    normalized_rows: list[list[str]] = []
    for row in segment:
        if len(row) < max(2, common_col_count - 1):
            continue
        normalized_rows.append(row[:common_col_count])
        if len(normalized_rows) >= 32:
            break

    if len(normalized_rows) < 2:
        return []
    return normalized_rows


def _score_space_aligned_segment(rows: list[list[str]]) -> float:
    if len(rows) < 2:
        return 0.0

    col_count = Counter(len(row) for row in rows).most_common(1)[0][0]
    stable_rows = [row[:col_count] for row in rows if len(row) >= max(2, col_count - 1)]
    if len(stable_rows) < 2:
        return 0.0

    header = stable_rows[0]
    header_non_empty = [str(cell or "").strip() for cell in header if str(cell or "").strip()]
    if len(header_non_empty) < 2:
        return 0.0

    header_text_hits = sum(1 for cell in header_non_empty if re.search(r"[A-Za-z一-鿿]", cell))
    header_numeric_hits = sum(1 for cell in header_non_empty if _is_layout_numeric_like_cell(cell))
    header_like = (
        header_text_hits >= max(1, len(header_non_empty) // 2)
        and header_numeric_hits <= max(1, len(header_non_empty) // 2)
    )

    data_rows = stable_rows[1:]
    numeric_rows = 0
    row_signal = 0.0
    paragraph_like_rows = 0
    compact_cells = 0
    sentence_like_cells = 0
    total_cells = 0

    for row in stable_rows:
        non_empty = [str(cell or "").strip() for cell in row if str(cell or "").strip()]
        if not non_empty:
            continue

        text_hits = sum(1 for cell in non_empty if re.search(r"[A-Za-z一-鿿]", cell))
        numeric_hits = sum(1 for cell in non_empty if _is_layout_numeric_like_cell(cell))
        sentence_hits = sum(1 for cell in non_empty if _is_layout_sentence_like_cell(cell))

        total_cells += len(non_empty)
        compact_cells += sum(1 for cell in non_empty if len(cell) <= 14)
        sentence_like_cells += sentence_hits

        if row is header:
            continue

        if numeric_hits >= 1:
            numeric_rows += 1
            row_signal += 2.2
        if numeric_hits >= 1 and text_hits >= 1:
            row_signal += 1.2
        if len(non_empty) >= max(2, col_count - 1):
            row_signal += 0.6

        if sentence_hits >= 1 and numeric_hits == 0 and any(len(cell) >= 18 for cell in non_empty):
            paragraph_like_rows += 1

    if len(data_rows) < 1:
        return 0.0

    if numeric_rows == 0 and col_count <= 2 and len(stable_rows) < 4:
        return 0.0
    if numeric_rows == 0 and col_count <= 2 and paragraph_like_rows >= 2:
        return 0.0

    col_consistency = float(sum(1 for row in rows if len(row) == col_count)) / float(max(1, len(rows)))
    compact_ratio = float(compact_cells) / float(max(1, total_cells))
    sentence_ratio = float(sentence_like_cells) / float(max(1, total_cells))

    score = 0.0
    score += float(len(stable_rows)) * 5.2
    score += float(col_count) * 2.8
    score += col_consistency * 14.0
    score += row_signal
    score += compact_ratio * 10.0

    if header_like:
        score += 10.0
    if numeric_rows == 0:
        score -= 8.0
    if paragraph_like_rows > 0:
        score -= float(paragraph_like_rows) * 8.0
    score -= sentence_ratio * 18.0

    return round(max(0.0, score), 3)


def _extract_space_aligned_table_rows(lines: list[str]) -> list[list[str]]:
    segments: list[list[list[str]]] = []
    current: list[list[str]] = []

    for raw_line in lines:
        raw_text = str(raw_line or "").replace("\x00", "")
        if not raw_text.strip():
            if len(current) >= 2:
                segments.append(current)
            current = []
            continue

        cells = _split_space_aligned_cells(raw_text)
        if not cells:
            if len(current) >= 2:
                segments.append(current)
            current = []
            continue

        if current and abs(len(cells) - len(current[-1])) > 1:
            if len(current) >= 2:
                segments.append(current)
            current = [cells]
            continue

        current.append(cells)

    if len(current) >= 2:
        segments.append(current)

    if not segments:
        return []

    candidates: list[tuple[int, list[list[str]], float, str]] = []

    for seg_idx, segment in enumerate(segments):
        normalized_rows = _normalize_space_aligned_segment(segment)
        if not normalized_rows:
            continue

        score = _score_space_aligned_segment(normalized_rows)
        if score <= 0:
            continue
        signature = _table_header_signature(normalized_rows)
        candidates.append((seg_idx, normalized_rows, score, signature))

    if not candidates:
        return []
    merged_rows = _merge_table_candidates(candidates, min_rows=2, score_floor=6.0)
    if len(merged_rows) < 2:
        return []
    return merged_rows[:32]



def _normalize_ruled_table_line(line: str) -> str:
    text = str(line or "").replace("\x00", "").strip()
    if not text:
        return ""

    normalized = text
    for ch in _CHAT_LAYOUT_PIPE_VERTICAL_CHARS:
        normalized = normalized.replace(ch, _CHAT_LAYOUT_PIPE_NORMALIZED)
    normalized = re.sub(r"\s*\|\s*", "|", normalized)
    normalized = re.sub(r"\|{2,}", "|", normalized)
    return normalized


def _is_ruled_table_border_line(line: str) -> bool:
    text = _normalize_ruled_table_line(line)
    if not text:
        return False

    compact = re.sub(r"\s+", "", text)
    if len(compact) < 3:
        return False

    if "|" not in compact and "+" not in compact and not re.search(r"[┼┬┴├┤┌┐└┘╔╗╚╝╦╩╠╣╬─━═]", compact):
        return False

    return bool(re.fullmatch(r"[+|=:_~.`\-┼┬┴├┤┌┐└┘╔╗╚╝╦╩╠╣╬─━═]+", compact))


def _split_ruled_table_cells(line: str) -> list[str]:
    normalized = _normalize_ruled_table_line(line)
    if not normalized or "|" not in normalized:
        return []
    if _is_ruled_table_border_line(normalized):
        return []

    parts = [part.strip() for part in normalized.split("|")]
    while parts and not parts[0]:
        parts.pop(0)
    while parts and not parts[-1]:
        parts.pop()

    if len(parts) < 2:
        return []

    non_empty = [cell for cell in parts if cell]
    if len(non_empty) < 2:
        return []

    # 仅保留结构化概率较高的行，避免长段落被误认为两列表格。
    if len(parts) <= 2:
        sentence_like = sum(1 for cell in non_empty if _is_layout_sentence_like_cell(cell))
        if sentence_like >= 1 and not any(_is_layout_numeric_like_cell(cell) for cell in non_empty):
            return []

    if len(non_empty) < max(2, len(parts) // 2):
        return []

    return parts[:10]



def _repair_ruled_table_row(cells: list[str], *, expected_cols: int) -> list[str]:
    normalized = [str(cell or "").strip() for cell in cells]
    if expected_cols <= 1 or len(normalized) < 2:
        return []

    if len(normalized) == expected_cols:
        return normalized[:expected_cols]

    if len(normalized) > expected_cols:
        kept = normalized[: expected_cols - 1]
        tail = " ".join(part for part in normalized[expected_cols - 1 :] if part).strip()
        kept.append(tail)
        return kept

    # 缺列时尽量补齐，避免 OCR 轻微断裂导致整行被丢弃。
    if len(normalized) >= max(2, expected_cols - 1):
        repaired = list(normalized)
        while len(repaired) < expected_cols:
            repaired.append("")
        return repaired[:expected_cols]

    return []


def _normalize_ruled_table_segment(segment: list[list[str]]) -> list[list[str]]:
    if not segment:
        return []

    col_counter: Counter[int] = Counter()
    for row in segment:
        if not isinstance(row, list):
            continue
        row_len = len([cell for cell in row if str(cell or "").strip()])
        if row_len >= 2:
            col_counter[row_len] += 1

    if not col_counter:
        return []

    expected_cols = col_counter.most_common(1)[0][0]
    expected_cols = max(2, min(10, expected_cols))

    normalized_rows: list[list[str]] = []
    for row in segment:
        repaired = _repair_ruled_table_row(row, expected_cols=expected_cols)
        if not repaired:
            continue

        non_empty = sum(1 for cell in repaired if str(cell or "").strip())
        if non_empty < max(2, expected_cols // 2):
            continue

        normalized_rows.append(repaired)
        if len(normalized_rows) >= 32:
            break

    if len(normalized_rows) < 2:
        return []
    return normalized_rows



def _split_markdown_table_cells(line: str) -> list[str]:
    raw = str(line or "").replace("\x00", "").strip()
    if not raw or "|" not in raw:
        return []

    # 代码块或注释片段中的竖线不参与 markdown 表格识别。
    if raw.startswith("```") or raw.startswith("~~~"):
        return []

    text = raw
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]

    parts = [part.strip() for part in text.split("|")]
    if len(parts) < 2:
        return []

    # 过滤空列，避免 "|||" 误命中
    non_empty = [cell for cell in parts if cell]
    if len(non_empty) < 2:
        return []

    # 两列表头如果都是长句且无数值，通常是说明文本而不是表格。
    if len(non_empty) <= 2:
        sentence_like = sum(1 for cell in non_empty if _is_layout_sentence_like_cell(cell))
        has_numeric = any(_is_layout_numeric_like_cell(cell) for cell in non_empty)
        if sentence_like >= len(non_empty) and not has_numeric:
            return []

    return parts[:10]


def _is_markdown_table_separator_row(cells: list[str]) -> bool:
    if not cells or len(cells) < 2:
        return False

    hit = 0
    for cell in cells:
        token = str(cell or "").strip().replace(" ", "")
        if not token:
            continue
        if re.fullmatch(r":?-{3,}:?", token):
            hit += 1
            continue
        return False

    return hit >= 2


def _is_markdown_table_header_like(cells: list[str]) -> bool:
    if not cells:
        return False

    non_empty = [str(cell or "").strip() for cell in cells if str(cell or "").strip()]
    if len(non_empty) < 2:
        return False

    text_hits = sum(1 for cell in non_empty if re.search(r"[A-Za-z一-鿿]", cell))
    numeric_hits = sum(1 for cell in non_empty if _is_layout_numeric_like_cell(cell))
    sentence_hits = sum(1 for cell in non_empty if _is_layout_sentence_like_cell(cell))

    if text_hits < max(1, len(non_empty) // 2):
        return False
    if numeric_hits > max(1, len(non_empty) // 2):
        return False
    if sentence_hits >= len(non_empty):
        return False
    return True


def _build_markdown_rows_from_separator(segment: list[list[str]], separator_idx: int) -> list[list[str]]:
    if separator_idx < 1 or separator_idx + 1 >= len(segment):
        return []

    header = segment[separator_idx - 1]
    if not _is_markdown_table_header_like(header):
        return []

    col_count = len([cell for cell in header if str(cell or "").strip()])
    col_count = max(2, min(10, col_count))

    normalized: list[list[str]] = []

    head = [str(cell or "").strip() for cell in header[:col_count]]
    if len(head) < col_count:
        head += [""] * (col_count - len(head))
    normalized.append(head)

    for row in segment[separator_idx + 1 :]:
        if _is_markdown_table_separator_row(row):
            break

        values = [str(cell or "").strip() for cell in row[:col_count]]
        if len(values) < col_count:
            values += [""] * (col_count - len(values))

        non_empty = sum(1 for value in values if value)
        if non_empty < max(2, col_count // 2):
            continue

        normalized.append(values)
        if len(normalized) >= 32:
            break

    if len(normalized) < 3:
        return []
    return normalized


def _normalize_markdown_table_segment(segment: list[list[str]]) -> list[list[str]]:
    if len(segment) < 3:
        return []

    best_rows: list[list[str]] = []
    best_score = 0.0

    for idx, row in enumerate(segment):
        if not _is_markdown_table_separator_row(row):
            continue

        candidate_rows = _build_markdown_rows_from_separator(segment, idx)
        if not candidate_rows:
            continue

        numeric_rows = 0
        for data_row in candidate_rows[1:]:
            if any(_is_layout_numeric_like_cell(cell) for cell in data_row):
                numeric_rows += 1

        score = _score_space_aligned_segment(candidate_rows)
        score += min(6.0, float(len(candidate_rows)) * 0.6)
        score += min(4.0, float(numeric_rows) * 0.8)

        if score <= best_score:
            continue

        best_score = score
        best_rows = candidate_rows

    if len(best_rows) < 3:
        return []
    return best_rows


def _table_header_signature(rows: list[list[str]]) -> str:
    if not rows:
        return ""

    header = rows[0]
    tokens: list[str] = []
    for cell in header:
        token = re.sub(r"\s+", "", str(cell or "").strip().lower())
        if token:
            tokens.append(token)
    if len(tokens) < 2:
        return ""
    return "|".join(tokens[:8])


def _table_numeric_profile(rows: list[list[str]]) -> list[float]:
    if not rows or len(rows) < 2:
        return []

    col_count = len(rows[0]) if rows[0] else 0
    if col_count <= 0:
        return []

    data_rows = rows[1:]
    if not data_rows:
        return []

    ratios: list[float] = []
    for col_idx in range(col_count):
        hits = 0
        total = 0
        for row in data_rows:
            if col_idx >= len(row):
                continue
            cell = str(row[col_idx] or "").strip()
            if not cell:
                continue
            total += 1
            if _is_layout_numeric_like_cell(cell):
                hits += 1

        if total <= 0:
            ratios.append(0.0)
        else:
            ratios.append(float(hits) / float(total))

    return ratios


def _table_profiles_compatible(anchor_rows: list[list[str]], candidate_rows: list[list[str]]) -> bool:
    if not anchor_rows or not candidate_rows:
        return False
    if len(anchor_rows[0]) != len(candidate_rows[0]):
        return False

    anchor_profile = _table_numeric_profile(anchor_rows)
    candidate_profile = _table_numeric_profile(candidate_rows)
    if not anchor_profile or not candidate_profile:
        return True

    mismatch = 0
    strong_anchor_numeric = 0
    for left, right in zip(anchor_profile, candidate_profile):
        if left >= 0.6:
            strong_anchor_numeric += 1
            if right <= 0.2:
                mismatch += 1
                continue
        if left <= 0.2 and right >= 0.8:
            mismatch += 1

    # 至少两列强语义冲突时判定为漂移，拒绝合并。
    if strong_anchor_numeric >= 2 and mismatch >= 2:
        return False
    return True


def _markdown_table_header_signature(rows: list[list[str]]) -> str:
    return _table_header_signature(rows)


def _markdown_table_numeric_profile(rows: list[list[str]]) -> list[float]:
    return _table_numeric_profile(rows)


def _markdown_table_profiles_compatible(anchor_rows: list[list[str]], candidate_rows: list[list[str]]) -> bool:
    return _table_profiles_compatible(anchor_rows, candidate_rows)


def _merge_table_candidates(
    candidates: list[tuple[int, list[list[str]], float, str]],
    *,
    min_rows: int,
    score_ratio: float = 0.42,
    score_floor: float = 8.0,
) -> list[list[str]]:
    if not candidates:
        return []

    anchor_idx, anchor_rows, anchor_score, anchor_sig = max(candidates, key=lambda x: x[2])
    if len(anchor_rows) < max(2, min_rows):
        return []

    merged: list[list[str]] = [anchor_rows[0]]
    seen_data_rows: set[tuple[str, ...]] = set()
    for row in anchor_rows[1:]:
        key = tuple(str(cell or "") for cell in row)
        if key in seen_data_rows:
            continue
        merged.append(row)
        seen_data_rows.add(key)
        if len(merged) >= 32:
            return merged[:32]

    score_threshold = max(float(score_floor), float(anchor_score) * float(score_ratio))

    for idx, rows, score, signature in sorted(candidates, key=lambda x: x[0]):
        if idx == anchor_idx:
            continue
        if anchor_sig and signature != anchor_sig:
            continue
        if score < score_threshold:
            continue
        if not _table_profiles_compatible(anchor_rows, rows):
            continue

        for row in rows[1:]:
            key = tuple(str(cell or "") for cell in row)
            if key in seen_data_rows:
                continue
            merged.append(row)
            seen_data_rows.add(key)
            if len(merged) >= 32:
                return merged[:32]

    return merged[:32]


def _merge_markdown_table_candidates(
    candidates: list[tuple[int, list[list[str]], float, str]],
) -> list[list[str]]:
    return _merge_table_candidates(candidates, min_rows=3, score_floor=8.0)


def _extract_markdown_table_rows(lines: list[str]) -> list[list[str]]:
    segments: list[list[list[str]]] = []
    current: list[list[str]] = []

    for raw_line in lines:
        cells = _split_markdown_table_cells(raw_line)
        if not cells:
            if len(current) >= 3:
                segments.append(current)
            current = []
            continue

        if current and abs(len(cells) - len(current[-1])) > 2:
            if len(current) >= 3:
                segments.append(current)
            current = [cells]
            continue

        current.append(cells)

    if len(current) >= 3:
        segments.append(current)

    if not segments:
        return []

    candidates: list[tuple[int, list[list[str]], float, str]] = []
    for seg_idx, segment in enumerate(segments):
        normalized_rows = _normalize_markdown_table_segment(segment)
        if not normalized_rows:
            continue

        signature = _table_header_signature(normalized_rows)
        score = _score_space_aligned_segment(normalized_rows) + 2.2
        candidates.append((seg_idx, normalized_rows, score, signature))

    if not candidates:
        return []

    merged_rows = _merge_markdown_table_candidates(candidates)
    if len(merged_rows) < 3:
        return []
    return merged_rows[:32]


def _extract_ruled_table_rows(lines: list[str]) -> list[list[str]]:

    segments: list[list[list[str]]] = []
    current: list[list[str]] = []

    for raw_line in lines:
        raw_text = str(raw_line or "").replace("\x00", "")
        if not raw_text.strip():
            if len(current) >= 2:
                segments.append(current)
            current = []
            continue

        if _is_ruled_table_border_line(raw_text):
            if current:
                continue
            continue

        cells = _split_ruled_table_cells(raw_text)
        if not cells:
            if len(current) >= 2:
                segments.append(current)
            current = []
            continue

        if current and abs(len(cells) - len(current[-1])) > 2:
            if len(current) >= 2:
                segments.append(current)
            current = [cells]
            continue

        current.append(cells)

    if len(current) >= 2:
        segments.append(current)

    if not segments:
        return []

    candidates: list[tuple[int, list[list[str]], float, str]] = []

    for seg_idx, segment in enumerate(segments):
        normalized_rows = _normalize_ruled_table_segment(segment)
        if not normalized_rows:
            continue

        score = _score_space_aligned_segment(normalized_rows) + 1.8
        if score <= 0:
            continue
        signature = _table_header_signature(normalized_rows)
        candidates.append((seg_idx, normalized_rows, score, signature))

    if not candidates:
        return []

    merged_rows = _merge_table_candidates(candidates, min_rows=2, score_floor=6.0)
    if len(merged_rows) < 2:
        return []
    return merged_rows[:32]


def _extract_layout_table_rows_with_source(lines: list[str]) -> tuple[list[list[str]], str]:
    space_rows = _extract_space_aligned_table_rows(lines)
    ruled_rows = _extract_ruled_table_rows(lines)
    markdown_rows = _extract_markdown_table_rows(lines)

    best_rows: list[list[str]] = []
    best_score = 0.0
    best_source = ""

    candidates = [
        (space_rows, 0.0, "space"),
        (ruled_rows, 0.5, "ruled"),
        (markdown_rows, 0.8, "markdown"),
    ]

    for rows, bonus, source in candidates:
        if len(rows) < 2:
            continue

        score = _score_space_aligned_segment(rows) + float(bonus)
        if score <= best_score:
            continue

        best_score = score
        best_rows = rows
        best_source = source

    if len(best_rows) < 2:
        return [], ""
    return best_rows[:32], best_source


def _extract_layout_table_rows(lines: list[str]) -> list[list[str]]:
    rows, _source = _extract_layout_table_rows_with_source(lines)
    return rows


def _append_layout_table_candidates_with_source(text: str) -> tuple[str, int, str]:
    lines = [line for line in str(text or "").splitlines() if _normalize_preview_line(line)]
    table_rows, layout_source = _extract_layout_table_rows_with_source(lines)
    if not table_rows:
        return str(text or ""), 0, ""

    rendered_rows = ["	".join(row[:8]) for row in table_rows[:24]]
    if not rendered_rows:
        return str(text or ""), 0, ""

    base = str(text or "").strip()
    candidate_block = "[结构化表格候选]\n" + "\n".join(rendered_rows)

    if base:
        if candidate_block in base:
            return base, len(rendered_rows), layout_source
        merged = f"{base}\n\n{candidate_block}".strip()
    else:
        merged = candidate_block

    return merged, len(rendered_rows), layout_source


def _append_layout_table_candidates(text: str) -> tuple[str, int]:
    merged, row_count, _layout_source = _append_layout_table_candidates_with_source(text)
    return merged, row_count


def _infer_layout_source_from_text(text: str) -> str:
    raw = str(text or "").strip()
    if not raw:
        return ""

    lines = [line for line in raw.splitlines() if _normalize_preview_line(line)]
    if not lines:
        return ""

    marker = "[结构化表格候选]"
    candidate_lines = lines
    if marker in lines:
        marker_idx = lines.index(marker)
        if marker_idx > 0:
            candidate_lines = lines[:marker_idx]
        else:
            candidate_lines = lines[marker_idx + 1 :]

    _rows, source = _extract_layout_table_rows_with_source(candidate_lines)
    if source:
        return source

    _rows2, source2 = _extract_layout_table_rows_with_source(lines)
    return source2 or ""


_PDF_TJ_LITERAL_RE = re.compile(rb"(\((?:\\.|[^\\()])*\))\s*Tj")
_PDF_TJ_ARRAY_RE = re.compile(rb"\[(.*?)\]\s*TJ", re.S)
_PDF_LITERAL_RE = re.compile(rb"\((?:\\.|[^\\()])*\)")


def _decode_pdf_literal_bytes(token: bytes) -> str:
    data = bytes(token or b"")
    if len(data) >= 2 and data[:1] == b"(" and data[-1:] == b")":
        data = data[1:-1]

    output = bytearray()
    idx = 0
    while idx < len(data):
        byte = data[idx]
        if byte != 0x5C:  # backslash
            output.append(byte)
            idx += 1
            continue

        idx += 1
        if idx >= len(data):
            break

        esc = data[idx]
        idx += 1

        if esc == 110:  # n
            output.append(10)
            continue
        if esc == 114:  # r
            output.append(13)
            continue
        if esc == 116:  # t
            output.append(9)
            continue
        if esc == 98:  # b
            output.append(8)
            continue
        if esc == 102:  # f
            output.append(12)
            continue
        if esc in (40, 41, 92):  # ( ) \
            output.append(esc)
            continue

        if 48 <= esc <= 55:  # octal
            oct_digits = [esc]
            for _ in range(2):
                if idx < len(data) and 48 <= data[idx] <= 55:
                    oct_digits.append(data[idx])
                    idx += 1
                else:
                    break
            try:
                output.append(int(bytes(oct_digits).decode("ascii"), 8) & 0xFF)
            except Exception:
                pass
            continue

        output.append(esc)

    text_value = output.decode("utf-8", errors="ignore").strip()
    if text_value:
        return text_value
    return output.decode("latin-1", errors="ignore").strip()


def _extract_pdf_stream_literal_text(raw: bytes) -> str:
    if not raw:
        return ""

    lines: list[str] = []
    for stream_match in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", raw, flags=re.S):
        stream_bytes = bytes(stream_match.group(1) or b"")
        prefix = raw[max(0, stream_match.start() - 260): stream_match.start()]

        if b"/FlateDecode" in prefix:
            try:
                stream_bytes = zlib.decompress(stream_bytes)
            except Exception:
                try:
                    stream_bytes = zlib.decompress(stream_bytes, -zlib.MAX_WBITS)
                except Exception:
                    continue

        for item in _PDF_TJ_LITERAL_RE.findall(stream_bytes):
            text_line = _decode_pdf_literal_bytes(item)
            text_line = re.sub(r"\s+", " ", text_line).strip()
            if text_line:
                lines.append(text_line)

        for arr_match in _PDF_TJ_ARRAY_RE.findall(stream_bytes):
            parts: list[str] = []
            for literal in _PDF_LITERAL_RE.findall(arr_match):
                literal_text = _decode_pdf_literal_bytes(literal)
                if literal_text:
                    parts.append(literal_text)
            merged = re.sub(r"\s+", " ", "".join(parts)).strip()
            if merged:
                lines.append(merged)

        if len(lines) >= 280:
            break

    deduped: list[str] = []
    seen: set[str] = set()
    for line in lines:
        normalized = str(line).strip()
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        deduped.append(normalized)
        if len(deduped) >= 180:
            break

    return "\n".join(deduped).strip()

def _extract_pdf_preview(raw: bytes) -> tuple[str, str, str]:
    # 优先 pypdf（纯 Python，部署最常见）；失败后自动降级
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(io.BytesIO(raw))
        pages_text: list[str] = []
        for page_idx, page in enumerate(reader.pages):
            if page_idx >= _CHAT_PDF_PREVIEW_MAX_PAGES:
                break
            text = str(page.extract_text() or "").strip()
            if text:
                pages_text.append(text)

        if pages_text:
            structured_text, meta = _build_structured_pdf_preview(pages_text)
            if structured_text:
                parser = "pypdf-structured" if (meta.get("section_count", 0) > 0 or meta.get("removed_noise_lines", 0) > 0 or meta.get("table_block_count", 0) > 0) else "pypdf"
                return structured_text, "parsed", parser
    except Exception:
        pass

    fallback_text = _extract_pdf_stream_literal_text(raw)
    if fallback_text:
        cleaned = _clean_ocr_text(fallback_text)
        if cleaned:
            return cleaned, "parsed", "pdf-stream-fallback"

    return "", "partial", ""

def _clean_ocr_text(raw_text: str) -> str:
    lines: list[str] = []
    for raw_line in str(raw_text or "").splitlines():
        line = _normalize_preview_line(raw_line)
        if not line:
            continue
        if _is_pdf_page_marker_line(line):
            continue
        lines.append(line)
    return "\n".join(lines).strip()


def _derive_chat_completions_url(raw_url: str) -> str:
    url = str(raw_url or "").strip().rstrip("/")
    if not url:
        return ""
    if url.endswith("/chat/completions"):
        return url
    if url.endswith("/v1") or url.endswith("/api/v3"):
        return f"{url}/chat/completions"
    return url


def _is_placeholder_api_key(raw_key: str) -> bool:
    token = str(raw_key or "").strip().lower()
    if not token:
        return True
    return any(marker in token for marker in _CHAT_DOUBAO_PLACEHOLDER_KEY_MARKERS)


def _split_model_candidates(raw_value: str) -> list[str]:
    tokens: list[str] = []
    for item in str(raw_value or "").replace(";", ",").split(","):
        value = str(item or "").strip()
        if value:
            tokens.append(value)
    return tokens


def _looks_like_non_chat_model(model_id: str) -> bool:
    token = str(model_id or "").strip().lower()
    if not token:
        return True
    blockers = (
        "embedding",
        "seedance",
        "seedream",
        "i2i",
        "t2i",
        "i2v",
        "t2v",
        "router",
        "rerank",
    )
    return any(marker in token for marker in blockers)


def _score_discovered_image_model(model_id: str, status: str) -> int:
    token = str(model_id or "").strip().lower()
    if not token:
        return -9999

    score = 0
    if "vision" in token:
        score += 100
    if token.startswith("doubao-seed-2-0-pro"):
        score += 96
    if token.startswith("doubao-seed-1-6-vision"):
        score += 92
    if token.startswith("doubao-1-5-vision-pro"):
        score += 88
    if token.startswith("doubao-1.5-vision"):
        score += 84
    if token.startswith("doubao-seed"):
        score += 20
    if token.startswith("doubao"):
        score += 12

    status_token = str(status or "").strip().lower()
    if "retiring" in status_token:
        score -= 20
    if "shutdown" in status_token:
        score -= 500

    if _looks_like_non_chat_model(token):
        score -= 200

    return score


def _discover_doubao_image_model_candidates(api_url: str, api_key: str, timeout_seconds: int) -> list[str]:
    if not _CHAT_DOUBAO_IMAGE_DISCOVER_MODELS:
        return []

    now = time.time()
    cache_ts = float(_CHAT_DOUBAO_IMAGE_DISCOVER_CACHE.get("ts") or 0.0)
    cached_models = _CHAT_DOUBAO_IMAGE_DISCOVER_CACHE.get("models")
    if isinstance(cached_models, list) and cached_models and (now - cache_ts) < float(_CHAT_DOUBAO_IMAGE_DISCOVER_TTL_SECONDS):
        return [str(x) for x in cached_models if str(x or "").strip()]

    model_list_url = str(api_url or "").strip()
    if model_list_url.endswith("/chat/completions"):
        model_list_url = model_list_url[: -len("/chat/completions")] + "/models"
    else:
        model_list_url = model_list_url.rstrip("/") + "/models"

    discovered: list[str] = []
    try:
        with httpx.Client(timeout=max(10, int(timeout_seconds)), trust_env=False) as client:
            resp = client.get(model_list_url, headers={"Authorization": f"Bearer {api_key}"})
        if int(resp.status_code or 0) != 200:
            raise RuntimeError(f"model_list_status={resp.status_code}")
        body = resp.json()
        rows = body.get("data") if isinstance(body, dict) else []
        if not isinstance(rows, list):
            rows = []

        scored: list[tuple[int, str]] = []
        for row in rows:
            if not isinstance(row, dict):
                continue
            model_id = str(row.get("id") or "").strip()
            status = str(row.get("status") or "").strip()
            if not model_id:
                continue
            score = _score_discovered_image_model(model_id, status)
            if score <= -80:
                continue
            scored.append((score, model_id))

        scored.sort(key=lambda x: (-x[0], x[1]))
        for _, model_id in scored:
            if model_id in discovered:
                continue
            discovered.append(model_id)
            if len(discovered) >= int(_CHAT_DOUBAO_IMAGE_DISCOVER_MAX_MODELS):
                break
    except Exception as exc:  # noqa: BLE001
        logger.debug("discover doubao models failed: %s", exc)

    _CHAT_DOUBAO_IMAGE_DISCOVER_CACHE["ts"] = now
    _CHAT_DOUBAO_IMAGE_DISCOVER_CACHE["models"] = discovered
    return discovered


def _resolve_doubao_image_runtime() -> tuple[str, str, list[str], int]:
    timeout_seconds = int(_CHAT_DOUBAO_IMAGE_TIMEOUT_SECONDS)
    if not _CHAT_DOUBAO_IMAGE_ENABLED:
        return "", "", [], timeout_seconds

    raw_url = str(os.getenv("DOUBAO_API_URL", getattr(cfg, "DOUBAO_API_URL", "")) or "").strip()
    api_url = _derive_chat_completions_url(raw_url)
    api_key = str(os.getenv("DOUBAO_API_KEY", getattr(cfg, "DOUBAO_API_KEY", "")) or "").strip()

    configured_candidates = _split_model_candidates(os.getenv("DOUBAO_IMAGE_MODEL_CANDIDATES", ""))
    model_candidates: list[str] = []
    for candidate in [
        str(_CHAT_DOUBAO_IMAGE_MODEL or "").strip(),
        str(getattr(cfg, "LLM_MODEL", "") or "").strip(),
        str(getattr(cfg, "DOUBAO_DEFAULT_MODEL", "") or "").strip(),
        *configured_candidates,
    ]:
        token = str(candidate or "").strip()
        if not token or token in model_candidates:
            continue
        model_candidates.append(token)

    if not api_url or not api_key:
        return "", "", [], timeout_seconds
    if _is_placeholder_api_key(api_key):
        return "", "", [], timeout_seconds

    discover_always = os.getenv("DOUBAO_IMAGE_DISCOVER_ALWAYS", "0") == "1"
    has_explicit_image_model = bool(str(_CHAT_DOUBAO_IMAGE_MODEL or "").strip()) or bool(configured_candidates)
    if discover_always or not has_explicit_image_model:
        for candidate in _discover_doubao_image_model_candidates(api_url, api_key, timeout_seconds):
            token = str(candidate or "").strip()
            if not token or token in model_candidates:
                continue
            model_candidates.append(token)

    if not model_candidates:
        return "", "", [], timeout_seconds

    max_models = max(1, int(_CHAT_DOUBAO_IMAGE_MAX_MODELS_PER_REQUEST))
    model_candidates = model_candidates[:max_models]
    return api_url, api_key, model_candidates, timeout_seconds

def _infer_image_mime_type(*, suffix: str, content_type: str, raw: bytes = b"") -> str:
    magic = bytes(raw[:16] if raw else b"")
    if magic.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if magic.startswith(b"\xff\xd8"):
        return "image/jpeg"
    if magic.startswith(b"GIF8"):
        return "image/gif"
    if len(magic) >= 12 and magic[:4] == b"RIFF" and magic[8:12] == b"WEBP":
        return "image/webp"
    if magic.startswith(b"BM"):
        return "image/bmp"

    normalized_type = str(content_type or "").strip().lower()
    if normalized_type.startswith("image/"):
        return normalized_type

    suffix_map = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
    }
    return suffix_map.get(str(suffix or "").lower(), "image/png")

def _extract_openai_message_text(payload: Any) -> str:
    if not isinstance(payload, dict):
        return ""

    choices = payload.get("choices")
    if not isinstance(choices, list) or not choices:
        return ""

    first_choice = choices[0] if isinstance(choices[0], dict) else {}
    message = first_choice.get("message") if isinstance(first_choice.get("message"), dict) else {}
    content = message.get("content")

    if isinstance(content, str):
        return content.strip()

    if isinstance(content, list):
        parts: list[str] = []
        for block in content:
            if isinstance(block, str) and block.strip():
                parts.append(block.strip())
                continue
            if not isinstance(block, dict):
                continue
            block_text = str(block.get("text") or "").strip()
            if block_text:
                parts.append(block_text)
        return "\n".join(parts).strip()

    return ""


def _extract_openai_error_info(payload: Any) -> tuple[str, str]:
    if not isinstance(payload, dict):
        return "", ""

    error_obj = payload.get("error")
    if isinstance(error_obj, dict):
        code = str(error_obj.get("code") or "").strip()
        message = str(error_obj.get("message") or "").strip()
        return code, message

    return "", ""


def _extract_image_doubao_preview(
    raw: bytes,
    *,
    suffix: str,
    content_type: str,
) -> tuple[str, str, str, str]:
    api_url, api_key, model_candidates, timeout_seconds = _resolve_doubao_image_runtime()
    if not api_url or not api_key or not model_candidates:
        return "", "partial", "", "config_missing"

    mime_type = _infer_image_mime_type(suffix=suffix, content_type=content_type, raw=raw)
    image_b64 = base64.b64encode(raw).decode("ascii")
    image_url = f"data:{mime_type};base64,{image_b64}"

    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {api_key}",
    }

    max_attempts_per_model = max(1, int(_CHAT_DOUBAO_IMAGE_RETRY_PER_MODEL))
    total_budget_seconds = max(
        float(timeout_seconds),
        float(_CHAT_DOUBAO_IMAGE_TOTAL_BUDGET_SECONDS),
    )
    deadline = time.monotonic() + total_budget_seconds
    last_reason = "unknown"

    for model in model_candidates:
        for attempt in range(1, max_attempts_per_model + 1):
            remaining_budget = deadline - time.monotonic()
            if remaining_budget <= 0:
                return "", "partial", "", "time_budget_exceeded"
            request_timeout = max(2.0, min(float(timeout_seconds), remaining_budget))
            payload = {
                "model": model,
                "temperature": 0.1,
                "max_tokens": int(_CHAT_DOUBAO_IMAGE_MAX_TOKENS),
                "messages": [
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": _CHAT_DOUBAO_IMAGE_PROMPT},
                            {"type": "image_url", "image_url": {"url": image_url}},
                        ],
                    }
                ],
            }

            try:
                with httpx.Client(timeout=request_timeout, trust_env=False) as client:
                    response = client.post(api_url, headers=headers, json=payload)
            except httpx.TimeoutException:
                last_reason = "network_timeout"
                if (deadline - time.monotonic()) <= 0:
                    return "", "partial", "", "time_budget_exceeded"
                if _CHAT_DOUBAO_IMAGE_FAIL_FAST_NETWORK:
                    return "", "partial", "", last_reason
                if attempt < max_attempts_per_model:
                    time.sleep(min(1.0, 0.25 * attempt))
                    continue
                break
            except Exception:
                last_reason = "network_error"
                if (deadline - time.monotonic()) <= 0:
                    return "", "partial", "", "time_budget_exceeded"
                if _CHAT_DOUBAO_IMAGE_FAIL_FAST_NETWORK:
                    return "", "partial", "", last_reason
                if attempt < max_attempts_per_model:
                    time.sleep(min(1.0, 0.25 * attempt))
                    continue
                break

            status_code = int(response.status_code or 0)
            body: Any = {}
            try:
                body = response.json()
            except Exception:
                body = {}

            if status_code >= 400:
                error_code, error_message = _extract_openai_error_info(body)
                code_lower = str(error_code or "").lower()
                message_lower = str(error_message or "").lower()

                if "invalidendpointormodel.notfound" in code_lower or "does not exist" in message_lower:
                    last_reason = "model_not_found"
                    break

                if "image dimensions are too small" in message_lower:
                    return "", "partial", "", "image_too_small"

                if "does not support" in message_lower and "image" in message_lower:
                    last_reason = "model_not_vision"
                    break

                if status_code >= 500 or status_code == 429:
                    last_reason = f"upstream_status_{status_code}"
                    if _CHAT_DOUBAO_IMAGE_FAIL_FAST_NETWORK:
                        return "", "partial", "", last_reason
                    if attempt < max_attempts_per_model:
                        time.sleep(min(1.0, 0.25 * attempt))
                        continue
                    break

                if status_code == 400 and ("invalidparameter" in code_lower or "unsupported" in message_lower):
                    last_reason = "invalid_image_request"
                    break

                last_reason = f"upstream_status_{status_code}"
                break

            extracted_text = _extract_openai_message_text(body)
            if not extracted_text:
                last_reason = "empty_response"
                if attempt < max_attempts_per_model:
                    time.sleep(min(1.0, 0.25 * attempt))
                    continue
                break

            cleaned_text = _clean_ocr_text(extracted_text)
            enhanced_text, table_rows, _layout_source = _append_layout_table_candidates_with_source(extracted_text)
            final_text = enhanced_text if table_rows > 0 else cleaned_text
            if not final_text:
                last_reason = "empty_response"
                if (deadline - time.monotonic()) <= 0:
                    return "", "partial", "", "time_budget_exceeded"
                if attempt < max_attempts_per_model:
                    time.sleep(min(1.0, 0.25 * attempt))
                    continue
                break

            if table_rows > 0:
                parser = f"doubao-vision-layout-{_layout_source}" if _layout_source else "doubao-vision-layout"
            else:
                parser = "doubao-vision"
            return final_text, "parsed", parser, ""

    if (deadline - time.monotonic()) <= 0:
        last_reason = "time_budget_exceeded"
    return "", "partial", "", last_reason

def _build_doubao_fallback_warning(reason: str) -> str:
    reason_key = str(reason or "").strip().lower()
    if reason_key == "time_budget_exceeded":
        return "【醒目提示】豆包读图耗时过长，已自动切换 OCR 兜底，识别质量可能下降。"
    if reason_key.startswith("network_"):
        return "【醒目提示】豆包读图暂时网络故障，已切换为 OCR 兜底，识别质量可能下降。"
    if reason_key.startswith("upstream_status_"):
        return "【醒目提示】豆包读图服务暂不可用，已切换为 OCR 兜底，识别质量可能下降。"
    if reason_key == "config_missing":
        return "【醒目提示】豆包读图未配置或凭证无效，已切换为 OCR 兜底。"
    if reason_key in {"model_not_found", "model_not_vision"}:
        return "【醒目提示】当前豆包模型不支持读图，已自动切换兼容链路并回退 OCR。"
    if reason_key == "image_too_small":
        return "【醒目提示】图片分辨率过低（最短边需≥14像素），已切换 OCR 兜底。"
    return "【醒目提示】豆包读图暂不可用，已切换为 OCR 兜底，识别质量可能下降。"


def _score_ocr_candidate(text: str) -> float:
    cleaned = _clean_ocr_text(text)
    if not cleaned:
        return 0.0

    lines = [line for line in cleaned.splitlines() if line.strip()]
    char_count = len(cleaned)
    cjk_count = len(re.findall(r"[一-鿿]", cleaned))
    latin_count = len(re.findall(r"[A-Za-z]", cleaned))
    digit_count = len(re.findall(r"\d", cleaned))

    unique_lines = len(set(lines))
    duplicate_penalty = max(0.0, float(len(lines) - unique_lines) * 0.8)

    layout_line_hits = 0
    ruled_hits = 0
    space_hits = 0
    layout_col_counts: list[int] = []

    for line in lines[:80]:
        text_line = str(line or "").strip()
        if not text_line:
            continue
        if _is_ruled_table_border_line(text_line):
            continue

        ruled_cells = _split_ruled_table_cells(text_line)
        if len(ruled_cells) >= 2:
            ruled_hits += 1
            layout_line_hits += 1
            layout_col_counts.append(len([cell for cell in ruled_cells if str(cell or "").strip()]))
            continue

        space_cells = _split_space_aligned_cells(text_line)
        if len(space_cells) >= 2:
            space_hits += 1
            layout_line_hits += 1
            layout_col_counts.append(len([cell for cell in space_cells if str(cell or "").strip()]))

    col_consistency = 0.0
    mode_col = 0
    mode_hits = 0
    if layout_col_counts:
        mode_col, mode_hits = Counter(layout_col_counts).most_common(1)[0]
        col_consistency = float(mode_hits) / float(max(1, len(layout_col_counts)))

    layout_bonus = 0.0
    if layout_line_hits >= 2:
        layout_bonus += min(10.0, float(layout_line_hits) * 2.2)
    if ruled_hits > 0:
        layout_bonus += min(6.0, float(ruled_hits) * 1.8)
    if space_hits > 0:
        layout_bonus += min(4.0, float(space_hits) * 1.0)
    if mode_col >= 2 and col_consistency >= 0.6 and layout_line_hits >= 2:
        layout_bonus += 4.0

    score = 0.0
    score += min(48.0, float(char_count) * 0.12)
    score += min(26.0, float(len(lines)) * 2.8)
    score += min(18.0, float(cjk_count) * 0.14 + float(latin_count) * 0.05 + float(digit_count) * 0.04)
    score += min(18.0, layout_bonus)
    score -= duplicate_penalty

    if len(lines) >= 2:
        score += 2.0
    return round(max(0.0, score), 3)


def _build_ocr_image_variants(image: Any) -> list[tuple[Any, str]]:
    variants: list[tuple[Any, str]] = []

    try:
        base = image.copy()
    except Exception:
        base = image

    try:
        if hasattr(base, "convert") and str(getattr(base, "mode", "")).upper() not in {"RGB", "L"}:
            base = base.convert("RGB")
    except Exception:
        pass

    variants.append((base, "orig"))

    # 复杂版面 OCR：增加灰度/对比度/二值化变体
    try:
        from PIL import ImageOps  # type: ignore

        gray = ImageOps.grayscale(base)
        variants.append((gray, "gray"))

        auto = ImageOps.autocontrast(gray)
        variants.append((auto, "autocontrast"))

        binary = auto.point(lambda p: 255 if p > 170 else 0)
        variants.append((binary, "binary"))
    except Exception:
        pass

    deduped: list[tuple[Any, str]] = []
    seen_labels: set[str] = set()
    for img, label in variants:
        if label in seen_labels:
            continue
        seen_labels.add(label)
        deduped.append((img, label))
        if len(deduped) >= _CHAT_OCR_MAX_VARIANTS:
            break
    return deduped


def _run_tesseract_pass(
    pytesseract_module: Any,
    image: Any,
    *,
    lang: str,
    config: str,
    timeout_seconds: int,
) -> str:
    timeout_value = max(1, int(timeout_seconds or 1))

    # Prefer full argument set first, then progressively downgrade for
    # simplified/mock pytesseract signatures while preserving config when possible.
    attempts = [
        {"lang": lang, "config": config, "timeout": float(timeout_value)},
        {"lang": lang, "config": config},
        {"config": config},
        {"lang": lang},
        {},
    ]

    last_type_error: TypeError | None = None
    for kwargs in attempts:
        try:
            return str(pytesseract_module.image_to_string(image, **kwargs) or "")
        except TypeError as exc:
            last_type_error = exc
            continue
        except RuntimeError as exc:
            if "timeout" in str(exc or "").lower():
                logger.warning(
                    "chat attachment OCR single-pass timeout (lang=%s, config=%s, timeout=%ss)",
                    lang,
                    config,
                    timeout_value,
                )
                return ""
            raise

    if last_type_error is not None:
        raise last_type_error
    return ""


def _extract_image_ocr_preview(raw: bytes) -> tuple[str, str, str]:
    # 可选 OCR：环境未安装时自动降级
    ocr_total_timed_out = False
    try:
        from PIL import Image  # type: ignore
        import pytesseract  # type: ignore

        image = Image.open(io.BytesIO(raw))
        variants = _build_ocr_image_variants(image)

        best_raw_text = ""
        best_score = 0.0
        best_variant = "orig"
        best_pass = ""

        started_at = time.perf_counter()
        total_timeout_seconds = max(6, int(_CHAT_OCR_TOTAL_TIMEOUT_SECONDS))
        single_pass_timeout_seconds = max(1, int(_CHAT_OCR_SINGLE_PASS_TIMEOUT_SECONDS))
        deadline = started_at + float(total_timeout_seconds)

        should_stop = False
        for variant_image, variant_label in variants:
            for lang, config in _CHAT_OCR_PASS_CONFIGS:
                if time.perf_counter() >= deadline:
                    ocr_total_timed_out = True
                    should_stop = True
                    break

                candidate_raw = _run_tesseract_pass(
                    pytesseract,
                    variant_image,
                    lang=lang,
                    config=config,
                    timeout_seconds=single_pass_timeout_seconds,
                )
                score = _score_ocr_candidate(candidate_raw)
                if score <= best_score:
                    continue

                best_raw_text = str(candidate_raw or "")
                best_score = score
                best_variant = variant_label
                best_pass = f"{lang}|{config}"

            if should_stop:
                break

        if ocr_total_timed_out:
            elapsed = max(0.0, time.perf_counter() - started_at)
            logger.warning(
                "chat attachment OCR total-timeout reached (timeout=%ss, elapsed=%.2fs)",
                total_timeout_seconds,
                elapsed,
            )

        if best_raw_text:
            cleaned_text = _clean_ocr_text(best_raw_text)
            enhanced_text, table_rows, _layout_source = _append_layout_table_candidates_with_source(best_raw_text)
            final_text = enhanced_text if table_rows > 0 else cleaned_text

            parser = "pytesseract"
            if best_variant != "orig" or "psm 6" not in best_pass:
                parser = "pytesseract-multi-pass"
            if table_rows > 0:
                parser = f"{parser}-layout-{_layout_source}" if _layout_source else f"{parser}-layout"
            return final_text, "parsed", parser
    except Exception:
        pass

    if ocr_total_timed_out:
        return "", "partial", "pytesseract-timeout"
    return "", "partial", ""


def _truncate_text(text: str, limit: int = _CHAT_ATTACHMENT_TEXT_CHAR_LIMIT) -> str:
    cleaned = str(text or "").replace("\x00", "").strip()
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[:limit] + "\n...[附件内容过长，已截断]"


def _extract_attachment_preview(*, filename: str, content_type: str, raw: bytes) -> dict:
    suffix = Path(filename).suffix.lower()
    notices: list[str] = []
    status = "parsed"
    text = ""
    parser = ""
    vision_engine = ""
    vision_fallback_reason = ""
    vision_warning = ""
    vision_is_degraded = False
    layout_source = ""

    if suffix and suffix not in _CHAT_ATTACHMENT_ALLOWED_SUFFIXES:
        status = "partial"
        notices.append("该格式暂不支持深度解析，已保留文件供后续引用。")

    if suffix in {".txt", ".md", ".json", ".yaml", ".yml", ".log", ".sql", ".xml", ".html", ".py", ".js", ".ts"}:
        text = _decode_text_bytes(raw)
        parser = "text-decoder"
    elif suffix == ".csv":
        text = _extract_csv_preview(raw)
        parser = "csv"
    elif suffix in {".xlsx", ".xls"}:
        text, status_xlsx, parser_xlsx = _extract_xlsx_preview(raw)
        if status_xlsx != "parsed":
            status = "partial"
            notices.append("Excel 解析能力受限，当前仅保留文件信息。")
        else:
            parser = parser_xlsx or "openpyxl"
    elif suffix == ".docx":
        text, status_docx, parser_docx = _extract_docx_preview(raw)
        if status_docx != "parsed":
            status = "partial"
            notices.append("DOCX 解析能力受限，建议补充关键段落文本。")
        else:
            parser = parser_docx or "python-docx"
    elif suffix == ".pdf":
        text, status_pdf, parser_pdf = _extract_pdf_preview(raw)
        if status_pdf != "parsed":
            status = "partial"
            notices.append("PDF 已上传，当前环境未启用全文提取；建议补充关键页文字。")
        else:
            parser = parser_pdf or "pypdf"
            if str(parser).startswith("pypdf-structured"):
                notices.append("PDF 已按章节/分页结构做预清洗，可直接用于跨页对齐分析。")
    elif suffix == ".pptx":
        text, status_pptx, parser_pptx = _extract_pptx_preview(raw)
        if status_pptx != "parsed":
            status = "partial"
            notices.append("PPTX 解析能力受限，建议转为 PDF 或补充关键页文本。")
        else:
            parser = parser_pptx or "pptx-xml-fallback"
            notices.append("PPTX 已提取关键页文本，可直接用于问题回答。")
    elif suffix == ".zip":
        text, status_zip, parser_zip = _extract_zip_preview(raw)
        if status_zip != "parsed":
            status = "partial"
            notices.append("ZIP 已提取文件清单，未识别到可解析文本；建议解压后上传关键文件。")
            parser = parser_zip or "zip-manifest"
        else:
            parser = parser_zip or "zip-text-bundle"
            notices.append("ZIP 已提取文件清单与可解析文本片段。")
            if parser == "zip-hybrid-bundle":
                notices.append("检测到 ZIP 内含 Office/嵌套压缩包，已自动提取其中可读内容。")
    elif suffix in _CHAT_IMAGE_SUFFIXES:
        text_doubao, status_doubao, parser_doubao, doubao_reason = _extract_image_doubao_preview(
            raw,
            suffix=suffix,
            content_type=content_type,
        )

        if status_doubao == "parsed":
            text = text_doubao
            parser = parser_doubao or "doubao-vision"
            vision_engine = "doubao"
            notices.append("图片已由豆包视觉模型解析。")
            parser_text = str(parser)
            if "layout" in parser_text:
                notices.append("图片中识别到结构化表格候选，已补充用于后续字段对齐。")
        else:
            text_ocr, status_ocr, parser_ocr = _extract_image_ocr_preview(raw)
            text = text_ocr
            vision_engine = "ocr-fallback"
            vision_fallback_reason = str(doubao_reason or "").strip().lower() or "unknown"
            vision_warning = _build_doubao_fallback_warning(vision_fallback_reason)
            vision_is_degraded = True
            notices.insert(0, vision_warning)
            if status_ocr != "parsed":
                status = "partial"
                if str(parser_ocr or "").strip().lower() == "pytesseract-timeout":
                    notices.append("图片已上传，但 OCR 解析超时；建议压缩/裁剪后重试，或稍后再试。")
                else:
                    notices.append("图片已上传，当前环境未启用 OCR；建议补充图片关键点描述。")
            else:
                parser = parser_ocr or "ocr"
                parser_text = str(parser)
                if "layout" in parser_text:
                    notices.append("图片中识别到结构化表格候选，已补充用于后续字段对齐。")
                if parser_text.startswith("pytesseract-multi-pass"):
                    notices.append("图片已执行多策略 OCR 预处理，可更稳健识别复杂版面文本。")
    elif content_type.startswith("text/"):

        text = _decode_text_bytes(raw)
        parser = "text-decoder"

    parser_layout_source = ""
    parser_match = re.search(r"-layout-(markdown|ruled|space)(?:$|[^a-z])", str(parser or "").lower())
    if parser_match:
        parser_layout_source = str(parser_match.group(1) or "").strip().lower()
        if parser_layout_source in {"markdown", "ruled", "space"}:
            layout_source = parser_layout_source

    if (not layout_source) and text and ("layout" in str(parser or "").lower() or "[结构化表格候选]" in text):
        layout_source = _infer_layout_source_from_text(text)

    text = _truncate_text(text)
    summary = ""
    if text:
        summary = _pick_preview_summary_line(text)
    if not summary:
        summary = f"{filename}（{len(raw)} bytes）"

    return {
        "status": status,
        "content": text,
        "summary": summary,
        "notices": notices,
        "parser": parser,
        "content_chars": len(text),
        "vision_engine": vision_engine,
        "vision_fallback_reason": vision_fallback_reason,
        "vision_warning": vision_warning,
        "vision_is_degraded": vision_is_degraded,
        "layout_source": layout_source,
    }


@router.post("/attachments/upload")
async def upload_chat_attachment(
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
):
    filename = _safe_filename(file.filename or "upload.bin")
    raw = await file.read()

    if not raw:
        raise HTTPException(status_code=400, detail="上传文件为空")

    max_bytes = max(1, int(UPLOAD_MAX_SIZE_MB)) * 1024 * 1024
    if len(raw) > max_bytes:
        raise HTTPException(status_code=413, detail=f"文件过大，请控制在 {UPLOAD_MAX_SIZE_MB}MB 以内")

    parse_result = _extract_attachment_preview(
        filename=filename,
        content_type=str(file.content_type or "").lower(),
        raw=raw,
    )

    token = uuid.uuid4().hex
    store_dir = _CHAT_UPLOAD_DIR / str(user.get("id") or "0")
    store_dir.mkdir(parents=True, exist_ok=True)
    stored_filename = f"{token[:12]}_{filename}"
    stored_path = store_dir / stored_filename
    try:
        stored_path.write_bytes(raw)
        stored = True
    except Exception:
        stored = False

    return {
        "attachment": {
            "id": token,
            "citation_tag": f"ATT-{token[:8].upper()}",
            "filename": filename,
            "file_type": str(file.content_type or "application/octet-stream"),
            "size": len(raw),
            "status": parse_result["status"],
            "summary": parse_result["summary"],
            "content": parse_result["content"],
            "notices": parse_result["notices"],
            "parser": parse_result.get("parser", ""),
            "content_chars": int(parse_result.get("content_chars") or len(parse_result["content"] or "")),
            "vision_engine": parse_result.get("vision_engine", ""),
            "vision_fallback_reason": parse_result.get("vision_fallback_reason", ""),
            "vision_warning": parse_result.get("vision_warning", ""),
            "vision_is_degraded": bool(parse_result.get("vision_is_degraded", False)),
            "layout_source": parse_result.get("layout_source", ""),
            "extracted_at": datetime.now(timezone.utc).isoformat(),
            "stored": stored,
        }
    }


@router.post("/voice/transcribe")
async def transcribe_chat_voice(
    file: UploadFile = File(...),
    language: str | None = Form(None),
    user: dict = Depends(get_current_user),
):
    filename = _safe_filename(file.filename or "voice.webm")
    suffix = Path(filename).suffix.lower()
    content_type = str(file.content_type or "").lower().strip()
    raw = await file.read()

    if not raw:
        raise HTTPException(status_code=400, detail="上传音频为空")

    max_bytes = max(1, int(UPLOAD_MAX_SIZE_MB)) * 1024 * 1024
    if len(raw) > max_bytes:
        raise HTTPException(status_code=413, detail=f"音频过大，请控制在 {UPLOAD_MAX_SIZE_MB}MB 以内")

    if not _is_allowed_voice_upload(suffix=suffix, content_type=content_type):
        raise HTTPException(status_code=415, detail="仅支持常见音频格式（webm/wav/mp3/m4a/ogg/aac/flac）")

    result = await _transcribe_audio_bytes(
        raw=raw,
        filename=filename,
        content_type=content_type,
        language=language,
    )
    logger.info(
        "Voice transcribe success user_id=%s size=%s model=%s",
        user.get("id"),
        len(raw),
        result.get("model", ""),
    )
    return {
        "transcript": result["text"],
        "model": result.get("model", ""),
        "provider_url": result.get("provider_url", ""),
        "content_type": content_type or "application/octet-stream",
        "size": len(raw),
        "language": str(language or "").strip().lower() or None,
    }


@router.post("/stream")
async def chat_stream_endpoint(
    req: ChatRequest,
    user: dict = Depends(get_current_user),
):
    """SSE流式聊天 — 推荐端点。

    事件类型: status / token / replace_reply / tool_call / tool_result / done / error
    后台事件（quality/trust/learning/alert）通过 GET /api/chat/events 获取。
    """
    normalized_message = _normalize_chat_message(req.message)
    if not normalized_message:
        raise HTTPException(status_code=400, detail="消息不能为空")

    conversation_id = req.conversation_id or str(uuid.uuid4())
    response_mode, is_valid_mode = _normalize_response_mode_input(req.response_mode)
    if not is_valid_mode:
        logger.warning(
            "Invalid response_mode on /api/chat/stream: %r (user_id=%s), fallback=%s",
            req.response_mode,
            user["id"],
            response_mode,
        )
        await _record_invalid_response_mode_metric(
            user_id=user["id"],
            endpoint="/api/chat/stream",
            raw_mode=req.response_mode,
            normalized_mode=response_mode,
        )

    merged_attachments = _merge_attachments_with_page_context(req.attachments or [], req.page_context)

    from src.core.chat_pipeline import chat_stream, PipelineEvent

    async def sse_generator():
        try:
            async for event in chat_stream(
                message=normalized_message,
                user_id=user["id"],
                role=req.role,
                role_lock=req.role_lock,
                conversation_id=conversation_id,
                product_id=req.product_id,
                product_ids=req.product_ids or [],
                workspace_id=req.workspace_id,
                response_mode=response_mode,
                learning_level=req.learning_level,
                collaboration_mode=req.collaboration_mode,
                hired_roles=req.hired_roles or [],
                attachments=merged_attachments,
            ):
                yield event.to_sse()
        except Exception as e:
            logger.error("SSE stream error: %s", e)
            err = PipelineEvent("error", {"message": str(e)})
            yield err.to_sse()

    return StreamingResponse(
        sse_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
            "X-Conversation-Id": conversation_id,
        },
    )


@router.post("", response_model=ChatResponse)
async def chat_endpoint(
    req: ChatRequest,
    user: dict = Depends(get_current_user),
):
    """非流式聊天 — 向后兼容。"""
    normalized_message = _normalize_chat_message(req.message)
    if not normalized_message:
        raise HTTPException(status_code=400, detail="消息不能为空")

    conversation_id = req.conversation_id or str(uuid.uuid4())
    response_mode, is_valid_mode = _normalize_response_mode_input(req.response_mode)
    if not is_valid_mode:
        logger.warning(
            "Invalid response_mode on /api/chat: %r (user_id=%s), fallback=%s",
            req.response_mode,
            user["id"],
            response_mode,
        )
        await _record_invalid_response_mode_metric(
            user_id=user["id"],
            endpoint="/api/chat",
            raw_mode=req.response_mode,
            normalized_mode=response_mode,
        )

    merged_attachments = _merge_attachments_with_page_context(req.attachments or [], req.page_context)

    from src.core.chat_pipeline import chat_simple
    result = await chat_simple(
        message=normalized_message,
        user_id=user["id"],
        role=req.role,
        role_lock=req.role_lock,
        conversation_id=conversation_id,
        product_id=req.product_id,
        product_ids=req.product_ids or [],
        workspace_id=req.workspace_id,
        response_mode=response_mode,
        learning_level=req.learning_level,
        collaboration_mode=req.collaboration_mode,
        hired_roles=req.hired_roles or [],
        attachments=merged_attachments,
    )

    return ChatResponse(
        reply=result["reply"],
        role=result["role"],
        metadata=result.get("metadata", {}),
    )


@router.get("/events")
async def get_background_events(
    user: dict = Depends(get_current_user),
    since_id: int = Query(0, description="返回ID大于此值的事件"),
    limit: int = Query(20, le=50),
):
    """轮询后台事件 — quality_check / trust_update / alert / learning。

    前端定期轮询此端点获取后台任务结果，不再阻塞SSE流。
    """
    db = await get_db()
    rows = await db.execute_fetchall(
        """
        SELECT id, event_type, data, created_at
        FROM background_events
        WHERE user_id = ? AND id > ? AND consumed = 0
        ORDER BY id ASC LIMIT ?
        """,
        (user["id"], since_id, limit),
    )

    events = []
    ids = []
    for row in rows:
        events.append({
            "id": row["id"],
            "event_type": row["event_type"],
            "data": json.loads(row["data"]) if isinstance(row["data"], str) else row["data"],
            "created_at": row["created_at"],
        })
        ids.append(row["id"])

    # 标记已消费
    if ids:
        placeholders = ",".join("?" * len(ids))
        await db.execute(
            f"UPDATE background_events SET consumed = 1 WHERE id IN ({placeholders})",
            ids,
        )
        await db.commit()

    return {"events": events}


# ═══════════════════════════════════════════════════════════════════════════
# 对话历史 API
# ═══════════════════════════════════════════════════════════════════════════

class ConversationRenameRequest(BaseModel):
    title: str = Field(default="", max_length=160)


class ConversationPinRequest(BaseModel):
    pinned: bool = Field(default=True)


def _normalize_conversation_title(raw_title: Any) -> str:
    text = str(raw_title or "").replace("\x00", "").strip()
    text = re.sub(r"\s+", " ", text)
    return text[:160]


async def _backfill_conversations_from_messages(db: Any, user_id: int) -> None:
    await db.execute(
        """
        INSERT INTO conversations (id, user_id, title, agent_role, created_at, updated_at)
        SELECT
            m.conversation_id AS id,
            m.user_id AS user_id,
            COALESCE(
                (SELECT substr(trim(m2.content), 1, 160)
                 FROM messages m2
                 WHERE m2.conversation_id = m.conversation_id
                   AND m2.user_id = m.user_id
                   AND m2.role = 'user'
                 ORDER BY m2.created_at ASC LIMIT 1),
                ''
            ) AS title,
            COALESCE(
                (SELECT CASE
                            WHEN json_valid(am.metadata) THEN json_extract(am.metadata, '$.role')
                            ELSE NULL
                        END
                 FROM messages am
                 WHERE am.conversation_id = m.conversation_id
                   AND am.user_id = m.user_id
                   AND am.role = 'assistant'
                 ORDER BY am.created_at DESC LIMIT 1),
                'ops'
            ) AS agent_role,
            MIN(m.created_at) AS created_at,
            MAX(m.created_at) AS updated_at
        FROM messages m
        WHERE m.user_id = ?
        GROUP BY m.conversation_id, m.user_id
        ON CONFLICT(id) DO UPDATE SET
            updated_at = CASE
                WHEN conversations.updated_at IS NULL OR conversations.updated_at < excluded.updated_at
                THEN excluded.updated_at
                ELSE conversations.updated_at
            END,
            agent_role = CASE
                WHEN trim(COALESCE(conversations.agent_role, '')) = '' AND trim(COALESCE(excluded.agent_role, '')) != ''
                THEN excluded.agent_role
                ELSE conversations.agent_role
            END,
            title = CASE
                WHEN trim(COALESCE(conversations.title, '')) = '' AND trim(COALESCE(excluded.title, '')) != ''
                THEN excluded.title
                ELSE conversations.title
            END
        WHERE conversations.user_id = excluded.user_id
        """,
        (user_id,),
    )
    await db.commit()


async def _ensure_conversation_row(db: Any, user_id: int, conversation_id: str) -> bool:
    row = await db.execute_fetchone(
        "SELECT id FROM conversations WHERE id = ? AND user_id = ?",
        (conversation_id, user_id),
    )
    if row:
        return True
    await _backfill_conversations_from_messages(db, user_id)
    row = await db.execute_fetchone(
        "SELECT id FROM conversations WHERE id = ? AND user_id = ?",
        (conversation_id, user_id),
    )
    return bool(row)


@conv_router.get("")
async def list_conversations(
    user: dict = Depends(get_current_user),
    role: str = Query("", description="按角色过滤"),
    limit: int = Query(50, le=200),
    offset: int = Query(0, ge=0),
):
    """获取当前用户的对话列表，置顶优先，再按最近更新时间倒序。"""
    db = await get_db()
    await _backfill_conversations_from_messages(db, user["id"])

    role_filter = str(role or "").strip().lower()
    role_sql = ""
    params: list[Any] = [user["id"]]
    if role_filter:
        role_sql = " AND COALESCE(NULLIF(TRIM(c.agent_role), ''), 'ops') = ?"
        params.append(role_filter)

    rows = await db.execute_fetchall(
        f"""
        SELECT
            c.id AS conversation_id,
            COALESCE(NULLIF(TRIM(c.title), ''),
                     (SELECT substr(trim(m2.content), 1, 60)
                      FROM messages m2
                      WHERE m2.conversation_id = c.id
                        AND m2.user_id = c.user_id
                        AND m2.role = 'user'
                      ORDER BY m2.created_at ASC LIMIT 1),
                     '（无标题）') AS title,
            COALESCE(NULLIF(TRIM(c.agent_role), ''),
                     (SELECT CASE
                                 WHEN json_valid(am.metadata) THEN json_extract(am.metadata, '$.role')
                                 ELSE NULL
                             END
                      FROM messages am
                      WHERE am.conversation_id = c.id
                        AND am.user_id = c.user_id
                        AND am.role = 'assistant'
                      ORDER BY am.created_at DESC LIMIT 1),
                     'ops') AS role,
            COALESCE(
                (SELECT MAX(mx.created_at)
                 FROM messages mx
                 WHERE mx.conversation_id = c.id
                   AND mx.user_id = c.user_id),
                c.updated_at,
                c.created_at
            ) AS last_at,
            (SELECT COUNT(*) FROM messages mc WHERE mc.conversation_id = c.id AND mc.user_id = c.user_id) AS msg_count,
            (SELECT CASE
                        WHEN json_valid(am2.metadata) THEN json_extract(am2.metadata, '$.response_mode')
                        ELSE NULL
                    END
             FROM messages am2
             WHERE am2.conversation_id = c.id
               AND am2.user_id = c.user_id
               AND am2.role = 'assistant'
             ORDER BY am2.created_at DESC LIMIT 1) AS response_mode,
            COALESCE(c.is_pinned, 0) AS is_pinned,
            c.pinned_at AS pinned_at
        FROM conversations c
        WHERE c.user_id = ?{role_sql}
        ORDER BY
            COALESCE(c.is_pinned, 0) DESC,
            CASE
                WHEN COALESCE(c.is_pinned, 0) = 1
                THEN COALESCE(c.pinned_at, c.updated_at, c.created_at)
                ELSE COALESCE(
                    (SELECT MAX(my.created_at)
                     FROM messages my
                     WHERE my.conversation_id = c.id
                       AND my.user_id = c.user_id),
                    c.updated_at,
                    c.created_at
                )
            END DESC
        LIMIT ? OFFSET ?
        """,
        (*params, limit, offset),
    )

    total_row = await db.execute_fetchone(
        f"""
        SELECT COUNT(*) AS cnt
        FROM conversations c
        WHERE c.user_id = ?{role_sql}
        """,
        tuple(params),
    )
    total = int(total_row["cnt"] or 0) if total_row else len(rows)

    return {
        "conversations": [
            {
                "id": r["conversation_id"],
                "role": r["role"] or "ops",
                "last_at": r["last_at"],
                "msg_count": r["msg_count"],
                "title": r["title"] or "（无标题）",
                "response_mode": _sanitize_response_mode(r["response_mode"]),
                "is_pinned": bool(int(r["is_pinned"] or 0)),
                "pinned_at": r["pinned_at"],
            }
            for r in rows
        ],
        "total": total,
    }


@conv_router.get("/{conversation_id}/messages")
async def get_conversation_messages(
    conversation_id: str,
    user: dict = Depends(get_current_user),
):
    """获取指定对话的全部消息。"""
    db = await get_db()
    rows = await db.execute_fetchall(
        """
        SELECT id, role, content, metadata, created_at
        FROM messages
        WHERE conversation_id = ? AND user_id = ?
        ORDER BY created_at ASC
        """,
        (conversation_id, user["id"]),
    )
    return {
        "conversation_id": conversation_id,
        "messages": [
            {
                "id": r["id"],
                "role": r["role"],
                "content": r["content"],
                "metadata": json.loads(r["metadata"]) if r["metadata"] else {},
                "created_at": r["created_at"],
            }
            for r in rows
        ],
    }


@conv_router.get("/{conversation_id}/export")
async def export_conversation(
    conversation_id: str,
    format: str = Query("raw", description="raw=原文markdown / report=AI总结报告"),
    user: dict = Depends(get_current_user),
):
    """导出对话：raw=原文Markdown，report=AI总结报告（触发LLM汇总）。"""
    db = await get_db()
    rows = await db.execute_fetchall(
        """
        SELECT role, content, metadata, created_at
        FROM messages
        WHERE conversation_id = ? AND user_id = ?
        ORDER BY created_at ASC
        """,
        (conversation_id, user["id"]),
    )
    if not rows:
        raise HTTPException(status_code=404, detail="对话不存在")

    messages = [
        {
            "role": r["role"],
            "content": r["content"] or "",
            "metadata": json.loads(r["metadata"]) if r["metadata"] else {},
            "created_at": r["created_at"],
        }
        for r in rows
    ]

    if format == "raw":
        # ── 原文导出：逐条拼接 Markdown ──
        lines = [f"# 对话记录 #{conversation_id}\n"]
        for m in messages:
            if m["role"] not in ("user", "assistant"):
                continue
            role_label = "**用户**" if m["role"] == "user" else "**AI助手**"
            ts = m["created_at"] or ""
            lines.append(f"### {role_label}  {ts}\n")
            lines.append(m["content"])
            lines.append("\n---\n")
        content_md = "\n".join(lines)
        from fastapi.responses import PlainTextResponse
        return PlainTextResponse(
            content=content_md,
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="conv_{conversation_id}.md"'},
        )

    elif format == "report":
        # ── 报告导出：调用 LLM 生成结构化总结 ──
        dialogue_text = ""
        for m in messages:
            if m["role"] not in ("user", "assistant"):
                continue
            label = "用户" if m["role"] == "user" else "AI助手"
            dialogue_text += f"{label}：{m['content'][:800]}\n\n"

        try:
            from src.llm_client import call_llm
            system_prompt = (
                "你是专业的电商运营总结助手。"
                "请将用户提供的对话内容整理成结构化业务报告，输出Markdown格式。"
                "报告包含：一、对话概述（1-3句）；二、核心问题/需求；三、AI建议与分析要点（分条列出）；"
                "四、关键数据/指标（如有）；五、行动建议（可直接执行的步骤）。"
                "语言简洁专业，中文。"
            )
            user_prompt = f"以下是一段AI电商助手对话，请整理成报告：\n\n{dialogue_text[:4000]}"
            report_md = await call_llm(
                messages=[{"role": "user", "content": user_prompt}],
                system=system_prompt,
                max_tokens=1200,
                temperature=0.5,
            ) or ""
        except Exception as e:
            report_md = f"# 对话报告生成失败\n\n错误：{e}\n\n## 原始对话摘要\n\n{dialogue_text[:2000]}"

        from fastapi.responses import PlainTextResponse
        return PlainTextResponse(
            content=report_md or "报告生成失败",
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="report_{conversation_id}.md"'},
        )

    raise HTTPException(status_code=400, detail="format 参数须为 raw 或 report")


@conv_router.patch("/{conversation_id}/title")
async def rename_conversation(
    conversation_id: str,
    req: ConversationRenameRequest,
    user: dict = Depends(get_current_user),
):
    # 重命名对话标题（持久化到 conversations.title）
    db = await get_db()
    if not await _ensure_conversation_row(db, user["id"], conversation_id):
        raise HTTPException(status_code=404, detail="对话不存在")

    title = _normalize_conversation_title(req.title)
    if not title:
        raise HTTPException(status_code=400, detail="标题不能为空")

    await db.execute(
        """
        UPDATE conversations
        SET title = ?, updated_at = CURRENT_TIMESTAMP
        WHERE id = ? AND user_id = ?
        """,
        (title, conversation_id, user["id"]),
    )
    await db.commit()
    return {"ok": True, "conversation_id": conversation_id, "title": title}


@conv_router.post("/{conversation_id}/title")
async def rename_conversation_compat(
    conversation_id: str,
    req: ConversationRenameRequest,
    user: dict = Depends(get_current_user),
):
    # 兼容部分代理/网关不放行 PATCH 的场景
    return await rename_conversation(conversation_id, req, user)


@conv_router.post("/{conversation_id}/pin")
async def pin_conversation(
    conversation_id: str,
    req: ConversationPinRequest,
    user: dict = Depends(get_current_user),
):
    # 设置/取消对话置顶
    db = await get_db()
    if not await _ensure_conversation_row(db, user["id"], conversation_id):
        raise HTTPException(status_code=404, detail="对话不存在")

    pinned = bool(req.pinned)
    if pinned:
        await db.execute(
            """
            UPDATE conversations
            SET is_pinned = 1,
                pinned_at = CURRENT_TIMESTAMP,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND user_id = ?
            """,
            (conversation_id, user["id"]),
        )
    else:
        await db.execute(
            """
            UPDATE conversations
            SET is_pinned = 0,
                pinned_at = NULL,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND user_id = ?
            """,
            (conversation_id, user["id"]),
        )
    await db.commit()
    return {"ok": True, "conversation_id": conversation_id, "is_pinned": pinned}


@conv_router.delete("/{conversation_id}")
async def delete_conversation(
    conversation_id: str,
    user: dict = Depends(get_current_user),
):
    # 删除指定对话（消息 + 会话元数据）
    db = await get_db()
    await db.execute(
        "DELETE FROM messages WHERE conversation_id = ? AND user_id = ?",
        (conversation_id, user["id"]),
    )
    await db.execute(
        "DELETE FROM conversations WHERE id = ? AND user_id = ?",
        (conversation_id, user["id"]),
    )
    await db.commit()
    return {"ok": True}


@conv_router.post("/{conversation_id}/teaching-submit")
async def submit_conversation_for_teaching_review(
    conversation_id: str,
    req: TeachingSubmissionRequest,
    user: dict = Depends(get_current_user),
):
    """学生提交会话给指定教师评阅。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "student":
        raise HTTPException(status_code=403, detail="仅学生账号可提交教学评阅")

    teacher_id = int(req.teacher_user_id)
    if teacher_id == int(user["id"]):
        raise HTTPException(status_code=400, detail="不能将会话提交给自己")

    exists = await db.execute_fetchone(
        "SELECT 1 FROM messages WHERE conversation_id = ? AND user_id = ? LIMIT 1",
        (conversation_id, user["id"]),
    )
    if not exists:
        raise HTTPException(status_code=404, detail="会话不存在或不属于当前学生")

    teacher_row = await db.execute_fetchone(
        "SELECT id, name, account_role FROM users WHERE id = ?",
        (teacher_id,),
    )
    if not teacher_row:
        raise HTTPException(status_code=404, detail="教师用户不存在")

    teacher_role = _normalize_account_role(
        teacher_row["account_role"] if "account_role" in teacher_row.keys() else "general"
    )
    if teacher_role != "teacher":
        raise HTTPException(status_code=400, detail="目标用户不是教师账号")

    note = str(req.note or "").strip()[:500]
    await db.execute(
        """
        INSERT INTO teaching_submissions
            (conversation_id, student_user_id, teacher_user_id, status, note, submitted_at, updated_at)
        VALUES (?, ?, ?, 'submitted', ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        ON CONFLICT(conversation_id, student_user_id, teacher_user_id)
        DO UPDATE SET
            status = 'submitted',
            note = excluded.note,
            submitted_at = CURRENT_TIMESTAMP,
            updated_at = CURRENT_TIMESTAMP
        """,
        (conversation_id, user["id"], teacher_id, note),
    )
    await db.commit()

    row = await db.execute_fetchone(
        """
        SELECT id, status, note, submitted_at, updated_at
        FROM teaching_submissions
        WHERE conversation_id = ? AND student_user_id = ? AND teacher_user_id = ?
        """,
        (conversation_id, user["id"], teacher_id),
    )
    if not row:
        raise HTTPException(status_code=500, detail="提交记录写入失败")

    return {
        "id": row["id"],
        "conversation_id": conversation_id,
        "student_user_id": user["id"],
        "teacher_user_id": teacher_id,
        "teacher_name": teacher_row["name"] or "",
        "status": row["status"],
        "note": row["note"] or "",
        "submitted_at": row["submitted_at"],
        "updated_at": row["updated_at"],
    }


@conv_router.get("/teaching/teachers")
async def list_teaching_teachers(
    user: dict = Depends(get_current_user),
    q: str = Query("", description="按姓名/邮箱模糊搜索"),
    limit: int = Query(50, ge=1, le=200),
):
    """教学评阅教师列表（供学生提交时选择）。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role not in {"student", "teacher"}:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看教师列表")

    keyword = f"%{str(q or '').strip().lower()}%"
    params: list[Any] = [int(user["id"])]
    sql = """
        SELECT id, name, email
        FROM users
        WHERE id != ?
          AND LOWER(COALESCE(account_role, 'general')) = 'teacher'
    """
    if str(q or "").strip():
        sql += " AND (LOWER(COALESCE(name, '')) LIKE ? OR LOWER(COALESCE(email, '')) LIKE ?)"
        params.extend([keyword, keyword])

    sql += " ORDER BY id DESC LIMIT ?"
    params.append(int(limit))

    rows = await db.execute_fetchall(sql, tuple(params))
    return {
        "teachers": [
            {
                "id": r["id"],
                "name": r["name"] or r["email"] or f"教师#{r['id']}",
                "email": r["email"] or "",
            }
            for r in rows
        ]
    }


@conv_router.get("/teaching/submissions")
async def list_teaching_submissions_for_teacher(
    user: dict = Depends(get_current_user),
    status: str = Query("", description="按评阅状态过滤"),
    student_user_id: int = Query(0, ge=0),
    limit: int = Query(50, le=200),
    offset: int = Query(0, ge=0),
):
    """教师查看分配给自己的学生提交列表。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可查看评阅列表")

    status_token = str(status or "").strip().lower()
    sql = """
        SELECT
            ts.id,
            ts.conversation_id,
            ts.student_user_id,
            su.name AS student_name,
            ts.teacher_user_id,
            tu.name AS teacher_name,
            ts.status,
            ts.note,
            ts.submitted_at,
            ts.updated_at,
            (SELECT content FROM messages m2
             WHERE m2.conversation_id = ts.conversation_id
               AND m2.user_id = ts.student_user_id
               AND m2.role = 'user'
             ORDER BY m2.created_at ASC LIMIT 1) AS first_user_msg,
            (SELECT MAX(created_at) FROM messages m3
             WHERE m3.conversation_id = ts.conversation_id
               AND m3.user_id = ts.student_user_id) AS last_at,
            (SELECT COUNT(*) FROM teaching_evaluations te
             WHERE te.submission_id = ts.id) AS evaluation_count
        FROM teaching_submissions ts
        JOIN users su ON su.id = ts.student_user_id
        JOIN users tu ON tu.id = ts.teacher_user_id
        WHERE ts.teacher_user_id = ?
    """
    params: list[Any] = [user["id"]]

    if status_token:
        sql += " AND ts.status = ?"
        params.append(status_token)
    if int(student_user_id or 0) > 0:
        sql += " AND ts.student_user_id = ?"
        params.append(int(student_user_id))

    sql += " ORDER BY ts.submitted_at DESC LIMIT ? OFFSET ?"
    params.extend([limit, offset])

    rows = await db.execute_fetchall(sql, tuple(params))

    count_sql = "SELECT COUNT(*) AS cnt FROM teaching_submissions WHERE teacher_user_id = ?"
    count_params: list[Any] = [user["id"]]
    if status_token:
        count_sql += " AND status = ?"
        count_params.append(status_token)
    if int(student_user_id or 0) > 0:
        count_sql += " AND student_user_id = ?"
        count_params.append(int(student_user_id))
    total_row = await db.execute_fetchone(count_sql, tuple(count_params))

    return {
        "submissions": [
            {
                "id": r["id"],
                "conversation_id": r["conversation_id"],
                "student_user_id": r["student_user_id"],
                "student_name": r["student_name"] or "",
                "teacher_user_id": r["teacher_user_id"],
                "teacher_name": r["teacher_name"] or "",
                "status": r["status"],
                "note": r["note"] or "",
                "title": (r["first_user_msg"] or "")[:60] or "（无标题）",
                "last_at": r["last_at"],
                "evaluation_count": int(r["evaluation_count"] or 0),
                "submitted_at": r["submitted_at"],
                "updated_at": r["updated_at"],
            }
            for r in rows
        ],
        "total": int(total_row["cnt"] or 0) if total_row else len(rows),
    }


@conv_router.get("/teaching/my-submissions")
async def list_my_teaching_submissions(
    user: dict = Depends(get_current_user),
    status: str = Query("", description="按评阅状态过滤"),
    limit: int = Query(50, le=200),
    offset: int = Query(0, ge=0),
):
    """学生查看自己提交的评阅记录。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "student":
        raise HTTPException(status_code=403, detail="仅学生账号可查看我的评阅提交")

    status_token = str(status or "").strip().lower()
    sql = """
        SELECT
            ts.id,
            ts.conversation_id,
            ts.student_user_id,
            ts.teacher_user_id,
            tu.name AS teacher_name,
            ts.status,
            ts.note,
            ts.submitted_at,
            ts.updated_at,
            (SELECT content FROM messages m2
             WHERE m2.conversation_id = ts.conversation_id
               AND m2.user_id = ts.student_user_id
               AND m2.role = 'user'
             ORDER BY m2.created_at ASC LIMIT 1) AS first_user_msg,
            (SELECT MAX(created_at) FROM messages m3
             WHERE m3.conversation_id = ts.conversation_id
               AND m3.user_id = ts.student_user_id) AS last_at,
            (SELECT COUNT(*) FROM teaching_evaluations te
             WHERE te.submission_id = ts.id) AS evaluation_count
        FROM teaching_submissions ts
        JOIN users tu ON tu.id = ts.teacher_user_id
        WHERE ts.student_user_id = ?
    """
    params: list[Any] = [user["id"]]

    if status_token:
        sql += " AND ts.status = ?"
        params.append(status_token)

    sql += " ORDER BY ts.submitted_at DESC LIMIT ? OFFSET ?"
    params.extend([limit, offset])

    rows = await db.execute_fetchall(sql, tuple(params))

    count_sql = "SELECT COUNT(*) AS cnt FROM teaching_submissions WHERE student_user_id = ?"
    count_params: list[Any] = [user["id"]]
    if status_token:
        count_sql += " AND status = ?"
        count_params.append(status_token)
    total_row = await db.execute_fetchone(count_sql, tuple(count_params))

    return {
        "submissions": [
            {
                "id": r["id"],
                "conversation_id": r["conversation_id"],
                "student_user_id": r["student_user_id"],
                "teacher_user_id": r["teacher_user_id"],
                "teacher_name": r["teacher_name"] or "",
                "status": r["status"],
                "note": r["note"] or "",
                "title": (r["first_user_msg"] or "")[:60] or "（无标题）",
                "last_at": r["last_at"],
                "evaluation_count": int(r["evaluation_count"] or 0),
                "submitted_at": r["submitted_at"],
                "updated_at": r["updated_at"],
            }
            for r in rows
        ],
        "total": int(total_row["cnt"] or 0) if total_row else len(rows),
    }


@conv_router.get("/teaching/submissions/{submission_id}/messages")
async def get_teaching_submission_messages(
    submission_id: int,
    user: dict = Depends(get_current_user),
):
    """教师或对应学生查看提交会话内容。"""
    db = await get_db()
    submission = await db.execute_fetchone(
        """
        SELECT id, conversation_id, student_user_id, teacher_user_id, status, note, submitted_at, updated_at
        FROM teaching_submissions
        WHERE id = ?
        """,
        (submission_id,),
    )
    if not submission:
        raise HTTPException(status_code=404, detail="提交记录不存在")

    account_role = await _resolve_requester_account_role(db, user)
    if not _can_access_teaching_submission(
        requester_id=int(user["id"]),
        requester_role=account_role,
        submission_row=submission,
    ):
        raise HTTPException(status_code=403, detail="无权查看该评阅提交")

    rows = await db.execute_fetchall(
        """
        SELECT id, role, content, metadata, created_at
        FROM messages
        WHERE conversation_id = ? AND user_id = ?
        ORDER BY created_at ASC
        """,
        (submission["conversation_id"], submission["student_user_id"]),
    )

    return {
        "submission": {
            "id": submission["id"],
            "conversation_id": submission["conversation_id"],
            "student_user_id": submission["student_user_id"],
            "teacher_user_id": submission["teacher_user_id"],
            "status": submission["status"],
            "note": submission["note"] or "",
            "submitted_at": submission["submitted_at"],
            "updated_at": submission["updated_at"],
        },
        "messages": [
            {
                "id": r["id"],
                "role": r["role"],
                "content": r["content"],
                "metadata": json.loads(r["metadata"]) if r["metadata"] else {},
                "created_at": r["created_at"],
            }
            for r in rows
        ],
    }


@conv_router.post("/teaching/submissions/{submission_id}/evaluations")
async def create_teaching_evaluation(
    submission_id: int,
    req: TeachingEvaluationRequest,
    user: dict = Depends(get_current_user),
):
    """教师创建评阅记录。"""
    db = await get_db()
    submission = await db.execute_fetchone(
        """
        SELECT id, conversation_id, student_user_id, teacher_user_id, status
        FROM teaching_submissions
        WHERE id = ?
        """,
        (submission_id,),
    )
    if not submission:
        raise HTTPException(status_code=404, detail="提交记录不存在")

    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher" or int(user["id"]) != int(submission["teacher_user_id"]):
        raise HTTPException(status_code=403, detail="仅被指派教师可创建评阅")

    target_message_id = int(req.message_id or 0)
    if target_message_id > 0:
        target_msg = await db.execute_fetchone(
            """
            SELECT id FROM messages
            WHERE id = ? AND conversation_id = ? AND user_id = ?
            """,
            (target_message_id, submission["conversation_id"], submission["student_user_id"]),
        )
        if not target_msg:
            raise HTTPException(status_code=400, detail="message_id 不属于该学生会话")

    score = float(req.score)
    feedback = str(req.feedback or "").strip()[:2000]
    rubric_json = json.dumps(req.rubric or {}, ensure_ascii=False)

    cur = await db.execute(
        """
        INSERT INTO teaching_evaluations
            (submission_id, conversation_id, teacher_user_id, student_user_id, message_id, score, feedback, rubric_json)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            submission_id,
            submission["conversation_id"],
            submission["teacher_user_id"],
            submission["student_user_id"],
            target_message_id if target_message_id > 0 else None,
            score,
            feedback,
            rubric_json,
        ),
    )
    await db.execute(
        "UPDATE teaching_submissions SET status = 'reviewed', updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        (submission_id,),
    )
    await db.commit()

    return {
        "id": int(cur.lastrowid or 0),
        "submission_id": submission_id,
        "conversation_id": submission["conversation_id"],
        "teacher_user_id": submission["teacher_user_id"],
        "student_user_id": submission["student_user_id"],
        "message_id": target_message_id if target_message_id > 0 else None,
        "score": score,
        "feedback": feedback,
        "rubric": req.rubric or {},
        "status": "reviewed",
    }


@conv_router.get("/teaching/submissions/{submission_id}/evaluations")
async def list_teaching_evaluations(
    submission_id: int,
    user: dict = Depends(get_current_user),
):
    """查看评阅记录（教师或对应学生）。"""
    db = await get_db()
    submission = await db.execute_fetchone(
        """
        SELECT id, conversation_id, student_user_id, teacher_user_id, status
        FROM teaching_submissions
        WHERE id = ?
        """,
        (submission_id,),
    )
    if not submission:
        raise HTTPException(status_code=404, detail="提交记录不存在")

    account_role = await _resolve_requester_account_role(db, user)
    if not _can_access_teaching_submission(
        requester_id=int(user["id"]),
        requester_role=account_role,
        submission_row=submission,
    ):
        raise HTTPException(status_code=403, detail="无权查看该评阅记录")

    rows = await db.execute_fetchall(
        """
        SELECT te.id, te.submission_id, te.conversation_id, te.teacher_user_id, te.student_user_id,
               te.message_id, te.score, te.feedback, te.rubric_json, te.created_at,
               u.name AS teacher_name
        FROM teaching_evaluations te
        LEFT JOIN users u ON u.id = te.teacher_user_id
        WHERE te.submission_id = ?
        ORDER BY te.created_at DESC
        """,
        (submission_id,),
    )

    evaluations = []
    for r in rows:
        rubric = {}
        try:
            if r["rubric_json"]:
                rubric = json.loads(r["rubric_json"])
        except Exception:
            rubric = {}
        evaluations.append(
            {
                "id": r["id"],
                "submission_id": r["submission_id"],
                "conversation_id": r["conversation_id"],
                "teacher_user_id": r["teacher_user_id"],
                "teacher_name": r["teacher_name"] or "",
                "student_user_id": r["student_user_id"],
                "message_id": r["message_id"],
                "score": float(r["score"] or 0),
                "feedback": r["feedback"] or "",
                "rubric": rubric,
                "created_at": r["created_at"],
            }
        )

    return {
        "submission_id": submission_id,
        "status": submission["status"],
        "evaluations": evaluations,
    }

@conv_router.post("/teaching/classes")
async def create_teaching_class(
    req: TeachingClassCreateRequest,
    user: dict = Depends(get_current_user),
):
    """教师创建班级。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可创建班级")

    name = str(req.name or "").strip()[:120]
    if not name:
        raise HTTPException(status_code=400, detail="班级名称不能为空")
    description = str(req.description or "").strip()[:1000]

    await db.execute(
        """
        INSERT INTO teaching_classes
            (teacher_user_id, name, description, status, created_at, updated_at)
        VALUES (?, ?, ?, 'active', CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        ON CONFLICT(teacher_user_id, name)
        DO UPDATE SET
            description = excluded.description,
            status = 'active',
            updated_at = CURRENT_TIMESTAMP
        """,
        (user["id"], name, description),
    )

    class_row = await db.execute_fetchone(
        """
        SELECT id, teacher_user_id, name, description, status, created_at, updated_at
        FROM teaching_classes
        WHERE teacher_user_id = ? AND name = ?
        """,
        (user["id"], name),
    )
    if not class_row:
        raise HTTPException(status_code=500, detail="班级创建失败")

    return {
        "id": class_row["id"],
        "teacher_user_id": class_row["teacher_user_id"],
        "name": class_row["name"],
        "description": class_row["description"] or "",
        "status": class_row["status"],
        "created_at": class_row["created_at"],
        "updated_at": class_row["updated_at"],
    }


@conv_router.get("/teaching/classes")
async def list_teaching_classes(
    user: dict = Depends(get_current_user),
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
):
    """教学班级列表：教师看自己创建的班级，学生看自己加入的班级。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)

    rows = []
    total = 0
    if account_role == "teacher":
        rows = await db.execute_fetchall(
            """
            SELECT
                tc.id,
                tc.teacher_user_id,
                tc.name,
                tc.description,
                tc.status,
                tc.created_at,
                tc.updated_at,
                (SELECT COUNT(*) FROM teaching_class_members tcm WHERE tcm.class_id = tc.id) AS member_count,
                (SELECT COUNT(*) FROM teaching_assignments ta WHERE ta.class_id = tc.id) AS assignment_count
            FROM teaching_classes tc
            WHERE tc.teacher_user_id = ?
            ORDER BY tc.created_at DESC
            LIMIT ? OFFSET ?
            """,
            (user["id"], limit, offset),
        )
        total_row = await db.execute_fetchone(
            "SELECT COUNT(*) AS cnt FROM teaching_classes WHERE teacher_user_id = ?",
            (user["id"],),
        )
        total = int(total_row["cnt"] or 0) if total_row else len(rows)
    elif account_role == "student":
        rows = await db.execute_fetchall(
            """
            SELECT
                tc.id,
                tc.teacher_user_id,
                tc.name,
                tc.description,
                tc.status,
                tc.created_at,
                tc.updated_at,
                tcm.joined_at,
                (SELECT COUNT(*) FROM teaching_class_members x WHERE x.class_id = tc.id) AS member_count,
                (SELECT COUNT(*) FROM teaching_assignments ta WHERE ta.class_id = tc.id) AS assignment_count
            FROM teaching_classes tc
            JOIN teaching_class_members tcm ON tcm.class_id = tc.id
            WHERE tcm.student_user_id = ?
            ORDER BY tc.created_at DESC
            LIMIT ? OFFSET ?
            """,
            (user["id"], limit, offset),
        )
        total_row = await db.execute_fetchone(
            "SELECT COUNT(*) AS cnt FROM teaching_class_members WHERE student_user_id = ?",
            (user["id"],),
        )
        total = int(total_row["cnt"] or 0) if total_row else len(rows)
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看教学班级")

    return {
        "classes": [
            {
                "id": r["id"],
                "teacher_user_id": r["teacher_user_id"],
                "name": r["name"] or "",
                "description": r["description"] or "",
                "status": r["status"],
                "member_count": int(r["member_count"] or 0),
                "assignment_count": int(r["assignment_count"] or 0),
                "joined_at": r["joined_at"] if "joined_at" in r.keys() else None,
                "created_at": r["created_at"],
                "updated_at": r["updated_at"],
            }
            for r in rows
        ],
        "total": total,
    }


@conv_router.post("/teaching/classes/{class_id}/members")
async def add_teaching_class_member(
    class_id: int,
    req: TeachingClassMemberAddRequest,
    user: dict = Depends(get_current_user),
):
    """教师向班级添加学生成员。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可管理班级成员")

    class_row = await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    student_id = int(req.student_user_id)
    if student_id == int(user["id"]):
        raise HTTPException(status_code=400, detail="教师不能将自己作为学生成员")

    student_row = await db.execute_fetchone(
        "SELECT id, name, email, account_role FROM users WHERE id = ?",
        (student_id,),
    )
    if not student_row:
        raise HTTPException(status_code=404, detail="学生账号不存在")

    student_role = _normalize_account_role(
        student_row["account_role"] if "account_role" in student_row.keys() else "general"
    )
    if student_role != "student":
        raise HTTPException(status_code=400, detail="目标用户不是学生账号")

    await db.execute(
        """
        INSERT OR IGNORE INTO teaching_class_members (class_id, student_user_id, joined_at)
        VALUES (?, ?, CURRENT_TIMESTAMP)
        """,
        (class_id, student_id),
    )
    await db.commit()

    member_row = await db.execute_fetchone(
        """
        SELECT tcm.class_id, tcm.student_user_id, tcm.joined_at, u.name AS student_name, u.email AS student_email
        FROM teaching_class_members tcm
        LEFT JOIN users u ON u.id = tcm.student_user_id
        WHERE tcm.class_id = ? AND tcm.student_user_id = ?
        """,
        (class_id, student_id),
    )
    if not member_row:
        raise HTTPException(status_code=500, detail="班级成员写入失败")

    return {
        "class_id": class_id,
        "class_name": class_row["name"] or "",
        "student_user_id": member_row["student_user_id"],
        "student_name": member_row["student_name"] or member_row["student_email"] or "",
        "student_email": member_row["student_email"] or "",
        "joined_at": member_row["joined_at"],
    }


@conv_router.get("/teaching/classes/{class_id}/members")
async def list_teaching_class_members(
    class_id: int,
    user: dict = Depends(get_current_user),
):
    """班级成员列表（班级教师或该班学生可查看）。"""
    db = await get_db()
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")

    account_role = await _resolve_requester_account_role(db, user)
    requester_id = int(user["id"])
    if account_role == "teacher":
        if int(class_row["teacher_user_id"] or 0) != requester_id:
            raise HTTPException(status_code=403, detail="无权查看该班级成员")
    elif account_role == "student":
        in_class = await _is_teaching_class_member(db, class_id=class_id, student_user_id=requester_id)
        if not in_class:
            raise HTTPException(status_code=403, detail="无权查看该班级成员")
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看班级成员")

    rows = await db.execute_fetchall(
        """
        SELECT tcm.class_id, tcm.student_user_id, tcm.joined_at, u.name AS student_name, u.email AS student_email
        FROM teaching_class_members tcm
        LEFT JOIN users u ON u.id = tcm.student_user_id
        WHERE tcm.class_id = ?
        ORDER BY tcm.joined_at DESC
        """,
        (class_id,),
    )

    return {
        "class": {
            "id": class_row["id"],
            "teacher_user_id": class_row["teacher_user_id"],
            "name": class_row["name"] or "",
            "description": class_row["description"] or "",
            "status": class_row["status"],
        },
        "members": [
            {
                "class_id": r["class_id"],
                "student_user_id": r["student_user_id"],
                "student_name": r["student_name"] or r["student_email"] or "",
                "student_email": r["student_email"] or "",
                "joined_at": r["joined_at"],
            }
            for r in rows
        ],
    }


@conv_router.delete("/teaching/classes/{class_id}/members/{student_user_id}")
async def remove_teaching_class_member(
    class_id: int,
    student_user_id: int,
    user: dict = Depends(get_current_user),
):
    """教师移除班级成员。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可管理班级成员")

    await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    existed = await db.execute_fetchone(
        "SELECT 1 FROM teaching_class_members WHERE class_id = ? AND student_user_id = ? LIMIT 1",
        (class_id, student_user_id),
    )
    if existed:
        await db.execute(
            "DELETE FROM teaching_class_members WHERE class_id = ? AND student_user_id = ?",
            (class_id, student_user_id),
        )
        await db.commit()

    return {
        "class_id": class_id,
        "student_user_id": student_user_id,
        "removed": bool(existed),
    }


@conv_router.post("/teaching/classes/{class_id}/assignments")
async def create_teaching_assignment(
    class_id: int,
    req: TeachingAssignmentCreateRequest,
    user: dict = Depends(get_current_user),
):
    """教师在班级内创建作业。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可创建作业")

    class_row = await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    title = str(req.title or "").strip()[:160]
    if not title:
        raise HTTPException(status_code=400, detail="作业标题不能为空")
    description = str(req.description or "").strip()[:4000]
    due_at = str(req.due_at or "").strip()[:64]

    cur = await db.execute(
        """
        INSERT INTO teaching_assignments
            (class_id, teacher_user_id, title, description, due_at, status, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, 'active', CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        """,
        (class_id, user["id"], title, description, due_at),
    )

    assignment_id = int(cur.lastrowid or 0)
    row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, title, description, due_at, status, created_at, updated_at
        FROM teaching_assignments
        WHERE id = ?
        """,
        (assignment_id,),
    )
    if not row:
        raise HTTPException(status_code=500, detail="作业创建失败")

    return {
        "id": row["id"],
        "class_id": row["class_id"],
        "class_name": class_row["name"] or "",
        "teacher_user_id": row["teacher_user_id"],
        "title": row["title"] or "",
        "description": row["description"] or "",
        "due_at": row["due_at"] or "",
        "status": row["status"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }


@conv_router.get("/teaching/classes/{class_id}/assignments")
async def list_teaching_assignments(
    class_id: int,
    user: dict = Depends(get_current_user),
    status: str = Query("", description="按作业状态过滤"),
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
):
    """班级作业列表（班级教师或该班学生可查看）。"""
    db = await get_db()
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")

    account_role = await _resolve_requester_account_role(db, user)
    requester_id = int(user["id"])
    if account_role == "teacher":
        if int(class_row["teacher_user_id"] or 0) != requester_id:
            raise HTTPException(status_code=403, detail="无权查看该班级作业")
    elif account_role == "student":
        in_class = await _is_teaching_class_member(db, class_id=class_id, student_user_id=requester_id)
        if not in_class:
            raise HTTPException(status_code=403, detail="无权查看该班级作业")
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看班级作业")

    status_token = str(status or "").strip().lower()
    sql = """
        SELECT
            ta.id,
            ta.class_id,
            ta.teacher_user_id,
            ta.title,
            ta.description,
            ta.due_at,
            ta.status,
            ta.created_at,
            ta.updated_at,
            (SELECT COUNT(*) FROM teaching_assignment_submissions tas WHERE tas.assignment_id = ta.id) AS submission_count,
            (SELECT tas2.status FROM teaching_assignment_submissions tas2
             WHERE tas2.assignment_id = ta.id AND tas2.student_user_id = ? LIMIT 1) AS my_submission_status,
            (SELECT tas2.submitted_at FROM teaching_assignment_submissions tas2
             WHERE tas2.assignment_id = ta.id AND tas2.student_user_id = ? LIMIT 1) AS my_submitted_at,
            (SELECT tas2.teaching_submission_id FROM teaching_assignment_submissions tas2
             WHERE tas2.assignment_id = ta.id AND tas2.student_user_id = ? LIMIT 1) AS my_teaching_submission_id
        FROM teaching_assignments ta
        WHERE ta.class_id = ?
    """
    params: list[Any] = [requester_id, requester_id, requester_id, class_id]
    if status_token:
        sql += " AND ta.status = ?"
        params.append(status_token)

    sql += " ORDER BY ta.created_at DESC LIMIT ? OFFSET ?"
    params.extend([limit, offset])

    rows = await db.execute_fetchall(sql, tuple(params))

    count_sql = "SELECT COUNT(*) AS cnt FROM teaching_assignments WHERE class_id = ?"
    count_params: list[Any] = [class_id]
    if status_token:
        count_sql += " AND status = ?"
        count_params.append(status_token)
    total_row = await db.execute_fetchone(count_sql, tuple(count_params))

    return {
        "class": {
            "id": class_row["id"],
            "teacher_user_id": class_row["teacher_user_id"],
            "name": class_row["name"] or "",
            "description": class_row["description"] or "",
            "status": class_row["status"],
        },
        "assignments": [
            {
                "id": r["id"],
                "class_id": r["class_id"],
                "teacher_user_id": r["teacher_user_id"],
                "title": r["title"] or "",
                "description": r["description"] or "",
                "due_at": r["due_at"] or "",
                "status": r["status"],
                "submission_count": int(r["submission_count"] or 0),
                "my_submission_status": r["my_submission_status"] or "",
                "my_submitted_at": r["my_submitted_at"],
                "my_teaching_submission_id": r["my_teaching_submission_id"],
                "created_at": r["created_at"],
                "updated_at": r["updated_at"],
            }
            for r in rows
        ],
        "total": int(total_row["cnt"] or 0) if total_row else len(rows),
    }


@conv_router.post("/{conversation_id}/teaching/assignments/{assignment_id}/submit")
async def submit_teaching_assignment_conversation(
    conversation_id: str,
    assignment_id: int,
    req: TeachingAssignmentSubmitRequest,
    user: dict = Depends(get_current_user),
):
    """学生提交班级作业（绑定现有教学评阅链路）。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "student":
        raise HTTPException(status_code=403, detail="仅学生账号可提交作业")

    assignment_row = await db.execute_fetchone(
        """
        SELECT
            ta.id,
            ta.class_id,
            ta.teacher_user_id,
            ta.title,
            ta.status,
            tc.name AS class_name,
            tc.teacher_user_id AS class_teacher_user_id
        FROM teaching_assignments ta
        JOIN teaching_classes tc ON tc.id = ta.class_id
        WHERE ta.id = ?
        """,
        (assignment_id,),
    )
    if not assignment_row:
        raise HTTPException(status_code=404, detail="作业不存在")

    class_id = int(assignment_row["class_id"])
    teacher_user_id = int(assignment_row["class_teacher_user_id"] or assignment_row["teacher_user_id"] or 0)
    if teacher_user_id <= 0:
        raise HTTPException(status_code=500, detail="作业教师配置异常")

    in_class = await _is_teaching_class_member(
        db,
        class_id=class_id,
        student_user_id=int(user["id"]),
    )
    if not in_class:
        raise HTTPException(status_code=403, detail="仅班级成员可提交该作业")

    exists = await db.execute_fetchone(
        "SELECT 1 FROM messages WHERE conversation_id = ? AND user_id = ? LIMIT 1",
        (conversation_id, user["id"]),
    )
    if not exists:
        raise HTTPException(status_code=404, detail="会话不存在或不属于当前学生")

    note = str(req.note or "").strip()[:500]

    await db.execute(
        """
        INSERT INTO teaching_submissions
            (conversation_id, student_user_id, teacher_user_id, status, note, submitted_at, updated_at)
        VALUES (?, ?, ?, 'submitted', ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        ON CONFLICT(conversation_id, student_user_id, teacher_user_id)
        DO UPDATE SET
            status = 'submitted',
            note = excluded.note,
            submitted_at = CURRENT_TIMESTAMP,
            updated_at = CURRENT_TIMESTAMP
        """,
        (conversation_id, user["id"], teacher_user_id, note),
    )

    submission_row = await db.execute_fetchone(
        """
        SELECT id, status, note, submitted_at, updated_at
        FROM teaching_submissions
        WHERE conversation_id = ? AND student_user_id = ? AND teacher_user_id = ?
        """,
        (conversation_id, user["id"], teacher_user_id),
    )
    if not submission_row:
        raise HTTPException(status_code=500, detail="教学评阅提交写入失败")

    await db.execute(
        """
        INSERT INTO teaching_assignment_submissions
            (assignment_id, class_id, teaching_submission_id, conversation_id, student_user_id, teacher_user_id,
             status, note, submitted_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, 'submitted', ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        ON CONFLICT(assignment_id, student_user_id)
        DO UPDATE SET
            teaching_submission_id = excluded.teaching_submission_id,
            conversation_id = excluded.conversation_id,
            status = 'submitted',
            note = excluded.note,
            submitted_at = CURRENT_TIMESTAMP,
            updated_at = CURRENT_TIMESTAMP
        """,
        (
            assignment_id,
            class_id,
            submission_row["id"],
            conversation_id,
            user["id"],
            teacher_user_id,
            note,
        ),
    )
    await db.commit()

    assignment_submission = await db.execute_fetchone(
        """
        SELECT
            id,
            assignment_id,
            class_id,
            teaching_submission_id,
            conversation_id,
            student_user_id,
            teacher_user_id,
            status,
            note,
            submitted_at,
            updated_at
        FROM teaching_assignment_submissions
        WHERE assignment_id = ? AND student_user_id = ?
        """,
        (assignment_id, user["id"]),
    )
    if not assignment_submission:
        raise HTTPException(status_code=500, detail="作业提交写入失败")

    return {
        "id": assignment_submission["id"],
        "assignment_id": assignment_submission["assignment_id"],
        "assignment_title": assignment_row["title"] or "",
        "class_id": assignment_submission["class_id"],
        "class_name": assignment_row["class_name"] or "",
        "teaching_submission_id": assignment_submission["teaching_submission_id"],
        "conversation_id": assignment_submission["conversation_id"],
        "student_user_id": assignment_submission["student_user_id"],
        "teacher_user_id": assignment_submission["teacher_user_id"],
        "status": assignment_submission["status"],
        "note": assignment_submission["note"] or "",
        "submitted_at": assignment_submission["submitted_at"],
        "updated_at": assignment_submission["updated_at"],
    }


@conv_router.get("/teaching/assignments/{assignment_id}/submissions")
async def list_teaching_assignment_submissions(
    assignment_id: int,
    user: dict = Depends(get_current_user),
    status: str = Query("", description="按提交状态过滤"),
    student_user_id: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
):
    """教师查看某个作业的提交列表。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可查看作业提交列表")

    assignment_row = await db.execute_fetchone(
        """
        SELECT ta.id, ta.class_id, ta.teacher_user_id, ta.title, ta.status, tc.name AS class_name
        FROM teaching_assignments ta
        JOIN teaching_classes tc ON tc.id = ta.class_id
        WHERE ta.id = ?
        """,
        (assignment_id,),
    )
    if not assignment_row:
        raise HTTPException(status_code=404, detail="作业不存在")
    if int(assignment_row["teacher_user_id"] or 0) != int(user["id"]):
        raise HTTPException(status_code=403, detail="无权查看该作业提交")

    status_token = str(status or "").strip().lower()
    sql = """
        SELECT
            tas.id,
            tas.assignment_id,
            tas.class_id,
            tas.teaching_submission_id,
            tas.conversation_id,
            tas.student_user_id,
            su.name AS student_name,
            su.email AS student_email,
            tas.teacher_user_id,
            tas.status,
            tas.note,
            tas.submitted_at,
            tas.updated_at,
            (SELECT COUNT(*) FROM teaching_evaluations te
             WHERE te.submission_id = tas.teaching_submission_id) AS evaluation_count,
            (SELECT content FROM messages m2
             WHERE m2.conversation_id = tas.conversation_id
               AND m2.user_id = tas.student_user_id
               AND m2.role = 'user'
             ORDER BY m2.created_at ASC LIMIT 1) AS first_user_msg,
            (SELECT MAX(created_at) FROM messages m3
             WHERE m3.conversation_id = tas.conversation_id
               AND m3.user_id = tas.student_user_id) AS last_at
        FROM teaching_assignment_submissions tas
        LEFT JOIN users su ON su.id = tas.student_user_id
        WHERE tas.assignment_id = ?
    """
    params: list[Any] = [assignment_id]
    if status_token:
        sql += " AND tas.status = ?"
        params.append(status_token)
    if int(student_user_id or 0) > 0:
        sql += " AND tas.student_user_id = ?"
        params.append(int(student_user_id))

    sql += " ORDER BY tas.submitted_at DESC LIMIT ? OFFSET ?"
    params.extend([limit, offset])

    rows = await db.execute_fetchall(sql, tuple(params))

    count_sql = "SELECT COUNT(*) AS cnt FROM teaching_assignment_submissions WHERE assignment_id = ?"
    count_params: list[Any] = [assignment_id]
    if status_token:
        count_sql += " AND status = ?"
        count_params.append(status_token)
    if int(student_user_id or 0) > 0:
        count_sql += " AND student_user_id = ?"
        count_params.append(int(student_user_id))
    total_row = await db.execute_fetchone(count_sql, tuple(count_params))

    return {
        "assignment": {
            "id": assignment_row["id"],
            "title": assignment_row["title"] or "",
            "status": assignment_row["status"],
            "class_id": assignment_row["class_id"],
            "class_name": assignment_row["class_name"] or "",
            "teacher_user_id": assignment_row["teacher_user_id"],
        },
        "submissions": [
            {
                "id": r["id"],
                "assignment_id": r["assignment_id"],
                "class_id": r["class_id"],
                "teaching_submission_id": r["teaching_submission_id"],
                "conversation_id": r["conversation_id"],
                "student_user_id": r["student_user_id"],
                "student_name": r["student_name"] or r["student_email"] or "",
                "teacher_user_id": r["teacher_user_id"],
                "status": r["status"],
                "note": r["note"] or "",
                "title": (r["first_user_msg"] or "")[:60] or "（无标题）",
                "last_at": r["last_at"],
                "evaluation_count": int(r["evaluation_count"] or 0),
                "submitted_at": r["submitted_at"],
                "updated_at": r["updated_at"],
            }
            for r in rows
        ],
        "total": int(total_row["cnt"] or 0) if total_row else len(rows),
    }

@conv_router.get("/teaching/classes/{class_id}/dashboard")
async def get_teaching_class_dashboard(
    class_id: int,
    user: dict = Depends(get_current_user),
    assignment_limit: int = Query(80, ge=1, le=300),
):
    """课堂治理看板：按班级聚合作业提交、评阅与评分分布。"""
    db = await get_db()
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")

    account_role = await _resolve_requester_account_role(db, user)
    requester_id = int(user["id"])

    if account_role == "teacher":
        if int(class_row["teacher_user_id"] or 0) != requester_id:
            raise HTTPException(status_code=403, detail="无权查看该班级看板")
    elif account_role == "student":
        in_class = await _is_teaching_class_member(db, class_id=class_id, student_user_id=requester_id)
        if not in_class:
            raise HTTPException(status_code=403, detail="无权查看该班级看板")
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看课堂看板")

    member_row = await db.execute_fetchone(
        "SELECT COUNT(*) AS cnt FROM teaching_class_members WHERE class_id = ?",
        (class_id,),
    )
    member_count = int(member_row["cnt"] or 0) if member_row else 0

    assignment_rows = await db.execute_fetchall(
        """
        SELECT
            ta.id,
            ta.class_id,
            ta.teacher_user_id,
            ta.title,
            ta.description,
            ta.due_at,
            ta.status,
            ta.created_at,
            ta.updated_at,
            (SELECT COUNT(*) FROM teaching_assignment_submissions tas
             WHERE tas.assignment_id = ta.id) AS submitted_count,
            (SELECT COUNT(*) FROM teaching_assignment_submissions tas
               JOIN teaching_submissions ts ON ts.id = tas.teaching_submission_id
             WHERE tas.assignment_id = ta.id AND ts.status = 'reviewed') AS reviewed_count,
            (SELECT MAX(tas.submitted_at) FROM teaching_assignment_submissions tas
             WHERE tas.assignment_id = ta.id) AS latest_submitted_at,
            (SELECT AVG(te.score) FROM teaching_assignment_submissions tas
               JOIN teaching_evaluations te ON te.submission_id = tas.teaching_submission_id
             WHERE tas.assignment_id = ta.id) AS avg_score
        FROM teaching_assignments ta
        WHERE ta.class_id = ?
        ORDER BY ta.created_at DESC
        LIMIT ?
        """,
        (class_id, assignment_limit),
    )

    assignment_count_row = await db.execute_fetchone(
        "SELECT COUNT(*) AS cnt FROM teaching_assignments WHERE class_id = ?",
        (class_id,),
    )
    assignment_count = int(assignment_count_row["cnt"] or 0) if assignment_count_row else len(assignment_rows)

    submitted_count_row = await db.execute_fetchone(
        "SELECT COUNT(*) AS cnt FROM teaching_assignment_submissions WHERE class_id = ?",
        (class_id,),
    )
    submitted_count = int(submitted_count_row["cnt"] or 0) if submitted_count_row else 0

    reviewed_count_row = await db.execute_fetchone(
        """
        SELECT COUNT(*) AS cnt
        FROM teaching_assignment_submissions tas
        JOIN teaching_submissions ts ON ts.id = tas.teaching_submission_id
        WHERE tas.class_id = ? AND ts.status = 'reviewed'
        """,
        (class_id,),
    )
    reviewed_count = int(reviewed_count_row["cnt"] or 0) if reviewed_count_row else 0

    score_rows = await db.execute_fetchall(
        """
        SELECT te.score AS score
        FROM teaching_assignment_submissions tas
        JOIN teaching_evaluations te ON te.submission_id = tas.teaching_submission_id
        WHERE tas.class_id = ?
        """,
        (class_id,),
    )

    score_distribution = {
        "lt60": 0,
        "s60_79": 0,
        "s80_89": 0,
        "s90_plus": 0,
        "count": 0,
        "avg_score": 0.0,
    }
    score_sum = 0.0
    for row in score_rows:
        score = float(row["score"] or 0)
        score_distribution["count"] += 1
        score_sum += score
        if score < 60:
            score_distribution["lt60"] += 1
        elif score < 80:
            score_distribution["s60_79"] += 1
        elif score < 90:
            score_distribution["s80_89"] += 1
        else:
            score_distribution["s90_plus"] += 1
    if score_distribution["count"] > 0:
        score_distribution["avg_score"] = round(score_sum / score_distribution["count"], 2)

    assignments = []
    for row in assignment_rows:
        submitted = int(row["submitted_count"] or 0)
        reviewed = int(row["reviewed_count"] or 0)
        pending = max(submitted - reviewed, 0)
        missing = max(member_count - submitted, 0)
        submission_rate = (submitted / member_count) if member_count > 0 else 0.0
        review_rate = (reviewed / submitted) if submitted > 0 else 0.0
        avg_score = float(row["avg_score"] or 0)
        assignments.append(
            {
                "id": row["id"],
                "class_id": row["class_id"],
                "teacher_user_id": row["teacher_user_id"],
                "title": row["title"] or "",
                "description": row["description"] or "",
                "due_at": row["due_at"] or "",
                "status": row["status"],
                "submitted_count": submitted,
                "reviewed_count": reviewed,
                "pending_review_count": pending,
                "missing_submission_count": missing,
                "submission_rate": round(submission_rate, 4),
                "review_completion_rate": round(review_rate, 4),
                "avg_score": round(avg_score, 2) if submitted > 0 else 0.0,
                "latest_submitted_at": row["latest_submitted_at"],
                "created_at": row["created_at"],
                "updated_at": row["updated_at"],
            }
        )

    expected_submissions = member_count * assignment_count
    pending_review_count = max(submitted_count - reviewed_count, 0)
    missing_submission_count = max(expected_submissions - submitted_count, 0)

    submission_gap_ratio = (missing_submission_count / expected_submissions) if expected_submissions > 0 else 0.0
    review_backlog_ratio = (pending_review_count / submitted_count) if submitted_count > 0 else 0.0
    low_score_ratio = (score_distribution["lt60"] / score_distribution["count"]) if score_distribution["count"] > 0 else 0.0

    interventions: list[dict[str, Any]] = []
    if expected_submissions > 0 and missing_submission_count > 0 and submission_gap_ratio >= 0.3:
        interventions.append(
            {
                "code": "submission_at_risk",
                "severity": "high" if submission_gap_ratio >= 0.5 else "medium",
                "title": "提交完成率偏低",
                "detail": f"当前缺交 {missing_submission_count}/{expected_submissions}，建议优先跟进未提交学生。",
                "action": "可按班级成员名单发送提醒，并设置补交截止时间。",
            }
        )

    if submitted_count > 0 and pending_review_count > 0 and review_backlog_ratio >= 0.35:
        interventions.append(
            {
                "code": "review_backlog",
                "severity": "high" if review_backlog_ratio >= 0.6 else "medium",
                "title": "评阅积压较高",
                "detail": f"待评阅 {pending_review_count}/{submitted_count}，建议本周优先清理积压。",
                "action": "可优先处理临近截止作业，并使用评分模板提升评阅效率。",
            }
        )

    if score_distribution["count"] >= 5 and low_score_ratio >= 0.25:
        interventions.append(
            {
                "code": "low_score_risk",
                "severity": "medium",
                "title": "低分占比偏高",
                "detail": f"当前低于60分占比 {round(low_score_ratio * 100)}%，建议安排集中讲解与补练。",
                "action": "按低分维度补充示例答案与评分标准说明，减少重复失分。",
            }
        )

    risky_assignment = None
    if assignments:
        risky_assignment = max(
            assignments,
            key=lambda item: (
                int(item.get("missing_submission_count") or 0),
                int(item.get("pending_review_count") or 0),
            ),
        )
    if risky_assignment and (
        int(risky_assignment.get("missing_submission_count") or 0) > 0
        or int(risky_assignment.get("pending_review_count") or 0) > 0
    ):
        interventions.append(
            {
                "code": "focus_assignment",
                "severity": "medium",
                "title": "建议优先关注高风险作业",
                "detail": (
                    f"“{str(risky_assignment.get('title') or '未命名作业')}”缺交"
                    f" {int(risky_assignment.get('missing_submission_count') or 0)}，"
                    f"待评阅 {int(risky_assignment.get('pending_review_count') or 0)}。"
                ),
                "action": "可先完成该作业的批量提醒与评阅，快速降低班级风险。",
                "assignment_id": int(risky_assignment.get("id") or 0),
            }
        )

    teacher_owner_id = int(class_row["teacher_user_id"] or 0)
    intervention_action_summary = await _build_teaching_intervention_action_summary(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
    )
    action_code_stat_map = {
        str(item.get("intervention_code") or "").strip().lower(): item
        for item in (intervention_action_summary.get("code_stats") or [])
        if str(item.get("intervention_code") or "").strip()
    }

    can_manage_goal_and_intervention = account_role == "teacher" and requester_id == teacher_owner_id
    for item in interventions:
        code = str(item.get("code") or "").strip().lower()
        stat = action_code_stat_map.get(code, {})
        item["executed_count"] = int(stat.get("action_count") or 0)
        item["executed_recent_7d"] = int(stat.get("recent_7d_count") or 0)
        item["last_executed_at"] = str(stat.get("last_action_at") or "")
        item["can_mark_executed"] = can_manage_goal_and_intervention

    goal_metric_values = _build_teaching_class_goal_metric_values(
        expected_submission_count=expected_submissions,
        submitted_count=submitted_count,
        reviewed_count=reviewed_count,
        score_distribution=score_distribution,
    )
    goal_rows = await _list_teaching_class_goal_rows(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        include_archived=False,
    )
    class_goals = [
        _serialize_teaching_class_goal_row(row, metric_values=goal_metric_values)
        for row in goal_rows
    ]
    for goal in class_goals:
        goal["can_manage"] = can_manage_goal_and_intervention
    class_goal_summary = _build_teaching_class_goal_summary(class_goals)
    goal_trend_bundle = await _build_teaching_goal_trend_bundle(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        goals=class_goals,
        window_days=30,
        snapshot_source="dashboard",
    )
    intervention_outcome_correlation = goal_trend_bundle.get("intervention_outcome_correlation") or {}
    strategy_bundle = _build_teaching_intervention_strategy_recommendations(
        interventions=interventions,
        intervention_action_summary=intervention_action_summary,
        intervention_outcome_correlation=intervention_outcome_correlation,
        goal_risk_summary=goal_trend_bundle.get("risk_summary") or {},
    )
    experiment_rows, _experiment_total = await _list_teaching_intervention_experiment_rows(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        limit=6,
        offset=0,
    )
    intervention_experiments = [
        _serialize_teaching_intervention_experiment_row(row, can_manage=can_manage_goal_and_intervention)
        for row in experiment_rows
    ]
    intervention_experiment_summary = _build_teaching_intervention_experiment_status_summary(
        intervention_experiments,
    )

    return {
        "class": {
            "id": class_row["id"],
            "teacher_user_id": class_row["teacher_user_id"],
            "name": class_row["name"] or "",
            "description": class_row["description"] or "",
            "status": class_row["status"],
            "created_at": class_row["created_at"],
            "updated_at": class_row["updated_at"],
        },
        "overview": {
            "member_count": member_count,
            "assignment_count": assignment_count,
            "expected_submission_count": expected_submissions,
            "submitted_count": submitted_count,
            "reviewed_count": reviewed_count,
            "pending_review_count": pending_review_count,
            "missing_submission_count": missing_submission_count,
            "review_completion_rate": round((reviewed_count / submitted_count), 4) if submitted_count > 0 else 0.0,
        },
        "score_distribution": score_distribution,
        "assignments": assignments,
        "interventions": interventions,
        "intervention_action_summary": intervention_action_summary,
        "intervention_strategy_recommendations": strategy_bundle.get("recommendations") or [],
        "intervention_strategy_summary": strategy_bundle.get("summary") or {},
        "intervention_experiments": intervention_experiments,
        "intervention_experiment_summary": intervention_experiment_summary,
        "class_goals": class_goals,
        "class_goal_summary": class_goal_summary,
        "class_goal_metric_values": goal_metric_values,
        "class_goal_trend_30d": goal_trend_bundle.get("goal_trends") or [],
        "class_goal_risk_summary": goal_trend_bundle.get("risk_summary") or {},
        "class_goal_trend_overview": goal_trend_bundle.get("overview") or {},
        "intervention_outcome_correlation": intervention_outcome_correlation,
    }

@conv_router.post("/teaching/classes/{class_id}/goals")
async def upsert_teaching_class_goal(
    class_id: int,
    req: TeachingClassGoalCreateRequest,
    user: dict = Depends(get_current_user),
):
    """教师创建或更新课堂目标（按 goal_code 幂等）。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可管理课堂目标")

    await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    goal_code = _normalize_teaching_class_goal_code(req.goal_code)
    goal_name = str(req.goal_name or "").strip()[:120]
    if not goal_name:
        raise HTTPException(status_code=400, detail="goal_name 不能为空")
    metric_type = _normalize_teaching_class_goal_metric_type(req.metric_type)
    target_value = float(req.target_value or 0)
    if target_value <= 0:
        raise HTTPException(status_code=400, detail="target_value 必须大于0")

    note = str(req.note or "").strip()[:1000]
    due_at = str(req.due_at or "").strip()[:64]

    await db.execute(
        """
        INSERT INTO teaching_class_goals
            (class_id, teacher_user_id, goal_code, goal_name, metric_type, target_value,
             status, note, due_at, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, 'active', ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        ON CONFLICT(class_id, teacher_user_id, goal_code)
        DO UPDATE SET
            goal_name = excluded.goal_name,
            metric_type = excluded.metric_type,
            target_value = excluded.target_value,
            status = 'active',
            note = excluded.note,
            due_at = excluded.due_at,
            updated_at = CURRENT_TIMESTAMP
        """,
        (
            class_id,
            int(user["id"]),
            goal_code,
            goal_name,
            metric_type,
            target_value,
            note,
            due_at,
        ),
    )
    await db.commit()

    row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, goal_code, goal_name, metric_type, target_value,
               status, note, due_at, created_at, updated_at
        FROM teaching_class_goals
        WHERE class_id = ? AND teacher_user_id = ? AND goal_code = ?
        LIMIT 1
        """,
        (class_id, int(user["id"]), goal_code),
    )
    if not row:
        raise HTTPException(status_code=500, detail="课堂目标保存失败")

    metric_values, _ctx = await _compute_teaching_class_goal_metric_values(db, class_id=class_id)
    goal = _serialize_teaching_class_goal_row(row, metric_values=metric_values)
    goal["can_manage"] = True

    rows = await _list_teaching_class_goal_rows(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
        include_archived=False,
    )
    goals = [_serialize_teaching_class_goal_row(item, metric_values=metric_values) for item in rows]
    summary = _build_teaching_class_goal_summary(goals)
    trend_bundle = await _build_teaching_goal_trend_bundle(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
        goals=goals,
        window_days=30,
        snapshot_source="goal_upsert",
    )

    return {
        "goal": goal,
        "summary": summary,
        "metric_values": metric_values,
        "goal_trends": trend_bundle.get("goal_trends") or [],
        "risk_summary": trend_bundle.get("risk_summary") or {},
        "trend_overview": trend_bundle.get("overview") or {},
        "window_days": int(trend_bundle.get("window_days") or 30),
    }


@conv_router.get("/teaching/classes/{class_id}/goals")
async def list_teaching_class_goals(
    class_id: int,
    user: dict = Depends(get_current_user),
    include_archived: bool = Query(False, description="是否包含 archived 目标，仅教师本人可用"),
):
    """查看课堂目标进度（教师可管理，学生可查看）。"""
    db = await get_db()
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")

    account_role = await _resolve_requester_account_role(db, user)
    requester_id = int(user["id"])
    teacher_owner_id = int(class_row["teacher_user_id"] or 0)

    if account_role == "teacher":
        if requester_id != teacher_owner_id:
            raise HTTPException(status_code=403, detail="无权查看该班级目标")
    elif account_role == "student":
        in_class = await _is_teaching_class_member(db, class_id=class_id, student_user_id=requester_id)
        if not in_class:
            raise HTTPException(status_code=403, detail="无权查看该班级目标")
        include_archived = False
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看课堂目标")

    metric_values, context = await _compute_teaching_class_goal_metric_values(db, class_id=class_id)
    rows = await _list_teaching_class_goal_rows(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        include_archived=bool(include_archived and account_role == "teacher" and requester_id == teacher_owner_id),
    )
    can_manage = account_role == "teacher" and requester_id == teacher_owner_id
    goals = []
    for row in rows:
        item = _serialize_teaching_class_goal_row(row, metric_values=metric_values)
        item["can_manage"] = can_manage
        goals.append(item)

    summary = _build_teaching_class_goal_summary(goals)
    trend_bundle = await _build_teaching_goal_trend_bundle(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        goals=goals,
        window_days=30,
        snapshot_source="goal_list",
    )
    return {
        "class_id": class_id,
        "goals": goals,
        "summary": summary,
        "metric_values": metric_values,
        "context": context,
        "goal_trends": trend_bundle.get("goal_trends") or [],
        "risk_summary": trend_bundle.get("risk_summary") or {},
        "trend_overview": trend_bundle.get("overview") or {},
        "window_days": int(trend_bundle.get("window_days") or 30),
    }


@conv_router.post("/teaching/classes/{class_id}/goals/{goal_id}/status")
async def update_teaching_class_goal_status(
    class_id: int,
    goal_id: int,
    req: TeachingClassGoalStatusUpdateRequest,
    user: dict = Depends(get_current_user),
):
    """教师更新课堂目标状态。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可更新课堂目标")

    await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, goal_code, goal_name, metric_type, target_value,
               status, note, due_at, created_at, updated_at
        FROM teaching_class_goals
        WHERE id = ? AND class_id = ? AND teacher_user_id = ?
        LIMIT 1
        """,
        (goal_id, class_id, int(user["id"])),
    )
    if not row:
        raise HTTPException(status_code=404, detail="课堂目标不存在")

    target_status = _normalize_teaching_class_goal_status(req.status)
    note = str(req.note or "").strip()[:1000]
    next_note = note if note else str(row["note"] or "")

    await db.execute(
        """
        UPDATE teaching_class_goals
        SET status = ?, note = ?, updated_at = CURRENT_TIMESTAMP
        WHERE id = ? AND class_id = ? AND teacher_user_id = ?
        """,
        (target_status, next_note, goal_id, class_id, int(user["id"])),
    )
    await db.commit()

    updated = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, goal_code, goal_name, metric_type, target_value,
               status, note, due_at, created_at, updated_at
        FROM teaching_class_goals
        WHERE id = ?
        LIMIT 1
        """,
        (goal_id,),
    )
    if not updated:
        raise HTTPException(status_code=500, detail="课堂目标状态更新失败")

    metric_values, _ctx = await _compute_teaching_class_goal_metric_values(db, class_id=class_id)
    goal = _serialize_teaching_class_goal_row(updated, metric_values=metric_values)
    goal["can_manage"] = True

    rows = await _list_teaching_class_goal_rows(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
        include_archived=False,
    )
    goals = [_serialize_teaching_class_goal_row(item, metric_values=metric_values) for item in rows]
    summary = _build_teaching_class_goal_summary(goals)
    trend_bundle = await _build_teaching_goal_trend_bundle(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
        goals=goals,
        window_days=30,
        snapshot_source="goal_status_update",
    )

    return {
        "goal": goal,
        "summary": summary,
        "metric_values": metric_values,
        "goal_trends": trend_bundle.get("goal_trends") or [],
        "risk_summary": trend_bundle.get("risk_summary") or {},
        "trend_overview": trend_bundle.get("overview") or {},
        "window_days": int(trend_bundle.get("window_days") or 30),
    }


@conv_router.get("/teaching/classes/{class_id}/goals/trend")
async def get_teaching_class_goal_trend(
    class_id: int,
    user: dict = Depends(get_current_user),
    days: int = Query(30, ge=7, le=180),
    include_archived: bool = Query(False, description="是否包含 archived 目标，仅教师本人可用"),
):
    """查看课堂目标趋势与干预-结果相关性。"""
    db = await get_db()
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")

    account_role = await _resolve_requester_account_role(db, user)
    requester_id = int(user["id"])
    teacher_owner_id = int(class_row["teacher_user_id"] or 0)

    if account_role == "teacher":
        if requester_id != teacher_owner_id:
            raise HTTPException(status_code=403, detail="无权查看该班级目标趋势")
    elif account_role == "student":
        in_class = await _is_teaching_class_member(db, class_id=class_id, student_user_id=requester_id)
        if not in_class:
            raise HTTPException(status_code=403, detail="无权查看该班级目标趋势")
        include_archived = False
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看课堂目标趋势")

    metric_values, context = await _compute_teaching_class_goal_metric_values(db, class_id=class_id)
    rows = await _list_teaching_class_goal_rows(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        include_archived=bool(include_archived and account_role == "teacher" and requester_id == teacher_owner_id),
    )

    can_manage = account_role == "teacher" and requester_id == teacher_owner_id
    goals: list[dict[str, Any]] = []
    for row in rows:
        item = _serialize_teaching_class_goal_row(row, metric_values=metric_values)
        item["can_manage"] = can_manage
        goals.append(item)

    summary = _build_teaching_class_goal_summary(goals)
    trend_bundle = await _build_teaching_goal_trend_bundle(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        goals=goals,
        window_days=days,
        snapshot_source="goal_trend_api",
    )

    return {
        "class_id": class_id,
        "goals": goals,
        "summary": summary,
        "metric_values": metric_values,
        "context": context,
        "window_days": int(trend_bundle.get("window_days") or _normalize_teaching_goal_trend_window_days(days)),
        "goal_trends": trend_bundle.get("goal_trends") or [],
        "risk_summary": trend_bundle.get("risk_summary") or {},
        "trend_overview": trend_bundle.get("overview") or {},
        "intervention_outcome_correlation": trend_bundle.get("intervention_outcome_correlation") or {},
    }


@conv_router.get("/teaching/classes/{class_id}/interventions/recommendations")
async def get_teaching_intervention_strategy_recommendations(
    class_id: int,
    user: dict = Depends(get_current_user),
    days: int = Query(30, ge=7, le=180),
    assignment_limit: int = Query(80, ge=1, le=300),
):
    """查看课堂干预策略优化建议（优先级 + 策略模式 + A/B分组建议）。"""
    dashboard_payload = await get_teaching_class_dashboard(
        class_id=class_id,
        user=user,
        assignment_limit=assignment_limit,
    )

    class_info = dashboard_payload.get("class") or {}
    teacher_owner_id = int(class_info.get("teacher_user_id") or 0)
    goals = dashboard_payload.get("class_goals") or []

    db = await get_db()
    trend_bundle = await _build_teaching_goal_trend_bundle(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        goals=goals,
        window_days=days,
        snapshot_source="strategy_recommendations_api",
    )
    intervention_outcome_correlation = trend_bundle.get("intervention_outcome_correlation") or {}
    strategy_bundle = _build_teaching_intervention_strategy_recommendations(
        interventions=dashboard_payload.get("interventions") or [],
        intervention_action_summary=dashboard_payload.get("intervention_action_summary") or {},
        intervention_outcome_correlation=intervention_outcome_correlation,
        goal_risk_summary=trend_bundle.get("risk_summary") or {},
    )

    return {
        "class_id": class_id,
        "class": class_info,
        "window_days": int(trend_bundle.get("window_days") or _normalize_teaching_goal_trend_window_days(days)),
        "recommendations": strategy_bundle.get("recommendations") or [],
        "summary": strategy_bundle.get("summary") or {},
        "risk_summary": trend_bundle.get("risk_summary") or {},
        "trend_overview": trend_bundle.get("overview") or {},
        "intervention_outcome_correlation": intervention_outcome_correlation,
    }


@conv_router.post("/teaching/classes/{class_id}/interventions/experiments/plan")
async def upsert_teaching_intervention_experiment_plan(
    class_id: int,
    req: TeachingInterventionExperimentPlanRequest,
    user: dict = Depends(get_current_user),
):
    """教师创建或更新干预实验计划。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可管理干预实验")

    await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    intervention_code = _normalize_teaching_intervention_code(req.intervention_code)
    intervention_title = str(req.intervention_title or "").strip()[:120]
    strategy_mode = _normalize_teaching_experiment_mode(req.strategy_mode)
    recommended_variant = _normalize_teaching_strategy_variant_token(req.recommended_variant) or "A"
    variants = _normalize_teaching_experiment_variants(req.variants, recommended_variant=recommended_variant)
    window_days = max(7, min(90, int(req.window_days or 14)))
    target_metric = _normalize_teaching_experiment_target_metric(req.target_metric)
    note = str(req.note or "").strip()[:500]

    experiment_seed = str(req.experiment_id or "").strip()
    if not experiment_seed:
        experiment_seed = f"{intervention_code}-{datetime.now(timezone.utc).strftime('%Y%m%d')}"
    experiment_code = _normalize_teaching_experiment_code(
        experiment_seed,
        fallback_intervention_code=intervention_code,
    )

    source_recommendation = req.source_recommendation if isinstance(req.source_recommendation, dict) else {}
    variant_plan_payload = {
        "variants": variants,
        "source_recommendation": source_recommendation,
    }

    await db.execute(
        """
        INSERT INTO teaching_intervention_experiments
            (class_id, teacher_user_id, experiment_code, intervention_code, intervention_title,
             strategy_mode, target_metric, recommended_variant, variant_plan_json,
             window_days, status, note, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'planned', ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        ON CONFLICT(class_id, teacher_user_id, experiment_code)
        DO UPDATE SET
            intervention_code = excluded.intervention_code,
            intervention_title = excluded.intervention_title,
            strategy_mode = excluded.strategy_mode,
            target_metric = excluded.target_metric,
            recommended_variant = excluded.recommended_variant,
            variant_plan_json = excluded.variant_plan_json,
            window_days = excluded.window_days,
            note = excluded.note,
            updated_at = CURRENT_TIMESTAMP
        """,
        (
            class_id,
            int(user["id"]),
            experiment_code,
            intervention_code,
            intervention_title,
            strategy_mode,
            target_metric,
            recommended_variant,
            json.dumps(variant_plan_payload, ensure_ascii=False),
            window_days,
            note,
        ),
    )
    await db.commit()

    row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, experiment_code, intervention_code, intervention_title,
               strategy_mode, target_metric, recommended_variant, variant_plan_json,
               window_days, status, note, created_at, updated_at
        FROM teaching_intervention_experiments
        WHERE class_id = ? AND teacher_user_id = ? AND experiment_code = ?
        LIMIT 1
        """,
        (class_id, int(user["id"]), experiment_code),
    )
    if not row:
        raise HTTPException(status_code=500, detail="干预实验计划保存失败")

    return {
        "experiment": _serialize_teaching_intervention_experiment_row(row, can_manage=True),
    }


@conv_router.get("/teaching/classes/{class_id}/interventions/experiments")
async def list_teaching_intervention_experiments(
    class_id: int,
    user: dict = Depends(get_current_user),
    status: str = Query("", description="可选，逗号分隔状态过滤"),
    limit: int = Query(20, ge=1, le=100),
    offset: int = Query(0, ge=0),
):
    """查看干预实验计划列表。"""
    db = await get_db()
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")

    account_role = await _resolve_requester_account_role(db, user)
    requester_id = int(user["id"])
    teacher_owner_id = int(class_row["teacher_user_id"] or 0)

    if account_role == "teacher":
        if requester_id != teacher_owner_id:
            raise HTTPException(status_code=403, detail="无权查看该班级干预实验")
    elif account_role == "student":
        in_class = await _is_teaching_class_member(db, class_id=class_id, student_user_id=requester_id)
        if not in_class:
            raise HTTPException(status_code=403, detail="无权查看该班级干预实验")
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看干预实验")

    status_filters = [item.strip() for item in str(status or "").split(",") if item.strip()]
    rows, total = await _list_teaching_intervention_experiment_rows(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        limit=limit,
        offset=offset,
        statuses=status_filters,
    )

    can_manage = account_role == "teacher" and requester_id == teacher_owner_id
    experiments = [
        _serialize_teaching_intervention_experiment_row(row, can_manage=can_manage)
        for row in rows
    ]
    summary = _build_teaching_intervention_experiment_status_summary(experiments)
    return {
        "class_id": class_id,
        "experiments": experiments,
        "summary": summary,
        "total": int(total),
        "limit": limit,
        "offset": offset,
    }


@conv_router.post("/teaching/classes/{class_id}/interventions/experiments/{experiment_code}/status")
async def update_teaching_intervention_experiment_status(
    class_id: int,
    experiment_code: str,
    req: TeachingInterventionExperimentStatusUpdateRequest,
    user: dict = Depends(get_current_user),
):
    """教师更新干预实验状态。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可管理干预实验")

    await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    normalized_code = _normalize_teaching_experiment_code(experiment_code)
    target_status = _normalize_teaching_experiment_status(req.status)
    note = str(req.note or "").strip()[:500]

    existing = await db.execute_fetchone(
        """
        SELECT id, note
        FROM teaching_intervention_experiments
        WHERE class_id = ? AND teacher_user_id = ? AND experiment_code = ?
        LIMIT 1
        """,
        (class_id, int(user["id"]), normalized_code),
    )
    if not existing:
        raise HTTPException(status_code=404, detail="干预实验不存在")

    next_note = note if note else str(existing["note"] or "")
    await db.execute(
        """
        UPDATE teaching_intervention_experiments
        SET status = ?, note = ?, updated_at = CURRENT_TIMESTAMP
        WHERE class_id = ? AND teacher_user_id = ? AND experiment_code = ?
        """,
        (target_status, next_note, class_id, int(user["id"]), normalized_code),
    )
    await db.commit()

    row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, experiment_code, intervention_code, intervention_title,
               strategy_mode, target_metric, recommended_variant, variant_plan_json,
               window_days, status, note, created_at, updated_at
        FROM teaching_intervention_experiments
        WHERE class_id = ? AND teacher_user_id = ? AND experiment_code = ?
        LIMIT 1
        """,
        (class_id, int(user["id"]), normalized_code),
    )
    if not row:
        raise HTTPException(status_code=500, detail="干预实验状态更新失败")

    return {
        "experiment": _serialize_teaching_intervention_experiment_row(row, can_manage=True),
    }


@conv_router.get("/teaching/classes/{class_id}/interventions/experiments/results")
async def list_teaching_intervention_experiment_results(
    class_id: int,
    user: dict = Depends(get_current_user),
    status: str = Query("", description="可选，逗号分隔状态过滤"),
    experiment_code: str = Query("", description="可选，指定实验编码"),
    days: int = Query(60, ge=7, le=180),
    limit: int = Query(20, ge=1, le=100),
    offset: int = Query(0, ge=0),
):
    """查看干预实验结果汇总与优胜分组建议。"""
    db = await get_db()
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")

    account_role = await _resolve_requester_account_role(db, user)
    requester_id = int(user["id"])
    teacher_owner_id = int(class_row["teacher_user_id"] or 0)

    if account_role == "teacher":
        if requester_id != teacher_owner_id:
            raise HTTPException(status_code=403, detail="无权查看该班级干预实验结果")
    elif account_role == "student":
        in_class = await _is_teaching_class_member(db, class_id=class_id, student_user_id=requester_id)
        if not in_class:
            raise HTTPException(status_code=403, detail="无权查看该班级干预实验结果")
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看干预实验结果")

    status_filters = [item.strip() for item in str(status or "").split(",") if item.strip()]
    experiment_codes: list[str] = []
    if str(experiment_code or "").strip():
        experiment_codes = [str(experiment_code).strip()]
        offset = 0
        limit = 1

    rows, total = await _list_teaching_intervention_experiment_rows(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        limit=limit,
        offset=offset,
        statuses=status_filters,
        experiment_codes=experiment_codes,
    )

    can_manage = account_role == "teacher" and requester_id == teacher_owner_id
    experiments = [
        _serialize_teaching_intervention_experiment_row(row, can_manage=can_manage)
        for row in rows
    ]

    result_map = await _build_teaching_intervention_experiment_result_map(
        db,
        class_id=class_id,
        teacher_user_id=teacher_owner_id,
        experiments=experiments,
        analysis_window_days=days,
    )

    normalized_days = max(7, min(180, int(days or 60)))
    experiment_results: list[dict[str, Any]] = []
    for item in experiments:
        code = str(item.get("experiment_code") or "")
        result = result_map.get(code) or {
            "analysis_window_days": normalized_days,
            "daily_progress_point_count": 0,
            "action_count": 0,
            "observed_action_count": 0,
            "variant_results": [],
            "winner_variant": str(item.get("recommended_variant") or "A"),
            "winner_source": "recommended_fallback",
            "confidence_score": 0.0,
            "guardrail_note": "暂无可用样本，建议先执行实验并记录分组。",
        }
        experiment_results.append(
            {
                **item,
                "result": result,
            }
        )

    low_confidence_count = len(
        [
            item
            for item in experiment_results
            if float((item.get("result") or {}).get("confidence_score") or 0) > 0
            and float((item.get("result") or {}).get("confidence_score") or 0) < 0.6
        ]
    )
    with_observation_count = len(
        [
            item
            for item in experiment_results
            if int((item.get("result") or {}).get("observed_action_count") or 0) > 0
        ]
    )

    summary = {
        "experiment_count": len(experiment_results),
        "with_observation_count": with_observation_count,
        "winner_ready_count": with_observation_count,
        "low_confidence_count": low_confidence_count,
        "analysis_window_days": normalized_days,
    }

    return {
        "class_id": class_id,
        "experiments": experiment_results,
        "summary": summary,
        "total": int(total),
        "limit": limit,
        "offset": offset,
    }


@conv_router.post("/teaching/classes/{class_id}/goals/term-archives")
async def upsert_teaching_goal_term_archive(
    class_id: int,
    req: TeachingGoalTermArchiveRequest,
    user: dict = Depends(get_current_user),
):
    """教师归档某学期目标状态（跨学期长期治理骨架）。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可归档学期目标")

    class_row = await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    term_code = _normalize_teaching_term_code(req.term_code)
    term_name = str(req.term_name or "").strip()[:120]
    term_start = str(req.term_start or "").strip()[:32]
    term_end = str(req.term_end or "").strip()[:32]
    note = str(req.note or "").strip()[:500]

    dashboard_payload = await get_teaching_class_dashboard(
        class_id=class_id,
        user=user,
        assignment_limit=120,
    )

    archive_payload = {
        "class": dashboard_payload.get("class") or {},
        "overview": dashboard_payload.get("overview") or {},
        "class_goal_summary": dashboard_payload.get("class_goal_summary") or {},
        "class_goal_metric_values": dashboard_payload.get("class_goal_metric_values") or {},
        "class_goal_trend_30d": dashboard_payload.get("class_goal_trend_30d") or [],
        "class_goal_risk_summary": dashboard_payload.get("class_goal_risk_summary") or {},
        "class_goal_trend_overview": dashboard_payload.get("class_goal_trend_overview") or {},
        "intervention_action_summary": dashboard_payload.get("intervention_action_summary") or {},
        "intervention_outcome_correlation": dashboard_payload.get("intervention_outcome_correlation") or {},
        "intervention_strategy_summary": dashboard_payload.get("intervention_strategy_summary") or {},
        "intervention_strategy_recommendations": dashboard_payload.get("intervention_strategy_recommendations") or [],
        "intervention_experiment_summary": dashboard_payload.get("intervention_experiment_summary") or {},
        "intervention_experiments": dashboard_payload.get("intervention_experiments") or [],
        "goals": dashboard_payload.get("class_goals") or [],
        "archived_at": datetime.now(timezone.utc).isoformat(),
    }

    snapshot_date = datetime.now(timezone.utc).date().isoformat()
    await db.execute(
        """
        INSERT INTO teaching_goal_term_archives
            (class_id, teacher_user_id, term_code, term_name, term_start, term_end,
             note, archive_json, snapshot_date, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        ON CONFLICT(class_id, teacher_user_id, term_code)
        DO UPDATE SET
            term_name = excluded.term_name,
            term_start = excluded.term_start,
            term_end = excluded.term_end,
            note = excluded.note,
            archive_json = excluded.archive_json,
            snapshot_date = excluded.snapshot_date,
            updated_at = CURRENT_TIMESTAMP
        """,
        (
            class_id,
            int(user["id"]),
            term_code,
            term_name,
            term_start,
            term_end,
            note,
            json.dumps(archive_payload, ensure_ascii=False),
            snapshot_date,
        ),
    )
    await db.commit()

    row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, term_code, term_name, term_start, term_end,
               note, archive_json, snapshot_date, created_at, updated_at
        FROM teaching_goal_term_archives
        WHERE class_id = ? AND teacher_user_id = ? AND term_code = ?
        LIMIT 1
        """,
        (class_id, int(user["id"]), term_code),
    )
    if not row:
        raise HTTPException(status_code=500, detail="学期归档保存失败")

    archive_json = _parse_teaching_template_payload(row["archive_json"])
    return {
        "archive": {
            "id": int(row["id"] or 0),
            "class_id": int(row["class_id"] or 0),
            "teacher_user_id": int(row["teacher_user_id"] or 0),
            "term_code": str(row["term_code"] or ""),
            "term_name": str(row["term_name"] or ""),
            "term_start": str(row["term_start"] or ""),
            "term_end": str(row["term_end"] or ""),
            "note": str(row["note"] or ""),
            "snapshot_date": str(row["snapshot_date"] or ""),
            "created_at": row["created_at"] or "",
            "updated_at": row["updated_at"] or "",
            "goal_summary": archive_json.get("class_goal_summary") or {},
            "risk_summary": archive_json.get("class_goal_risk_summary") or {},
            "strategy_summary": archive_json.get("intervention_strategy_summary") or {},
        }
    }


@conv_router.get("/teaching/classes/{class_id}/goals/term-archives")
async def list_teaching_goal_term_archives(
    class_id: int,
    user: dict = Depends(get_current_user),
    limit: int = Query(20, ge=1, le=100),
    offset: int = Query(0, ge=0),
    include_detail: bool = Query(False, description="是否返回归档明细，仅教师本人可用"),
):
    """查看课堂学期归档记录。"""
    db = await get_db()
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")

    account_role = await _resolve_requester_account_role(db, user)
    requester_id = int(user["id"])
    teacher_owner_id = int(class_row["teacher_user_id"] or 0)

    if account_role == "teacher":
        if requester_id != teacher_owner_id:
            raise HTTPException(status_code=403, detail="无权查看该班级学期归档")
    elif account_role == "student":
        in_class = await _is_teaching_class_member(db, class_id=class_id, student_user_id=requester_id)
        if not in_class:
            raise HTTPException(status_code=403, detail="无权查看该班级学期归档")
        include_detail = False
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看学期归档")

    rows = await db.execute_fetchall(
        """
        SELECT id, class_id, teacher_user_id, term_code, term_name, term_start, term_end,
               note, archive_json, snapshot_date, created_at, updated_at
        FROM teaching_goal_term_archives
        WHERE class_id = ? AND teacher_user_id = ?
        ORDER BY snapshot_date DESC, id DESC
        LIMIT ? OFFSET ?
        """,
        (class_id, teacher_owner_id, limit, offset),
    )
    total_row = await db.execute_fetchone(
        """
        SELECT COUNT(1) AS cnt
        FROM teaching_goal_term_archives
        WHERE class_id = ? AND teacher_user_id = ?
        """,
        (class_id, teacher_owner_id),
    )

    can_include_detail = bool(include_detail and account_role == "teacher" and requester_id == teacher_owner_id)
    archives: list[dict[str, Any]] = []
    for row in rows:
        archive_json = _parse_teaching_template_payload(row["archive_json"])
        item = {
            "id": int(row["id"] or 0),
            "class_id": int(row["class_id"] or 0),
            "teacher_user_id": int(row["teacher_user_id"] or 0),
            "term_code": str(row["term_code"] or ""),
            "term_name": str(row["term_name"] or ""),
            "term_start": str(row["term_start"] or ""),
            "term_end": str(row["term_end"] or ""),
            "note": str(row["note"] or ""),
            "snapshot_date": str(row["snapshot_date"] or ""),
            "created_at": row["created_at"] or "",
            "updated_at": row["updated_at"] or "",
            "goal_summary": archive_json.get("class_goal_summary") or {},
            "risk_summary": archive_json.get("class_goal_risk_summary") or {},
            "strategy_summary": archive_json.get("intervention_strategy_summary") or {},
        }
        if can_include_detail:
            item["archive_detail"] = archive_json
        archives.append(item)

    return {
        "class_id": class_id,
        "archives": archives,
        "total": int(total_row["cnt"] or 0) if total_row else len(archives),
        "limit": limit,
        "offset": offset,
    }


@conv_router.get("/teaching/classes/{class_id}/goals/term-archives/compare")
async def compare_teaching_goal_term_archives(
    class_id: int,
    user: dict = Depends(get_current_user),
    base: str = Query("", description="基线学期编码"),
    target: str = Query("", description="目标学期编码"),
):
    """对比两个学期归档的目标、风险、关键指标与策略差异。"""
    db = await get_db()
    class_row = await _fetch_teaching_class(db, class_id)
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")

    account_role = await _resolve_requester_account_role(db, user)
    requester_id = int(user["id"])
    teacher_owner_id = int(class_row["teacher_user_id"] or 0)

    if account_role == "teacher":
        if requester_id != teacher_owner_id:
            raise HTTPException(status_code=403, detail="无权查看该班级学期归档对比")
    elif account_role == "student":
        in_class = await _is_teaching_class_member(db, class_id=class_id, student_user_id=requester_id)
        if not in_class:
            raise HTTPException(status_code=403, detail="无权查看该班级学期归档对比")
    else:
        raise HTTPException(status_code=403, detail="仅教师或学生账号可查看学期归档对比")

    base_term_code = _normalize_teaching_term_code(base)
    target_term_code = _normalize_teaching_term_code(target)
    if base_term_code == target_term_code:
        raise HTTPException(status_code=400, detail="base 与 target 不能相同")

    base_row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, term_code, term_name, term_start, term_end,
               note, archive_json, snapshot_date, created_at, updated_at
        FROM teaching_goal_term_archives
        WHERE class_id = ? AND teacher_user_id = ? AND term_code = ?
        LIMIT 1
        """,
        (class_id, teacher_owner_id, base_term_code),
    )
    if not base_row:
        raise HTTPException(status_code=404, detail=f"base 学期归档不存在: {base_term_code}")

    target_row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, term_code, term_name, term_start, term_end,
               note, archive_json, snapshot_date, created_at, updated_at
        FROM teaching_goal_term_archives
        WHERE class_id = ? AND teacher_user_id = ? AND term_code = ?
        LIMIT 1
        """,
        (class_id, teacher_owner_id, target_term_code),
    )
    if not target_row:
        raise HTTPException(status_code=404, detail=f"target 学期归档不存在: {target_term_code}")

    base_archive = _parse_teaching_template_payload(base_row["archive_json"])
    target_archive = _parse_teaching_template_payload(target_row["archive_json"])

    def _to_number(value: Any) -> float:
        try:
            num = float(value)
            return num if num == num else 0.0
        except Exception:
            return 0.0

    def _delta_summary(
        fields: list[str],
        base_obj: dict[str, Any],
        target_obj: dict[str, Any],
        *,
        as_int: bool,
    ) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for field in fields:
            before = _to_number(base_obj.get(field))
            after = _to_number(target_obj.get(field))
            delta = after - before
            if as_int:
                result[field] = int(round(delta))
            else:
                result[field] = round(delta, 4)
        return result

    base_goal_summary = base_archive.get("class_goal_summary") if isinstance(base_archive.get("class_goal_summary"), dict) else {}
    target_goal_summary = target_archive.get("class_goal_summary") if isinstance(target_archive.get("class_goal_summary"), dict) else {}
    base_risk_summary = base_archive.get("class_goal_risk_summary") if isinstance(base_archive.get("class_goal_risk_summary"), dict) else {}
    target_risk_summary = target_archive.get("class_goal_risk_summary") if isinstance(target_archive.get("class_goal_risk_summary"), dict) else {}
    base_metric_values = base_archive.get("class_goal_metric_values") if isinstance(base_archive.get("class_goal_metric_values"), dict) else {}
    target_metric_values = target_archive.get("class_goal_metric_values") if isinstance(target_archive.get("class_goal_metric_values"), dict) else {}
    base_strategy_summary = base_archive.get("intervention_strategy_summary") if isinstance(base_archive.get("intervention_strategy_summary"), dict) else {}
    target_strategy_summary = target_archive.get("intervention_strategy_summary") if isinstance(target_archive.get("intervention_strategy_summary"), dict) else {}

    goal_delta = _delta_summary(
        ["total_goal_count", "achieved_goal_count", "active_goal_count", "overdue_goal_count"],
        base_goal_summary,
        target_goal_summary,
        as_int=True,
    )
    risk_delta = _delta_summary(
        ["high_risk_count", "medium_risk_count", "low_risk_count"],
        base_risk_summary,
        target_risk_summary,
        as_int=True,
    )
    metric_delta = _delta_summary(
        ["submission_rate", "review_completion_rate", "avg_score"],
        base_metric_values,
        target_metric_values,
        as_int=False,
    )
    strategy_delta = _delta_summary(
        ["recommendation_count", "p0_count", "p1_count", "p2_count"],
        base_strategy_summary,
        target_strategy_summary,
        as_int=True,
    )

    highlights: list[str] = []
    if int(goal_delta.get("achieved_goal_count") or 0) > 0:
        highlights.append("目标达成数提升")
    if int(goal_delta.get("overdue_goal_count") or 0) < 0:
        highlights.append("逾期目标减少")
    if int(risk_delta.get("high_risk_count") or 0) < 0:
        highlights.append("高风险目标减少")
    if float(metric_delta.get("submission_rate") or 0) > 0:
        highlights.append("提交率提升")
    if float(metric_delta.get("review_completion_rate") or 0) > 0:
        highlights.append("评阅完成率提升")
    if float(metric_delta.get("avg_score") or 0) > 0:
        highlights.append("平均分提升")
    if not highlights:
        highlights.append("暂无显著正向变化，建议继续跟踪。")

    return {
        "class_id": class_id,
        "base": {
            "term_code": str(base_row["term_code"] or ""),
            "term_name": str(base_row["term_name"] or ""),
            "snapshot_date": str(base_row["snapshot_date"] or ""),
        },
        "target": {
            "term_code": str(target_row["term_code"] or ""),
            "term_name": str(target_row["term_name"] or ""),
            "snapshot_date": str(target_row["snapshot_date"] or ""),
        },
        "delta": {
            "goal_summary": goal_delta,
            "risk_summary": risk_delta,
            "metric_values": metric_delta,
            "strategy_summary": strategy_delta,
            "highlights": highlights,
        },
    }


@conv_router.post("/teaching/classes/{class_id}/interventions/actions")
async def create_teaching_intervention_action(
    class_id: int,
    req: TeachingInterventionActionCreateRequest,
    user: dict = Depends(get_current_user),
):
    """教师记录课堂干预建议的执行动作。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可记录干预执行")

    await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    intervention_code = _normalize_teaching_intervention_code(req.intervention_code)
    intervention_title = str(req.intervention_title or "").strip()[:120]
    intervention_severity = _normalize_teaching_intervention_severity(req.intervention_severity)
    note = str(req.note or "").strip()[:1200]
    assignment_id = int(req.assignment_id or 0)

    if assignment_id > 0:
        assignment_row = await db.execute_fetchone(
            """
            SELECT id
            FROM teaching_assignments
            WHERE id = ? AND class_id = ? AND teacher_user_id = ?
            """,
            (assignment_id, class_id, int(user["id"])),
        )
        if not assignment_row:
            raise HTTPException(status_code=400, detail="assignment_id 不属于当前班级")

    metadata = dict(req.metadata or {}) if isinstance(req.metadata, dict) else {}
    strategy_variant = _normalize_teaching_strategy_variant_token(req.strategy_variant)
    raw_strategy_experiment_id = str(req.strategy_experiment_id or "").strip()
    strategy_experiment_id = (
        _normalize_teaching_experiment_code(
            raw_strategy_experiment_id,
            fallback_intervention_code=intervention_code,
        )
        if raw_strategy_experiment_id
        else ""
    )
    strategy_note = str(req.strategy_note or "").strip()[:200]
    if strategy_variant:
        metadata["strategy_variant"] = strategy_variant
    if strategy_experiment_id:
        metadata["strategy_experiment_id"] = strategy_experiment_id
    if strategy_note:
        metadata["strategy_note"] = strategy_note

    cur = await db.execute(
        """
        INSERT INTO teaching_intervention_actions
            (class_id, teacher_user_id, intervention_code, intervention_title, intervention_severity,
             assignment_id, note, metadata_json, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        """,
        (
            class_id,
            int(user["id"]),
            intervention_code,
            intervention_title,
            intervention_severity,
            assignment_id,
            note,
            json.dumps(metadata, ensure_ascii=False),
        ),
    )
    await db.commit()

    action_row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, intervention_code, intervention_title, intervention_severity,
               assignment_id, note, metadata_json, created_at
        FROM teaching_intervention_actions
        WHERE id = ?
        """,
        (int(cur.lastrowid or 0),),
    )
    if not action_row:
        raise HTTPException(status_code=500, detail="干预执行记录写入失败")

    summary = await _build_teaching_intervention_action_summary(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )
    return {
        "action": _serialize_teaching_intervention_action_row(action_row),
        "summary": summary,
    }


@conv_router.get("/teaching/classes/{class_id}/interventions/actions")
async def list_teaching_intervention_actions(
    class_id: int,
    user: dict = Depends(get_current_user),
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
):
    """教师查看课堂干预执行历史。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可查看干预执行历史")

    await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    rows = await db.execute_fetchall(
        """
        SELECT id, class_id, teacher_user_id, intervention_code, intervention_title, intervention_severity,
               assignment_id, note, metadata_json, created_at
        FROM teaching_intervention_actions
        WHERE class_id = ? AND teacher_user_id = ?
        ORDER BY created_at DESC, id DESC
        LIMIT ? OFFSET ?
        """,
        (class_id, int(user["id"]), limit, offset),
    )
    total_row = await db.execute_fetchone(
        """
        SELECT COUNT(1) AS cnt
        FROM teaching_intervention_actions
        WHERE class_id = ? AND teacher_user_id = ?
        """,
        (class_id, int(user["id"])),
    )
    summary = await _build_teaching_intervention_action_summary(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    return {
        "class_id": class_id,
        "actions": [_serialize_teaching_intervention_action_row(row) for row in rows],
        "total": int(total_row["cnt"] or 0) if total_row else len(rows),
        "limit": limit,
        "offset": offset,
        "summary": summary,
    }


@conv_router.post("/teaching/templates")
async def create_teaching_template(
    req: TeachingTemplateCreateRequest,
    user: dict = Depends(get_current_user),
):
    """教师创建实训模板（班级/作业/评分量表）。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可创建模板")

    template_type = _normalize_teaching_template_type(req.template_type)
    name = str(req.name or "").strip()[:120]
    if not name:
        raise HTTPException(status_code=400, detail="模板名称不能为空")

    description = str(req.description or "").strip()[:1000]
    payload = req.payload if isinstance(req.payload, dict) else {}
    payload_json = json.dumps(payload, ensure_ascii=False)
    initial_status = _normalize_teaching_template_status(req.initial_status, allow_active_alias=True)

    existing_row = await db.execute_fetchone(
        """
        SELECT id FROM teaching_templates
        WHERE teacher_user_id = ? AND template_type = ? AND name = ?
        LIMIT 1
        """,
        (user["id"], template_type, name),
    )

    if existing_row:
        change_type = "update"
        await db.execute(
            """
            UPDATE teaching_templates
            SET description = ?, payload_json = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (description, payload_json, int(existing_row["id"])),
        )
    else:
        change_type = "create"
        await db.execute(
            """
            INSERT INTO teaching_templates
                (teacher_user_id, template_type, name, description, payload_json, status, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
            """,
            (user["id"], template_type, name, description, payload_json, initial_status),
        )

    row = await db.execute_fetchone(
        """
        SELECT id, teacher_user_id, template_type, name, description, payload_json, status, created_at, updated_at
        FROM teaching_templates
        WHERE teacher_user_id = ? AND template_type = ? AND name = ?
        """,
        (user["id"], template_type, name),
    )
    if not row:
        raise HTTPException(status_code=500, detail="模板创建失败")

    latest_version_no = await _append_teaching_template_version(
        db,
        template_row=row,
        change_type=change_type,
        change_note="create_or_upsert",
    )
    await db.commit()

    usage_meta_map = await _list_teaching_template_usage_meta(
        db,
        teacher_user_id=int(user["id"]),
        template_ids=[int(row["id"] or 0)],
    )
    usage_meta = usage_meta_map.get(int(row["id"] or 0), {})

    return _serialize_teaching_template_row(
        row,
        latest_version_no=latest_version_no,
        version_count=latest_version_no,
        usage_count=int(usage_meta.get("usage_count") or 0),
        class_usage_count=int(usage_meta.get("class_usage_count") or 0),
        assignment_usage_count=int(usage_meta.get("assignment_usage_count") or 0),
        usage_recent_7d=int(usage_meta.get("usage_recent_7d") or 0),
        last_used_at=str(usage_meta.get("last_used_at") or ""),
    )


@conv_router.get("/teaching/templates")
async def list_teaching_templates(
    user: dict = Depends(get_current_user),
    template_type: str = Query("", description="模板类型 class/assignment/rubric"),
    limit: int = Query(80, ge=1, le=300),
):
    """教师查看自己的模板列表。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可查看模板")

    type_filter = str(template_type or "").strip().lower()
    params: list[Any] = [user["id"]]
    sql = """
        SELECT id, teacher_user_id, template_type, name, description, payload_json, status, created_at, updated_at
        FROM teaching_templates
        WHERE teacher_user_id = ?
    """
    if type_filter:
        type_filter = _normalize_teaching_template_type(type_filter)
        sql += " AND template_type = ?"
        params.append(type_filter)

    sql += " ORDER BY updated_at DESC LIMIT ?"
    params.append(limit)

    rows = await db.execute_fetchall(sql, tuple(params))
    version_meta_map = await _list_teaching_template_version_meta(db, teacher_user_id=int(user["id"]))
    usage_meta_map = await _list_teaching_template_usage_meta(
        db,
        teacher_user_id=int(user["id"]),
        template_ids=[int(row["id"] or 0) for row in rows],
    )

    templates = []
    for row in rows:
        tid = int(row["id"] or 0)
        meta = version_meta_map.get(tid, {})
        usage_meta = usage_meta_map.get(tid, {})
        templates.append(
            _serialize_teaching_template_row(
                row,
                latest_version_no=int(meta.get("latest_version_no") or 0),
                version_count=int(meta.get("version_count") or 0),
                usage_count=int(usage_meta.get("usage_count") or 0),
                class_usage_count=int(usage_meta.get("class_usage_count") or 0),
                assignment_usage_count=int(usage_meta.get("assignment_usage_count") or 0),
                usage_recent_7d=int(usage_meta.get("usage_recent_7d") or 0),
                last_used_at=str(usage_meta.get("last_used_at") or ""),
            )
        )

    return {"templates": templates}


@conv_router.get("/teaching/templates/insights")
async def get_teaching_template_insights(
    user: dict = Depends(get_current_user),
    template_type: str = Query("", description="模板类型 class/assignment/rubric"),
    top_limit: int = Query(8, ge=1, le=30),
):
    """教师查看模板运营洞察（使用统计 + 治理建议）。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可查看模板洞察")

    type_filter = str(template_type or "").strip().lower()
    params: list[Any] = [user["id"]]
    sql = """
        SELECT id, teacher_user_id, template_type, name, description, payload_json, status, created_at, updated_at
        FROM teaching_templates
        WHERE teacher_user_id = ?
    """
    if type_filter:
        type_filter = _normalize_teaching_template_type(type_filter)
        sql += " AND template_type = ?"
        params.append(type_filter)

    sql += " ORDER BY updated_at DESC"
    rows = await db.execute_fetchall(sql, tuple(params))

    usage_meta_map = await _list_teaching_template_usage_meta(
        db,
        teacher_user_id=int(user["id"]),
        template_ids=[int(row["id"] or 0) for row in rows],
    )
    assignment_template_ids = [
        int(row["id"] or 0)
        for row in rows
        if str(row["template_type"] or "").strip().lower() == "assignment"
    ]
    effect_snapshot_map = await _build_assignment_template_effect_snapshots(
        db,
        teacher_user_id=int(user["id"]),
        template_ids=assignment_template_ids,
    )

    template_items: list[dict[str, Any]] = []
    status_counter: Counter[str] = Counter()
    usage_count_total = 0
    usage_recent_7d_total = 0
    used_template_count = 0
    last_used_at = ""
    effect_template_count = 0
    effect_assignment_count_total = 0

    for row in rows:
        tid = int(row["id"] or 0)
        usage_meta = usage_meta_map.get(tid, {})
        usage_count = int(usage_meta.get("usage_count") or 0)
        class_usage_count = int(usage_meta.get("class_usage_count") or 0)
        assignment_usage_count = int(usage_meta.get("assignment_usage_count") or 0)
        usage_recent_7d = int(usage_meta.get("usage_recent_7d") or 0)
        used_at = str(usage_meta.get("last_used_at") or "")

        status_counter[_canonical_teaching_template_status(row["status"])] += 1
        usage_count_total += usage_count
        usage_recent_7d_total += usage_recent_7d
        if usage_count > 0:
            used_template_count += 1
        if used_at and (not last_used_at or used_at > last_used_at):
            last_used_at = used_at

        template_type_token = str(row["template_type"] or "").strip().lower()
        effect_snapshot = effect_snapshot_map.get(tid, {}) if template_type_token == "assignment" else {}
        if effect_snapshot:
            effect_template_count += 1
            effect_assignment_count_total += int(effect_snapshot.get("assignment_count") or 0)

        template_items.append(
            _serialize_teaching_template_row(
                row,
                usage_count=usage_count,
                class_usage_count=class_usage_count,
                assignment_usage_count=assignment_usage_count,
                usage_recent_7d=usage_recent_7d,
                last_used_at=used_at,
                effect_snapshot=effect_snapshot,
            )
        )

    total_templates = len(template_items)
    unused_template_count = max(total_templates - used_template_count, 0)
    publishable_template_count = int(status_counter.get("approved", 0))

    suggestions: list[dict[str, Any]] = []
    if status_counter.get("review", 0) > 0:
        suggestions.append(
            {
                "code": "pending_approval_backlog",
                "severity": "high",
                "title": "存在待审核模板积压",
                "detail": f"当前有 {int(status_counter.get('review', 0))} 个模板处于待审核。",
                "action": "建议优先完成模板审核，减少教学执行时的临时改稿。",
            }
        )
    if status_counter.get("draft", 0) > 0:
        suggestions.append(
            {
                "code": "draft_backlog",
                "severity": "medium",
                "title": "草稿模板较多",
                "detail": f"当前有 {int(status_counter.get('draft', 0))} 个草稿模板未进入发布链路。",
                "action": "建议按课程节奏清理或提审草稿模板，保持模板池可用性。",
            }
        )
    if total_templates >= 4 and unused_template_count >= max(2, (total_templates // 2)):
        suggestions.append(
            {
                "code": "low_template_activation",
                "severity": "medium",
                "title": "模板激活率偏低",
                "detail": f"已使用模板 {used_template_count}/{total_templates}，存在较多未使用模板。",
                "action": "建议下线长期未使用模板或在课堂中安排演练，提高模板复用效率。",
            }
        )
    if usage_count_total > 0 and usage_recent_7d_total <= 0:
        suggestions.append(
            {
                "code": "template_usage_cooling",
                "severity": "medium",
                "title": "近7天模板使用趋冷",
                "detail": "模板历史有使用记录，但最近7天没有新的套用。",
                "action": "建议检查教学计划变更，并提前准备下周高频模板。",
            }
        )

    top_templates = sorted(
        template_items,
        key=lambda item: (
            int(item.get("usage_count") or 0),
            int(item.get("usage_recent_7d") or 0),
            str(item.get("updated_at") or ""),
        ),
        reverse=True,
    )[: int(top_limit)]

    return {
        "template_type": type_filter or "all",
        "summary": {
            "total_templates": total_templates,
            "publishable_template_count": publishable_template_count,
            "draft_count": int(status_counter.get("draft", 0)),
            "review_count": int(status_counter.get("review", 0)),
            "approved_count": int(status_counter.get("approved", 0)),
            "used_template_count": used_template_count,
            "unused_template_count": unused_template_count,
            "usage_count_total": usage_count_total,
            "usage_recent_7d_total": usage_recent_7d_total,
            "last_used_at": last_used_at,
            "effect_template_count": effect_template_count,
            "effect_assignment_count_total": effect_assignment_count_total,
        },
        "top_templates": top_templates,
        "suggestions": suggestions,
    }


@conv_router.delete("/teaching/templates/{template_id}")
async def delete_teaching_template(
    template_id: int,
    user: dict = Depends(get_current_user),
):
    """教师删除自己的模板。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可删除模板")

    template_row = await _require_teacher_owned_template(
        db,
        template_id=template_id,
        teacher_user_id=int(user["id"]),
    )

    await db.execute("DELETE FROM teaching_templates WHERE id = ?", (template_id,))
    await db.commit()

    return {
        "id": template_id,
        "template_type": template_row["template_type"],
        "deleted": True,
    }


@conv_router.post("/teaching/templates/{template_id}/status")
async def update_teaching_template_status(
    template_id: int,
    req: TeachingTemplateStatusUpdateRequest,
    user: dict = Depends(get_current_user),
):
    """教师更新模板状态（draft/review/approved）。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可更新模板状态")

    template_row = await _require_teacher_owned_template(
        db,
        template_id=template_id,
        teacher_user_id=int(user["id"]),
    )

    current_canonical = _canonical_teaching_template_status(template_row["status"])
    target_status = _ensure_teaching_template_status_transition(template_row["status"], req.target_status)

    await db.execute(
        """
        UPDATE teaching_templates
        SET status = ?, updated_at = CURRENT_TIMESTAMP
        WHERE id = ? AND teacher_user_id = ?
        """,
        (target_status, template_id, int(user["id"])),
    )

    updated_row = await _fetch_teaching_template(db, template_id)
    if not updated_row:
        raise HTTPException(status_code=500, detail="模板状态更新失败")

    note = str(req.note or "").strip()
    if not note:
        note = f"status:{current_canonical}->{target_status}"

    latest_version_no = await _append_teaching_template_version(
        db,
        template_row=updated_row,
        change_type="status_change",
        change_note=note,
    )
    await db.commit()

    usage_meta_map = await _list_teaching_template_usage_meta(
        db,
        teacher_user_id=int(user["id"]),
        template_ids=[int(updated_row["id"] or 0)],
    )
    usage_meta = usage_meta_map.get(int(updated_row["id"] or 0), {})

    return {
        "template": _serialize_teaching_template_row(
            updated_row,
            latest_version_no=latest_version_no,
            version_count=latest_version_no,
            usage_count=int(usage_meta.get("usage_count") or 0),
            class_usage_count=int(usage_meta.get("class_usage_count") or 0),
            assignment_usage_count=int(usage_meta.get("assignment_usage_count") or 0),
            usage_recent_7d=int(usage_meta.get("usage_recent_7d") or 0),
            last_used_at=str(usage_meta.get("last_used_at") or ""),
        ),
        "transition": {
            "from": current_canonical,
            "to": target_status,
        },
    }

@conv_router.get("/teaching/templates/{template_id}/versions")
async def list_teaching_template_versions(
    template_id: int,
    user: dict = Depends(get_current_user),
    limit: int = Query(50, ge=1, le=300),
):
    """教师查看模板版本历史。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可查看模板版本")

    template_row = await _require_teacher_owned_template(
        db,
        template_id=template_id,
        teacher_user_id=int(user["id"]),
    )

    rows = await db.execute_fetchall(
        """
        SELECT id, template_id, teacher_user_id, template_type, version_no, change_type, change_note,
               source_version_id, name, description, payload_json, created_at
        FROM teaching_template_versions
        WHERE template_id = ? AND teacher_user_id = ?
        ORDER BY version_no DESC
        LIMIT ?
        """,
        (template_id, int(user["id"]), limit),
    )

    versions = []
    for row in rows:
        versions.append(
            {
                "id": row["id"],
                "template_id": row["template_id"],
                "teacher_user_id": row["teacher_user_id"],
                "template_type": row["template_type"],
                "version_no": int(row["version_no"] or 0),
                "change_type": row["change_type"] or "",
                "change_note": row["change_note"] or "",
                "source_version_id": row["source_version_id"],
                "name": row["name"] or "",
                "description": row["description"] or "",
                "payload": _parse_teaching_template_payload(row["payload_json"]),
                "created_at": row["created_at"],
            }
        )

    count_row = await db.execute_fetchone(
        """
        SELECT COUNT(1) AS total
        FROM teaching_template_versions
        WHERE template_id = ? AND teacher_user_id = ?
        """,
        (template_id, int(user["id"])),
    )
    total = int(count_row["total"] or 0) if count_row else len(versions)
    latest_version_no = int(versions[0]["version_no"] or 0) if versions else 0
    usage_meta_map = await _list_teaching_template_usage_meta(
        db,
        teacher_user_id=int(user["id"]),
        template_ids=[int(template_row["id"] or 0)],
    )
    usage_meta = usage_meta_map.get(int(template_row["id"] or 0), {})

    return {
        "template": _serialize_teaching_template_row(
            template_row,
            latest_version_no=latest_version_no,
            version_count=total,
            usage_count=int(usage_meta.get("usage_count") or 0),
            class_usage_count=int(usage_meta.get("class_usage_count") or 0),
            assignment_usage_count=int(usage_meta.get("assignment_usage_count") or 0),
            usage_recent_7d=int(usage_meta.get("usage_recent_7d") or 0),
            last_used_at=str(usage_meta.get("last_used_at") or ""),
        ),
        "versions": versions,
        "total": total,
    }


@conv_router.post("/teaching/templates/{template_id}/rollback")
async def rollback_teaching_template(
    template_id: int,
    req: TeachingTemplateRollbackRequest,
    user: dict = Depends(get_current_user),
):
    """教师将模板回滚到指定历史版本。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可回滚模板")

    template_row = await _require_teacher_owned_template(
        db,
        template_id=template_id,
        teacher_user_id=int(user["id"]),
    )

    version_id = int(req.version_id or 0)
    version_no = int(req.version_no or 0)
    if version_id <= 0 and version_no <= 0:
        raise HTTPException(status_code=400, detail="请提供 version_id 或 version_no")

    if version_id > 0:
        target = await db.execute_fetchone(
            """
            SELECT id, template_id, teacher_user_id, template_type, version_no, name, description, payload_json, created_at
            FROM teaching_template_versions
            WHERE template_id = ? AND teacher_user_id = ? AND id = ?
            LIMIT 1
            """,
            (template_id, int(user["id"]), version_id),
        )
    else:
        target = await db.execute_fetchone(
            """
            SELECT id, template_id, teacher_user_id, template_type, version_no, name, description, payload_json, created_at
            FROM teaching_template_versions
            WHERE template_id = ? AND teacher_user_id = ? AND version_no = ?
            LIMIT 1
            """,
            (template_id, int(user["id"]), version_no),
        )

    if not target:
        raise HTTPException(status_code=404, detail="目标模板版本不存在")

    try:
        await db.execute(
            """
            UPDATE teaching_templates
            SET
                template_type = ?,
                name = ?,
                description = ?,
                payload_json = ?,
                status = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND teacher_user_id = ?
            """,
            (
                str(target["template_type"] or "").strip().lower(),
                str(target["name"] or "")[:120],
                str(target["description"] or "")[:1000],
                str(target["payload_json"] or "{}"),
                _normalize_teaching_template_status(template_row["status"], allow_active_alias=True),
                template_id,
                int(user["id"]),
            ),
        )
    except Exception as e:
        msg = str(e)
        if "UNIQUE" in msg.upper():
            raise HTTPException(status_code=409, detail="回滚失败：目标名称与现有模板冲突")
        raise HTTPException(status_code=500, detail=f"回滚失败: {msg}")

    updated_template = await _fetch_teaching_template(db, template_id)
    if not updated_template:
        raise HTTPException(status_code=500, detail="回滚后模板不存在")

    rollback_note = str(req.note or "").strip()
    if not rollback_note:
        rollback_note = f"rollback_to_v{int(target['version_no'] or 0)}"

    latest_version_no = await _append_teaching_template_version(
        db,
        template_row=updated_template,
        change_type="rollback",
        change_note=rollback_note,
        source_version_id=int(target["id"]),
    )
    await db.commit()

    usage_meta_map = await _list_teaching_template_usage_meta(
        db,
        teacher_user_id=int(user["id"]),
        template_ids=[int(updated_template["id"] or 0)],
    )
    usage_meta = usage_meta_map.get(int(updated_template["id"] or 0), {})

    return {
        "template": _serialize_teaching_template_row(
            updated_template,
            latest_version_no=latest_version_no,
            version_count=latest_version_no,
            usage_count=int(usage_meta.get("usage_count") or 0),
            class_usage_count=int(usage_meta.get("class_usage_count") or 0),
            assignment_usage_count=int(usage_meta.get("assignment_usage_count") or 0),
            usage_recent_7d=int(usage_meta.get("usage_recent_7d") or 0),
            last_used_at=str(usage_meta.get("last_used_at") or ""),
        ),
        "rollback": {
            "target_version_id": int(target["id"]),
            "target_version_no": int(target["version_no"] or 0),
            "new_version_no": latest_version_no,
        },
    }


@conv_router.post("/teaching/templates/{template_id}/instantiate-class")
async def instantiate_teaching_class_from_template(
    template_id: int,
    req: TeachingClassTemplateInstantiateRequest,
    user: dict = Depends(get_current_user),
):
    """按班级模板创建班级。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可套用班级模板")

    template_row = await _require_teacher_owned_template(
        db,
        template_id=template_id,
        teacher_user_id=int(user["id"]),
    )
    if str(template_row["template_type"] or "").strip().lower() != "class":
        raise HTTPException(status_code=400, detail="仅 class 模板可用于创建班级")
    if not _is_teaching_template_publishable(template_row["status"]):
        raise HTTPException(status_code=409, detail="模板未发布，请先将模板状态切换为 approved")

    payload = _parse_teaching_template_payload(template_row["payload_json"])
    base_name = str(payload.get("name") or payload.get("class_name") or "").strip()
    base_description = str(payload.get("description") or payload.get("class_description") or "").strip()

    name = str(req.name_override or "").strip() or base_name
    description = str(req.description_override or "").strip() or base_description

    name = name[:120]
    description = description[:1000]
    if not name:
        raise HTTPException(status_code=400, detail="模板未提供有效班级名称，请在模板 payload 中配置 name")

    await db.execute(
        """
        INSERT INTO teaching_classes
            (teacher_user_id, name, description, status, created_at, updated_at)
        VALUES (?, ?, ?, 'active', CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        ON CONFLICT(teacher_user_id, name)
        DO UPDATE SET
            description = excluded.description,
            status = 'active',
            updated_at = CURRENT_TIMESTAMP
        """,
        (user["id"], name, description),
    )
    await db.commit()

    class_row = await db.execute_fetchone(
        """
        SELECT id, teacher_user_id, name, description, status, created_at, updated_at
        FROM teaching_classes
        WHERE teacher_user_id = ? AND name = ?
        """,
        (user["id"], name),
    )
    if not class_row:
        raise HTTPException(status_code=500, detail="班级模板套用失败")

    try:
        await _append_teaching_template_usage_log(
            db,
            template_row=template_row,
            usage_scene="class_instantiate",
            target_id=int(class_row["id"] or 0),
            metadata={
                "class_name": class_row["name"] or "",
                "class_id": int(class_row["id"] or 0),
            },
        )
    except Exception as e:
        logger.warning("failed to append class template usage log: %s", e)

    await db.commit()

    return {
        "template_id": template_id,
        "class": {
            "id": class_row["id"],
            "teacher_user_id": class_row["teacher_user_id"],
            "name": class_row["name"] or "",
            "description": class_row["description"] or "",
            "status": class_row["status"],
            "created_at": class_row["created_at"],
            "updated_at": class_row["updated_at"],
        },
    }


@conv_router.post("/teaching/classes/{class_id}/templates/{template_id}/assignments/apply")
async def apply_teaching_assignment_template_to_class(
    class_id: int,
    template_id: int,
    req: TeachingAssignmentTemplateApplyRequest,
    user: dict = Depends(get_current_user),
):
    """将作业模板应用到指定班级。"""
    db = await get_db()
    account_role = await _resolve_requester_account_role(db, user)
    if account_role != "teacher":
        raise HTTPException(status_code=403, detail="仅教师账号可套用作业模板")

    class_row = await _require_teacher_owned_class(
        db,
        class_id=class_id,
        teacher_user_id=int(user["id"]),
    )

    template_row = await _require_teacher_owned_template(
        db,
        template_id=template_id,
        teacher_user_id=int(user["id"]),
    )
    if str(template_row["template_type"] or "").strip().lower() != "assignment":
        raise HTTPException(status_code=400, detail="仅 assignment 模板可用于创建作业")
    if not _is_teaching_template_publishable(template_row["status"]):
        raise HTTPException(status_code=409, detail="模板未发布，请先将模板状态切换为 approved")

    payload = _parse_teaching_template_payload(template_row["payload_json"])
    title = str(req.title_override or "").strip() or str(payload.get("title") or "").strip()
    description = str(req.description_override or "").strip() or str(payload.get("description") or "").strip()
    due_at = str(req.due_at_override or "").strip() or str(payload.get("due_at") or "").strip()

    title = title[:160]
    description = description[:4000]
    due_at = due_at[:64]

    if not title:
        raise HTTPException(status_code=400, detail="模板未提供有效作业标题，请在模板 payload 中配置 title")

    cur = await db.execute(
        """
        INSERT INTO teaching_assignments
            (class_id, teacher_user_id, title, description, due_at, status, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, 'active', CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        """,
        (class_id, user["id"], title, description, due_at),
    )
    await db.commit()

    assignment_id = int(cur.lastrowid or 0)
    row = await db.execute_fetchone(
        """
        SELECT id, class_id, teacher_user_id, title, description, due_at, status, created_at, updated_at
        FROM teaching_assignments
        WHERE id = ?
        """,
        (assignment_id,),
    )
    if not row:
        raise HTTPException(status_code=500, detail="作业模板套用失败")

    try:
        await _append_teaching_template_usage_log(
            db,
            template_row=template_row,
            usage_scene="assignment_apply",
            target_id=int(row["id"] or 0),
            metadata={
                "class_id": int(class_row["id"] or 0),
                "assignment_id": int(row["id"] or 0),
                "assignment_title": row["title"] or "",
            },
        )
    except Exception as e:
        logger.warning("failed to append assignment template usage log: %s", e)

    await db.commit()

    return {
        "template_id": template_id,
        "class": {
            "id": class_row["id"],
            "name": class_row["name"] or "",
        },
        "assignment": {
            "id": row["id"],
            "class_id": row["class_id"],
            "teacher_user_id": row["teacher_user_id"],
            "title": row["title"] or "",
            "description": row["description"] or "",
            "due_at": row["due_at"] or "",
            "status": row["status"],
            "created_at": row["created_at"],
            "updated_at": row["updated_at"],
        },
    }



















